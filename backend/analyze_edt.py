import openpyxl, os, re

folder = '/Users/tobidesis/Downloads/données'

def find_edt_data(ws):
    all_rows = list(ws.iter_rows(values_only=True))
    grade_info = groupe_info = site_info = formation_info = ''
    header_row_idx = None

    for i, row in enumerate(all_rows):
        vals = [str(v).strip().replace('\xa0', ' ') if v else '' for v in row]
        full = ' '.join(vals).upper()

        if 'FORMATION EN' in full and not formation_info:
            formation_info = next((v for v in vals if 'FORMATION' in v.upper()), '')[:80]
        if 'EMPLOI DU TEMPS' in full and ('GRADE' in full or 'CATEGORIE' in full):
            grade_info = full[:100]
        if 'GROUPE' in full:
            for v in vals:
                if 'GROUPE' in v.upper() and not groupe_info:
                    groupe_info = v.strip()
        if 'SITE' in full or 'CPFAE' in full:
            for v in vals:
                if ('SITE' in v.upper() or 'CPFAE' in v.upper()) and not site_info:
                    site_info = v.strip()
        if 'MATIERES' in full or 'MATIERE' in full:
            header_row_idx = i
            break

    return header_row_idx, formation_info, grade_info, groupe_info, site_info, all_rows

for fname in sorted(os.listdir(folder)):
    if 'Emploi' not in fname:
        continue
    path = os.path.join(folder, fname)
    print(f"\n{'='*60}")
    print(f"FICHIER: {fname}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for shname in wb.sheetnames:
        ws = wb[shname]
        hi, formation, grade, groupe, site, all_rows = find_edt_data(ws)
        print(f"\n  Feuille '{shname}':")
        print(f"    Formation : {formation}")
        print(f"    Grade     : {grade[:80]}")
        print(f"    Groupe    : {groupe}")
        print(f"    Site      : {site}")
        print(f"    Header@L  : {hi+1 if hi is not None else 'NON TROUVE'}")

        if hi is not None:
            print(f"    En-tete   : {[str(v)[:22] for v in all_rows[hi] if v][:8]}")
            for j in range(hi+1, min(hi+6, len(all_rows))):
                row = all_rows[j]
                vals = [v for v in row if v is not None and str(v).strip()]
                if vals:
                    print(f"    Data L{j+1:02d} : {[str(v)[:25] for v in row if v is not None and str(v).strip()][:8]}")
    wb.close()
