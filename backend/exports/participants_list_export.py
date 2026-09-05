"""Export PDF / Excel — liste de classe auditeurs par groupe."""
import io
import re
from collections import defaultdict
from datetime import datetime

from django.db.models import Q

from formations.models import Participant

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except Exception:
    OPENPYXL_AVAILABLE = False


_HEADER_LINES = [
    "RÉPUBLIQUE DE CÔTE D'IVOIRE",
    "Ministère de la Fonction Publique et de la Modernisation de l'Administration",
    "Direction de la Formation et du Renforcement des Capacités (DFRC)",
    "Centre de Perfectionnement des Fonctionnaires et Agents de l'État — CPFAE",
]

_TABLE_HEADERS = ['N°', 'Matricule', 'Nom', 'Prénom', 'Sexe', 'Grade', 'Téléphone', 'Type concours', 'Moyenne', 'Temps', 'Décision']

_DECISION_LABELS = {
    'ADMIS': 'Admis',
    'AJOURNE': 'Ajourné',
    'EXCLUSION': 'Exclusion',
    'EN_ATTENTE': 'En attente',
}


def _decision_moyenne(decision):
    if decision is None or decision.moyenne_generale is None:
        return '—'
    return f"{decision.moyenne_generale}/20"


def _decision_temps(decision):
    if decision is None or decision.taux_presence is None:
        return '—'
    return f"{decision.taux_presence}%"


def _decision_label(decision):
    if decision is None:
        return '—'
    return _DECISION_LABELS.get(decision.decision, '—')


def _normalize_groupe_value(value):
    raw = str(value or '').strip()
    if not raw:
        return ''
    compact = re.sub(r'\s+', ' ', raw).strip()
    match = re.fullmatch(r'(?:GROUPE\s*)?0*(\d+)', compact, flags=re.IGNORECASE)
    if match:
        return f"GROUPE {int(match.group(1))}"
    return compact.upper()


def _groupe_sort_key(value):
    match = re.fullmatch(r'GROUPE (\d+)', value or '')
    if match:
        return (0, int(match.group(1)))
    return (1, value or '')


def _sexe_label(participant):
    labels = dict(Participant.Sexe.choices)
    return labels.get(participant.sexe, '—') if participant.sexe else '—'


def _participant_table_row(index, participant):
    decision = getattr(participant, '_decision', None)
    return [
        index,
        participant.matricule or '—',
        (participant.nom or '').upper(),
        (participant.prenom or '').title(),
        _sexe_label(participant),
        participant.grade or '—',
        participant.telephone or '—',
        participant.type_concours or '—',
        _decision_moyenne(decision),
        _decision_temps(decision),
        _decision_label(decision),
    ]


def _apply_participant_filters(queryset, params):
    """Applique les filtres query string (aligné sur participant_list_api)."""
    qs = queryset

    search = (params.get('search') or '').strip()
    if search:
        qs = qs.filter(
            Q(nom__icontains=search)
            | Q(prenom__icontains=search)
            | Q(matricule__icontains=search)
            | Q(email__icontains=search)
            | Q(type_concours__icontains=search)
            | Q(grade__icontains=search)
            | Q(groupe__icontains=search)
        )

    secretariat = (params.get('secretariat') or '').strip()
    if secretariat:
        qs = qs.filter(
            Q(secretariat__nom__icontains=secretariat)
            | Q(secretariat__type__libelle__icontains=secretariat)
        )

    grade = (params.get('grade') or '').strip()
    if grade:
        qs = qs.filter(grade__icontains=grade)

    type_concours = (params.get('type_concours') or '').strip()
    if type_concours:
        qs = qs.filter(type_concours__icontains=type_concours)

    vague = (params.get('vague') or '').strip()
    if vague:
        qs = qs.filter(vague__icontains=vague)

    sexe = (params.get('sexe') or '').strip()
    if sexe:
        qs = qs.filter(sexe=sexe)

    groupe = (params.get('groupe') or '').strip()
    if groupe:
        groupe_norm = _normalize_groupe_value(groupe)
        matching_ids = [
            pk
            for pk, participant_groupe in qs.values_list('id', 'groupe')
            if _normalize_groupe_value(participant_groupe) == groupe_norm
        ]
        qs = qs.filter(pk__in=matching_ids)

    return qs.select_related('secretariat').order_by('nom', 'prenom', 'matricule')


def _attach_decisions(participants):
    """Attache à chaque participant sa décision pédagogique la plus récente (_decision)."""
    if not participants:
        return
    try:
        from suiviEvaluation.models import DecisionPedagogique
    except Exception:
        return
    ids = [p.id for p in participants]
    latest = {}
    for dec in (
        DecisionPedagogique.objects
        .filter(participant_id__in=ids)
        .order_by('participant_id', '-updated_at')
    ):
        if dec.participant_id not in latest:
            latest[dec.participant_id] = dec
    for p in participants:
        p._decision = latest.get(p.id)


def build_liste_classe_groups(queryset, params):
    """
    Retourne une liste de blocs {groupe, grade, secretariat, participants, effectif}.
    Si params['groupe'] est fourni, un seul bloc ; sinon tous les groupes distincts.
    """
    filtered = _apply_participant_filters(queryset, params)
    participants = list(filtered)
    _attach_decisions(participants)

    buckets = defaultdict(list)
    for p in participants:
        groupe_key = _normalize_groupe_value(p.groupe) or '(Sans groupe)'
        buckets[groupe_key].append(p)

    blocks = []
    for groupe_key in sorted(buckets.keys(), key=_groupe_sort_key):
        members = sorted(buckets[groupe_key], key=lambda p: ((p.nom or '').upper(), (p.prenom or '').upper()))
        grades = sorted({m.grade for m in members if m.grade})
        secretariats = sorted({
            m.secretariat.nom for m in members if m.secretariat_id and m.secretariat
        })
        blocks.append({
            'groupe': groupe_key,
            'grade': ', '.join(grades) or '—',
            'secretariat': ', '.join(secretariats) or '—',
            'participants': members,
            'effectif': len(members),
        })
    return blocks


def _safe_sheet_title(title, used):
    base = re.sub(r'[\\/*?:\[\]]', ' ', str(title or 'Groupe'))[:28].strip() or 'Groupe'
    name = base
    n = 2
    while name in used:
        suffix = f' ({n})'
        name = f"{base[: max(1, 28 - len(suffix))]}{suffix}"
        n += 1
    used.add(name)
    return name


def export_liste_classe_pdf(blocks, meta=None):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('ReportLab non disponible')

    meta = meta or {}
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.0 * cm,
        bottomMargin=1.0 * cm,
    )
    styles = getSampleStyleSheet()
    style_inst = ParagraphStyle(
        'Inst', parent=styles['Normal'],
        fontSize=8, alignment=TA_CENTER, textColor=colors.HexColor('#555555'), spaceAfter=1,
    )
    style_title = ParagraphStyle(
        'Title', parent=styles['Title'],
        fontSize=14, textColor=colors.HexColor('#1A68AC'),
        alignment=TA_CENTER, spaceAfter=4,
    )
    style_sub = ParagraphStyle(
        'Sub', parent=styles['Normal'],
        fontSize=9, alignment=TA_CENTER, textColor=colors.HexColor('#444444'), spaceAfter=2,
    )
    cell_normal = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=8, leading=10)
    cell_header = ParagraphStyle(
        'Hdr', parent=cell_normal, fontName='Helvetica-Bold',
        textColor=colors.white, alignment=TA_CENTER,
    )

    def _p(text, hdr=False):
        safe = str(text).replace('&', '&amp;').replace('<', '&lt;')
        return Paragraph(safe, cell_header if hdr else cell_normal)

    elements = []
    for line in _HEADER_LINES:
        elements.append(Paragraph(line, style_inst))
    elements.append(Spacer(1, 0.2 * cm))
    elements.append(Paragraph('LISTE DE CLASSE — AUDITEURS', style_title))
    if meta.get('filtres_label'):
        elements.append(Paragraph(meta['filtres_label'], style_sub))
    elements.append(Paragraph(
        f"Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
        style_sub,
    ))
    elements.append(Spacer(1, 0.25 * cm))

    if not blocks:
        elements.append(Paragraph('Aucun étudiant trouvé pour les critères sélectionnés.', style_sub))
    else:
        for idx, block in enumerate(blocks):
            if idx > 0:
                elements.append(Spacer(1, 0.4 * cm))
            elements.append(Paragraph(
                f"<b>{block['groupe']}</b> — Grade : {block['grade']} — "
                f"Effectif : {block['effectif']} étudiant(s)",
                style_sub,
            ))
            if block.get('secretariat') and block['secretariat'] != '—':
                elements.append(Paragraph(f"Secrétariat : {block['secretariat']}", style_sub))

            data = [[_p(h, hdr=True) for h in _TABLE_HEADERS]]
            for num, participant in enumerate(block['participants'], start=1):
                data.append([_p(v) for v in _participant_table_row(num, participant)])

            col_widths = [
                0.7 * cm, 2.0 * cm, 2.6 * cm, 2.6 * cm, 1.0 * cm, 1.3 * cm,
                2.0 * cm, 2.2 * cm, 1.3 * cm, 1.3 * cm, 1.6 * cm,
            ]
            table = Table(data, repeatRows=1, colWidths=col_widths)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A68AC')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CCCCCC')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 1), (0, -1), 'CENTER'),
            ]))
            elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _write_liste_classe_sheet(ws, block, meta=None):
    meta = meta or {}
    green_fill = PatternFill(start_color='1A68AC', end_color='1A68AC', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    title_font = Font(bold=True, size=13, color='1A68AC')
    sub_font = Font(size=10, color='444444')

    row = 1
    for line in _HEADER_LINES:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_TABLE_HEADERS))
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = Font(size=9, color='555555')
        cell.alignment = Alignment(horizontal='center')
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_TABLE_HEADERS))
    ws.cell(row=row, column=1, value='LISTE DE CLASSE — AUDITEURS').font = title_font
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_TABLE_HEADERS))
    ws.cell(
        row=row, column=1,
        value=f"{block['groupe']} — Grade : {block['grade']} — Effectif : {block['effectif']}",
    ).font = sub_font
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    row += 1

    if meta.get('filtres_label'):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_TABLE_HEADERS))
        ws.cell(row=row, column=1, value=meta['filtres_label']).font = sub_font
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
        row += 1

    row += 1
    for col, header in enumerate(_TABLE_HEADERS, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = green_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    row += 1

    for num, participant in enumerate(block['participants'], start=1):
        for col, value in enumerate(_participant_table_row(num, participant), start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    widths = [5, 14, 18, 18, 8, 10, 14, 18, 10, 10, 12]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def export_liste_classe_excel(blocks, meta=None):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('OpenPyXL non disponible')

    meta = meta or {}
    wb = Workbook()
    used_titles = set()

    if not blocks:
        ws = wb.active
        ws.title = 'Liste de classe'
        ws['A1'] = 'Aucun étudiant trouvé pour les critères sélectionnés.'
    elif len(blocks) == 1:
        ws = wb.active
        ws.title = _safe_sheet_title(blocks[0]['groupe'], used_titles)
        _write_liste_classe_sheet(ws, blocks[0], meta)
    else:
        wb.remove(wb.active)
        for block in blocks:
            title = _safe_sheet_title(block['groupe'], used_titles)
            ws = wb.create_sheet(title=title)
            _write_liste_classe_sheet(ws, block, meta)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_liste_classe_meta(params):
    parts = []
    mapping = [
        ('secretariat', 'Secrétariat'),
        ('grade', 'Grade'),
        ('groupe', 'Groupe'),
        ('type_concours', 'Type concours'),
        ('vague', 'Vague'),
        ('sexe', 'Sexe'),
    ]
    for key, label in mapping:
        val = (params.get(key) or '').strip()
        if val:
            parts.append(f"{label} : {val}")
    if not parts:
        return {'filtres_label': 'Tous les groupes'}
    return {'filtres_label': ' — '.join(parts)}


def liste_classe_filename(blocks, fmt):
    if len(blocks) == 1:
        groupe_slug = re.sub(r'[^\w\-]+', '_', blocks[0]['groupe']).strip('_') or 'groupe'
        ext = 'pdf' if fmt == 'pdf' else 'xlsx'
        return f"liste_classe_{groupe_slug}.{ext}"
    ext = 'pdf' if fmt == 'pdf' else 'xlsx'
    return f"listes_classe_auditeurs.{ext}"
