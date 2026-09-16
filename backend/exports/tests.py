"""Tests exports présences et finance (accès, contenu PDF/Excel)."""
from datetime import timedelta
from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient, APIRequestFactory

from authentication.models import User
from formations.models import (
    Formation,
    Formateur,
    Module,
    ModuleFormateur,
    ModuleParticipant,
    Participant,
    Secretariat,
    SessionModule,
)
from presences.models import Pointage

from .views import (
    _calculer_duree_export,
    _check_export_access,
    _check_finance_export_access,
    _check_liste_classe_export_access,
    _formation_meta,
    _get_formation_data,
)


def make_user(username, role='CPFAE_ADMIN', **kwargs):
    return User.objects.create_user(username=username, password='pass', role=role, **kwargs)


class ExportFixtureMixin:
    """Formation + module + séance + pointage pour exports présences."""

    @classmethod
    def setUpTestData(cls):
        cls.secretariat_a = Secretariat.objects.create(nom='Secrétariat A')
        cls.secretariat_b = Secretariat.objects.create(nom='Secrétariat B')

        cls.formation = Formation.objects.create(formation='Cycle export')
        cls.module_a = Module.objects.create(
            formation=cls.formation,
            intitule='Module A',
            statut='EN_COURS',
            secretariat=cls.secretariat_a,
            grade='A4',
            groupe='G1',
        )
        cls.module_b = Module.objects.create(
            formation=cls.formation,
            intitule='Module B',
            statut='PLANIFIEE',
            secretariat=cls.secretariat_b,
        )

        cls.participant = Participant.objects.create(
            matricule='EXP001', nom='Koné', prenom='Awa', grade='A4', groupe='GROUPE 1',
        )
        cls.participant_g2 = Participant.objects.create(
            matricule='EXP002', nom='Diallo', prenom='Moussa', grade='A4', groupe='GROUPE 2',
        )
        ModuleParticipant.objects.create(module=cls.module_a, participant=cls.participant)

        cls.formateur = Formateur.objects.create(
            numerobadge='FEXP01', nom='Traoré', prenom='Ibrahim', specialite='Droit',
        )
        cls.formateur.secretariats.add(cls.secretariat_a)
        ModuleFormateur.objects.create(module=cls.module_a, formateur=cls.formateur)

        cls.today = timezone.localdate()
        now = timezone.now()
        cls.session = SessionModule.objects.create(
            module=cls.module_a,
            date_journee=cls.today,
            numero=1,
            demarree_le=now - timedelta(hours=2),
            terminee_le=now,
        )
        cls.pointage = Pointage.objects.create(
            participant=cls.participant,
            session=cls.session,
            date_journee=cls.today,
            timestamp_entree=now - timedelta(hours=2),
            timestamp_sortie=now,
            statut=Pointage.Statut.TERMINE,
            duree_presence_minutes=120,
        )

        cls.admin = make_user('export_admin', role='CPFAE_ADMIN')
        cls.direction = make_user('export_direction', role='DIRECTION')
        cls.finance = make_user('export_finance', role='FINANCE')
        cls.encadrant = make_user('export_encadrant', role='ENCADRANT')
        cls.module_a.superviseur = cls.encadrant
        cls.module_a.save(update_fields=['superviseur'])

        cls.secretariat_user = make_user('export_secretariat', role='SECRETARIAT')
        cls.secretariat_user.secretariat = cls.secretariat_a
        cls.secretariat_user.save(update_fields=['secretariat'])

        cls.other_secretariat = make_user('export_secretariat_b', role='SECRETARIAT')
        cls.other_secretariat.secretariat = cls.secretariat_b
        cls.other_secretariat.save(update_fields=['secretariat'])

        cls.formation_isolated = Formation.objects.create(formation='Cycle isolé A')
        cls.module_isolated = Module.objects.create(
            formation=cls.formation_isolated,
            intitule='Module isolé',
            statut='PLANIFIEE',
            secretariat=cls.secretariat_a,
        )

        cls.auditeur = make_user('export_auditeur', role='AUDITEUR')
        cls.archive = make_user('export_archive', role='ARCHIVE')


class ExportAccessTest(ExportFixtureMixin, TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_unauthenticated_presence_export_returns_401(self):
        urls = [
            f'/api/exports/formation/{self.formation.pk}/pdf/',
            f'/api/exports/module/{self.module_a.pk}/pdf/',
            f'/api/exports/session/{self.session.pk}/pdf/',
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_finance_blocked_on_presence_exports(self):
        self.client.force_authenticate(user=self.finance)
        urls = [
            f'/api/exports/formation/{self.formation.pk}/pdf/',
            f'/api/exports/formation/{self.formation.pk}/excel/',
            f'/api/exports/module/{self.module_a.pk}/pdf/',
            f'/api/exports/session/{self.session.pk}/excel/',
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_auditeur_blocked_on_presence_exports(self):
        self.client.force_authenticate(user=self.auditeur)
        res = self.client.get(f'/api/exports/formation/{self.formation.pk}/pdf/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_secretariat_scoped_to_own_modules(self):
        self.client.force_authenticate(user=self.secretariat_user)
        ok = self.client.get(f'/api/exports/formation/{self.formation.pk}/pdf/')
        self.assertEqual(ok.status_code, status.HTTP_200_OK)

        self.client.force_authenticate(user=self.other_secretariat)
        ok_shared = self.client.get(f'/api/exports/formation/{self.formation.pk}/pdf/')
        self.assertEqual(ok_shared.status_code, status.HTTP_200_OK)

        denied = self.client.get(f'/api/exports/formation/{self.formation_isolated.pk}/pdf/')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

    def test_encadrant_scoped_to_supervised_modules(self):
        self.client.force_authenticate(user=self.encadrant)
        ok = self.client.get(f'/api/exports/module/{self.module_a.pk}/pdf/')
        self.assertEqual(ok.status_code, status.HTTP_200_OK)

        other_encadrant = make_user('export_encadrant2', role='ENCADRANT')
        self.client.force_authenticate(user=other_encadrant)
        denied = self.client.get(f'/api/exports/module/{self.module_a.pk}/pdf/')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_export_all_presence_formats(self):
        self.client.force_authenticate(user=self.admin)
        endpoints = [
            (f'/api/exports/formation/{self.formation.pk}/pdf/', 'application/pdf'),
            (f'/api/exports/formation/{self.formation.pk}/excel/',
             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            (f'/api/exports/module/{self.module_a.pk}/pdf/', 'application/pdf'),
            (f'/api/exports/session/{self.session.pk}/excel/',
             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        ]
        for url, content_type in endpoints:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_200_OK)
                self.assertEqual(res['Content-Type'], content_type)
                self.assertGreater(len(res.content), 500)
                if content_type == 'application/pdf':
                    self.assertTrue(res.content.startswith(b'%PDF'))
                else:
                    self.assertTrue(res.content.startswith(b'PK'))

    def test_check_export_access_matrix(self):
        factory = APIRequestFactory()

        def req(user):
            request = factory.get('/')
            request.user = user
            return request

        self.assertTrue(_check_export_access(req(self.admin), self.formation))
        self.assertTrue(_check_export_access(req(self.direction), self.formation))
        self.assertTrue(_check_export_access(req(self.archive), self.formation))
        self.assertTrue(_check_export_access(req(self.secretariat_user), self.formation))
        self.assertTrue(_check_export_access(req(self.other_secretariat), self.formation))
        self.assertFalse(_check_export_access(req(self.other_secretariat), self.formation_isolated))
        self.assertTrue(_check_export_access(req(self.encadrant), self.formation))
        self.assertFalse(_check_export_access(req(self.finance), self.formation))
        self.assertFalse(_check_export_access(req(self.auditeur), self.formation))


class ExportFinanceAccessTest(ExportFixtureMixin, TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_finance_allowed_on_finance_exports(self):
        self.client.force_authenticate(user=self.finance)
        urls = [
            f'/api/exports/formateur/{self.formateur.pk}/pdf/',
            f'/api/exports/formateur/{self.formateur.pk}/excel/',
            '/api/exports/finance/synthese/pdf/',
            '/api/exports/finance/synthese/excel/',
            '/api/exports/finance/encadrants/pdf/',
            '/api/exports/finance/encadrants/excel/',
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_200_OK, msg=url)
                self.assertIn(
                    res['Content-Type'],
                    ('application/pdf',
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                )
                self.assertGreater(len(res.content), 500)

    def test_direction_allowed_on_finance_exports(self):
        self.client.force_authenticate(user=self.direction)
        res = self.client.get(f'/api/exports/formateur/{self.formateur.pk}/pdf/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.content.startswith(b'%PDF'))

    def test_archive_allowed_on_finance_and_presence_exports(self):
        self.client.force_authenticate(user=self.archive)
        urls = [
            f'/api/exports/formation/{self.formation.pk}/pdf/',
            f'/api/exports/formateur/{self.formateur.pk}/pdf/',
            '/api/exports/finance/encadrants/pdf/',
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_200_OK, msg=url)
                self.assertGreater(len(res.content), 500)

    def test_secretariat_blocked_on_finance_exports(self):
        self.client.force_authenticate(user=self.secretariat_user)
        urls = [
            f'/api/exports/formateur/{self.formateur.pk}/pdf/',
            '/api/exports/finance/synthese/pdf/',
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_check_finance_export_access_matrix(self):
        factory = APIRequestFactory()

        def req(user):
            request = factory.get('/')
            request.user = user
            return request

        self.assertTrue(_check_finance_export_access(req(self.finance)))
        self.assertTrue(_check_finance_export_access(req(self.direction)))
        self.assertTrue(_check_finance_export_access(req(self.archive)))
        self.assertFalse(_check_finance_export_access(req(self.admin)))
        self.assertFalse(_check_finance_export_access(req(self.secretariat_user)))


class ExportDataHelpersTest(ExportFixtureMixin, TestCase):
    def test_formation_meta_includes_superviseurs(self):
        meta = _formation_meta(self.formation)
        self.assertEqual(meta['grade'], 'A4')
        self.assertEqual(meta['groupe'], 'G1')
        self.assertEqual(len(meta['superviseurs']), 1)
        self.assertEqual(meta['superviseurs'][0].username, 'export_encadrant')

    def test_get_formation_data_includes_pointage_row(self):
        _, rows, stats = _get_formation_data(self.formation.pk)
        participant_rows = [r for r in rows if r.get('role') == 'Étudiant']
        self.assertGreaterEqual(len(participant_rows), 1)
        self.assertGreaterEqual(stats['nb_presents'], 1)
        self.assertGreater(stats['taux'], 0)

    def test_get_formation_data_module_filter(self):
        _, rows_all, _ = _get_formation_data(self.formation.pk)
        _, rows_module, _ = _get_formation_data(self.formation.pk, module_pk=self.module_a.pk)
        self.assertGreater(len(rows_all), 0)
        self.assertGreater(len(rows_module), 0)
        self.assertLessEqual(len(rows_module), len(rows_all))

    def test_calculer_duree_export_closed_pointage(self):
        minutes, duree_str, entree, sortie = _calculer_duree_export(self.pointage)
        self.assertEqual(minutes, 120.0)
        self.assertIn('h', duree_str)
        self.assertRegex(entree, r'^\d{2}:\d{2}$')
        self.assertRegex(sortie, r'^\d{2}:\d{2}$')

    def test_calculer_duree_export_open_pointage_shows_en_cours(self):
        open_pt = Pointage.objects.create(
            participant=self.participant,
            session=self.session,
            date_journee=self.today,
            timestamp_entree=timezone.now() - timedelta(minutes=30),
            statut=Pointage.Statut.EN_COURS,
        )
        _, _, _, sortie = _calculer_duree_export(open_pt)
        self.assertEqual(sortie, 'En cours')

    def test_finance_formateur_pdf_contains_formateur_name(self):
        client = APIClient()
        client.force_authenticate(user=self.finance)
        res = client.get(f'/api/exports/formateur/{self.formateur.pk}/pdf/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Vérification légère du contenu PDF (texte compressé — recherche du nom encodé possible)
        self.assertTrue(res.content.startswith(b'%PDF'))
        self.assertGreater(len(res.content), 1000)

    def test_finance_synthese_excel_is_valid_workbook(self):
        from openpyxl import load_workbook

        client = APIClient()
        client.force_authenticate(user=self.finance)
        res = client.get('/api/exports/finance/synthese/excel/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = load_workbook(BytesIO(res.content))
        self.assertGreaterEqual(len(wb.sheetnames), 1)

    def test_presence_excel_contains_participant_name(self):
        from openpyxl import load_workbook

        client = APIClient()
        client.force_authenticate(user=self.admin)
        res = client.get(f'/api/exports/formation/{self.formation.pk}/excel/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = load_workbook(BytesIO(res.content))
        ws = wb.active
        values = [str(cell.value or '') for row in ws.iter_rows() for cell in row]
        joined = ' '.join(values)
        self.assertIn('Koné', joined)
        self.assertIn('Awa', joined)


class ListeClasseExportTest(ExportFixtureMixin, TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_check_liste_classe_export_access_matrix(self):
        factory = APIRequestFactory()

        def req(user):
            request = factory.get('/')
            request.user = user
            return request

        self.assertTrue(_check_liste_classe_export_access(req(self.archive)))
        self.assertTrue(_check_liste_classe_export_access(req(self.admin)))
        self.assertTrue(_check_liste_classe_export_access(req(self.direction)))
        self.assertTrue(_check_liste_classe_export_access(req(self.secretariat_user)))
        self.assertTrue(_check_liste_classe_export_access(req(self.encadrant)))
        self.assertFalse(_check_liste_classe_export_access(req(self.finance)))
        self.assertFalse(_check_liste_classe_export_access(req(self.auditeur)))

    def test_archive_can_export_liste_classe_by_groupe(self):
        self.client.force_authenticate(user=self.archive)
        res = self.client.get('/api/exports/participants/liste-classe/pdf/?groupe=GROUPE%201')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res['Content-Type'], 'application/pdf')
        self.assertTrue(res.content.startswith(b'%PDF'))
        self.assertGreater(len(res.content), 500)

    def test_liste_classe_excel_single_groupe(self):
        from openpyxl import load_workbook

        self.client.force_authenticate(user=self.archive)
        res = self.client.get('/api/exports/participants/liste-classe/excel/?groupe=GROUPE%201')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = load_workbook(BytesIO(res.content))
        self.assertEqual(len(wb.sheetnames), 1)
        values = [str(cell.value or '') for row in wb.active.iter_rows() for cell in row]
        joined = ' '.join(values)
        self.assertIn('KONÉ', joined)
        self.assertNotIn('Diallo', joined)

    def test_liste_classe_excel_all_groupes(self):
        from openpyxl import load_workbook

        self.client.force_authenticate(user=self.archive)
        res = self.client.get('/api/exports/participants/liste-classe/excel/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = load_workbook(BytesIO(res.content))
        self.assertGreaterEqual(len(wb.sheetnames), 2)
