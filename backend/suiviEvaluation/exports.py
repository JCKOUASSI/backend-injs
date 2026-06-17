"""Exports helpers for suiviEvaluation: PDF & Excel for fiches auditeur/formateur."""
import io
from django.core.files.base import ContentFile
from .models import ExportRapport

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    OPENPYXL_AVAILABLE = True
except Exception:
    OPENPYXL_AVAILABLE = False


def _build_fiche_auditeur_table_rows(fiche):
    rows = []
    for sm in fiche.suivi_modules.all():
        rows.append([
            sm.module.intitule,
            str(sm.heures_presence),
            str(sm.heures_prevues),
            str(sm.taux_presence or ''),
            str(sm.moyenne_module or ''),
        ])
    return rows


def export_fiche_auditeur_pdf(fiche):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('reportlab non disponible')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elems = []
    title = f"Fiche académique — {fiche.participant.nom} {fiche.participant.prenom}"
    elems.append(Paragraph(title, styles['Title']))
    elems.append(Spacer(1, 12))
    meta = f"Formation: {fiche.formation.formation} — Année: {fiche.annee_academique or ''}"
    elems.append(Paragraph(meta, styles['Normal']))
    elems.append(Spacer(1, 12))
    data = [["Module", "Heures présence", "Heures prévues", "Taux %", "Moyenne"]]
    data += _build_fiche_auditeur_table_rows(fiche)
    tbl = Table(data, colWidths=[9*cm, 3*cm, 3*cm, 2*cm, 2*cm])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
    ]))
    elems.append(tbl)
    elems.append(Spacer(1, 12))
    decision = fiche.decision_finale.get_decision_display() if fiche.decision_finale else ''
    elems.append(Paragraph(f"Moyenne générale: {fiche.moyenne_generale or ''} — Décision: {decision}", styles['Normal']))
    doc.build(elems)
    buffer.seek(0)
    return buffer


def export_fiche_auditeur_excel(fiche):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('openpyxl non disponible')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Fiche Auditeur'
    ws['A1'] = f"Fiche académique — {fiche.participant.nom} {fiche.participant.prenom}"
    ws['A2'] = f"Formation: {fiche.formation.formation}"
    ws['A1'].font = Font(bold=True)
    headers = ["Module", "Heures présence", "Heures prévues", "Taux %", "Moyenne"]
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in _build_fiche_auditeur_table_rows(fiche):
        ws.append(row)
    ws.append([])
    ws.append(["Moyenne générale", fiche.moyenne_generale or ''])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def save_export_report(file_buffer, filename, type_export, type_rapport, user=None, formation=None, module=None, participant=None):
    content = ContentFile(file_buffer.getvalue())
    obj = ExportRapport.objects.create(
        type_export=type_export,
        type_rapport=type_rapport,
        parametres={},
        genere_par=user,
        formation=formation,
        module=module,
        participant=participant,
    )
    obj.fichier.save(filename, content)
    obj.save()
    return obj


def _build_fiche_formateur_rows(fiche):
    return [[
        fiche.module.intitule if fiche.module else (fiche.formation.formation if fiche.formation else ''),
        str(fiche.heures_prevues),
        str(fiche.heures_effectuees),
        str(fiche.taux_presence or ''),
        str(fiche.nb_seances),
        str(fiche.nb_seances_realisees),
        str(fiche.satisfaction_auditeurs or ''),
    ]]


def export_fiche_formateur_pdf(fiche):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('reportlab non disponible')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elems = []
    title = f"Fiche formateur — {fiche.formateur.nom} {fiche.formateur.prenom}"
    elems.append(Paragraph(title, styles['Title']))
    elems.append(Spacer(1, 12))
    meta = f"Module: {fiche.module.intitule if fiche.module else ''} — Formation: {fiche.formation.formation if fiche.formation else ''}"
    elems.append(Paragraph(meta, styles['Normal']))
    elems.append(Spacer(1, 12))
    data = [["Module/Formation", "Heures prévues", "Heures effectuées", "Taux %", "Nb séances", "Nb réalisées", "Satisfaction"]]
    data += _build_fiche_formateur_rows(fiche)
    tbl = Table(data, colWidths=[7*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
    ]))
    elems.append(tbl)
    elems.append(Spacer(1, 12))
    elems.append(Paragraph(f"Appréciation: {fiche.appreciation_pedagogique or ''}", styles['Normal']))
    doc.build(elems)
    buffer.seek(0)
    return buffer


def export_fiche_formateur_excel(fiche):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('openpyxl non disponible')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Fiche Formateur'
    ws['A1'] = f"Fiche formateur — {fiche.formateur.nom} {fiche.formateur.prenom}"
    ws['A2'] = f"Module/Formation: {fiche.module.intitule if fiche.module else ''} / {fiche.formation.formation if fiche.formation else ''}"
    ws['A1'].font = Font(bold=True)
    headers = ["Module/Formation", "Heures prévues", "Heures effectuées", "Taux %", "Nb séances", "Nb réalisées", "Satisfaction"]
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in _build_fiche_formateur_rows(fiche):
        ws.append(row)
    ws.append([])
    ws.append(["Appréciation", fiche.appreciation_pedagogique or ''])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
