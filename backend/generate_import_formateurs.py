"""
Script de génération du fichier d'import formateurs.

Source : LISTE DES FORMATEURS FAB 2026_Prod V2.xlsx
Sortie : import_formateurs.xlsx  (feuille "Formateurs" au format attendu par import_excel.py)

Nettoyages appliqués :
  - Numéros de téléphone multi-lignes → seul le premier numéro est conservé
  - Espaces en début/fin de chaque champ
  - Emails vides/espaces → chaîne vide
  - Organisation None → chaîne vide
"""

import sys
import os

SOURCE = '/Users/tobidesis/Downloads/données/LISTE DES FORMATEURS FAB 2026_Prod V2.xlsx'
OUTPUT = os.path.join(os.path.dirname(__file__), 'import_formateurs.xlsx')

try:
    import openpyxl
except ImportError:
    print("openpyxl manquant. Installez-le : pip install openpyxl")
    sys.exit(1)

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def clean(val, default=''):
    if val is None:
        return default
    return str(val).strip()


def clean_phone(val):
    """Garde uniquement le premier numéro si plusieurs sont séparés par \\n."""
    s = clean(val)
    if not s:
        return ''
    first = s.split('\n')[0].strip()
    first = first.replace(' ', '')
    return first


def clean_email(val):
    s = clean(val)
    return s if '@' in s else ''


def main():
    print(f"Lecture de : {SOURCE}")
    wb_src = load_workbook(SOURCE)
    ws_src = wb_src['Formateurs']
    rows = list(ws_src.iter_rows(values_only=True))

    header = rows[0]
    print(f"En-têtes source : {list(header)}")
    print(f"Nombre de lignes de données : {len(rows) - 1}")

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = 'Formateurs'

    # En-tête de sortie (noms reconnus par import_excel.py _rows())
    out_header = ['Numéro', 'Nom', 'Prénoms', 'Email', 'Téléphone', 'Spécialité', 'Organisation']
    ws_out.append(out_header)

    # Style en-tête
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in ws_out[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    stats = {'total': 0, 'phone_fixed': 0, 'email_empty': 0, 'skipped': 0}

    for i, row in enumerate(rows[1:], 2):
        numero  = clean(row[0])
        nom     = clean(row[1])
        prenom  = clean(row[2])
        email   = clean_email(row[3])
        phone   = clean_phone(row[4])
        specialite  = clean(row[5])
        organisation = clean(row[6])

        if not nom or not prenom:
            print(f"  ⚠  Ligne {i} ignorée (nom/prénom vide) : {row}")
            stats['skipped'] += 1
            continue

        if row[4] and '\n' in str(row[4]):
            stats['phone_fixed'] += 1

        if not email:
            stats['email_empty'] += 1

        ws_out.append([numero, nom, prenom, email, phone, specialite, organisation])
        stats['total'] += 1

    # Ajuster largeurs colonnes
    col_widths = [10, 22, 30, 35, 18, 35, 20]
    for col_idx, width in enumerate(col_widths, 1):
        ws_out.column_dimensions[ws_out.cell(1, col_idx).column_letter].width = width

    wb_out.save(OUTPUT)
    print(f"\n✅ Fichier généré : {OUTPUT}")
    print(f"   Formateurs écrits  : {stats['total']}")
    print(f"   Téléphones nettoyés : {stats['phone_fixed']} (multi-lignes → 1er numéro)")
    print(f"   Emails vides       : {stats['email_empty']}")
    if stats['skipped']:
        print(f"   Lignes ignorées    : {stats['skipped']}")
    print(f"\nUsage import :")
    print(f"  python manage.py import_excel import_formateurs.xlsx")
    print(f"  python manage.py import_excel import_formateurs.xlsx --dry-run")


if __name__ == '__main__':
    main()
