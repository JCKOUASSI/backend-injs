"""
Analyse complète du dossier données :
- Structure de chaque fichier source (feuilles, colonnes, nb lignes)
- Couverture par grade/groupe
- Modules présents
- Dates couvertes
- Données manquantes
"""
import openpyxl, os, re
from collections import defaultdict

SRC = 'donnees'

MONTHS_FR = {'JANVIER':1,'FEVRIER':2,'MARS':3,'AVRIL':4,'MAI':5,'JUIN':6,
             'JUILLET':7,'AOUT':8,'SEPTEMBRE':9,'OCTOBRE':10,'NOVEMBRE':11,'DECEMBRE':12}

def clean(v):
    if v is None: return ''
    return str(v).strip().replace('\xa0',' ').strip()

def parse_date_fr(s):
    if not s: return None
    s = re.sub(r'^(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE)\s*','',s.upper()).strip()
    m = re.match(r'(\d{1,2})\s+(\w+)\s+(\d{4})',s)
    if m:
        from datetime import date
        day, mon, year = int(m.group(1)), MONTHS_FR.get(m.group(2)), int(m.group(3))
        if mon:
            try: return date(year,mon,day)
            except: pass
    return None

src_files = [f for f in os.listdir(SRC) if f.endswith('.xlsx') and not f.startswith('import_')]
src_files.sort()

print(f"\n{'='*65}")
print("ANALYSE COMPLÈTE DU DOSSIER DONNÉES")
print(f"{'='*65}")
print(f"\n{len(src_files)} fichier(s) source(s) :\n")

# ─── FICHIER FAB COMPLET ─────────────────────────────────────────────────────
fab_file = next((f for f in src_files if 'COMPLET' in f.upper()), None)
if fab_file:
    print(f"{'─'*65}")
    print(f"📊 {fab_file}")
    print(f"{'─'*65}")
    wb = openpyxl.load_workbook(os.path.join(SRC, fab_file), read_only=True, data_only=True)
    skip = ('TCD','RECAP')
    total_p = 0
    grades_info = defaultdict(list)

    for shname in wb.sheetnames:
        if shname in skip: continue
        ws = wb[shname]
        rows = list(ws.iter_rows(values_only=True, max_row=3))
        if len(rows) < 2: continue
        headers = [clean(h) for h in rows[0]]
        row1 = rows[1]
        if not row1 or all(v is None for v in row1): continue

        def gcol(pat):
            for i,h in enumerate(headers):
                if pat.lower() in h.lower(): return i
            return None

        grade  = clean(row1[gcol('GRADE')]  if gcol('GRADE')  is not None else '')
        groupe = clean(row1[gcol('GROUPE')] if gcol('GROUPE') is not None else '')
        cat    = clean(row1[gcol('CATEG')]  if gcol('CATEG')  is not None else '')
        site   = clean(row1[gcol('SITE')]   if gcol('SITE')   is not None else '')
        salle  = clean(row1[gcol('SALLE')]  if gcol('SALLE')  is not None else '')
        nb = sum(1 for r in ws.iter_rows(min_row=2,values_only=True) if r and any(v for v in r))
        total_p += nb
        grades_info[grade].append({'groupe':groupe,'cat':cat,'site':site,'salle':salle,'nb':nb,'sheet':shname})

    wb.close()

    for grade in sorted(grades_info.keys()):
        glist = grades_info[grade]
        sous_total = sum(g['nb'] for g in glist)
        print(f"\n  Grade {grade} — {len(glist)} groupes — {sous_total} participants")
        for g in sorted(glist, key=lambda x: x['groupe']):
            print(f"    {g['groupe']:15s}  cat:{g['cat']:2s}  salle:{g['salle'][:20]:20s}  ({g['nb']:4d} participants)")
    print(f"\n  TOTAL : {sum(len(v) for v in grades_info.values())} groupes — {total_p} participants")
    print(f"  Colonnes: {[h for h in headers if h][:10]}")

# ─── FICHIERS EDT ─────────────────────────────────────────────────────────────
edt_files = [f for f in src_files if 'Emploi' in f]
print(f"\n{'─'*65}")
print(f"📅 FICHIERS EMPLOI DU TEMPS ({len(edt_files)} fichiers)")
print(f"{'─'*65}")

all_groupes_edt = set()
all_modules_edt = set()
couverture = defaultdict(lambda: defaultdict(set))  # grade->groupe->modules

for fname in sorted(edt_files):
    path = os.path.join(SRC, fname)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"\n  📄 {fname}")

    for shname in wb.sheetnames:
        ws = wb[shname]
        all_rows = list(ws.iter_rows(values_only=True))

        grade_str = groupe_str = ''
        header_idx = None
        for i, row in enumerate(all_rows):
            vals = [clean(v) for v in row]
            full = ' '.join(vals).upper()
            if ('EMPLOI DU TEMPS' in full or 'GRADE' in full) and not grade_str:
                grade_str = full
            if 'GROUPE' in full and not groupe_str:
                for v in vals:
                    if re.search(r'GROUPE\s*[:\s]*0*\d+', v.upper()):
                        groupe_str = v.strip()
                        break
            if 'MATIERES' in full:
                header_idx = i
                break

        # Grade
        grade = ''
        m = re.search(r'GRADE\s+([A-Z]\d)', grade_str.upper())
        if m: grade = m.group(1)
        if not grade:
            m2 = re.search(r'([A-Z]\d)\s*[-/]?\s*GROUPE', groupe_str.upper())
            if m2: grade = m2.group(1)
        if not grade and 'CATEGORIE B' in grade_str.upper(): grade = 'B3'
        if not grade and 'CATEGORIE A' in grade_str.upper():
            # essayer de déduire A3 ou A4
            m3 = re.search(r'A(\d)', grade_str.upper())
            if m3: grade = f'A{m3.group(1)}'

        # Groupe
        groupe = ''
        m4 = re.search(r'GROUPE\s*[:\s]*0*(\d+)', groupe_str.upper())
        if m4: groupe = f'GROUPE {int(m4.group(1))}'

        # Modules et dates
        modules = set()
        dates = set()
        if header_idx is not None:
            data_start = header_idx + 1
            if data_start < len(all_rows):
                nv = [clean(v).upper() for v in all_rows[data_start]]
                if any('DU MODULE' in v or 'RESTANT' in v for v in nv):
                    data_start += 1

            current = ''
            noise_kw = ('MINISTERE','DIRECTION','MATIERES','EMPLOI DU TEMPS',
                        'DU MODULE','RESTANT','FORMATION EN','CPFAE','GROUPE',
                        'SITE','BATIMENT','SALLE','REPUBLIQUE')
            for row in all_rows[data_start:]:
                vals = [clean(v) for v in row]
                non_empty = [v for v in vals if v]
                if not non_empty: continue
                full2 = ' '.join(non_empty).upper()
                if any(k in full2 for k in noise_kw): continue

                first = vals[0] if vals else ''
                if first and not any(k in first.upper() for k in noise_kw):
                    if not re.match(r'^\d+\s*H?\s*$', first.upper()):
                        if not re.search(r'(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE)', first.upper()):
                            if not any(mo in first.upper() for mo in ('JANVIER','FEVRIER','MARS','AVRIL','MAI','JUIN','JUILLET','AOUT','SEPTEMBRE','OCTOBRE','NOVEMBRE','DECEMBRE')):
                                current = first.strip()
                                if re.match(r'SIGFAE\s+ET\s*$', current.upper()): current = 'SIGFAE'

                for v in non_empty:
                    d = parse_date_fr(v)
                    if d and current:
                        modules.add(current)
                        dates.add(d)
                        couverture[grade][groupe].add(current)

        all_groupes_edt.add((grade, groupe))
        all_modules_edt.update(modules)

        dates_sorted = sorted(dates)
        date_range = f"{dates_sorted[0]} → {dates_sorted[-1]}" if dates_sorted else "aucune date"
        print(f"    Feuille '{shname:12s}'  Grade:{grade or '?':4s}  {groupe or '?':12s}  "
              f"{len(modules):2d} modules  {len(dates):3d} jours  [{date_range}]")
        for mod in sorted(modules):
            print(f"      • {mod}")

    wb.close()

# ─── RÉCAP COUVERTURE ─────────────────────────────────────────────────────────
print(f"\n{'='*65}")
print("RÉCAPITULATIF COUVERTURE")
print(f"{'='*65}")

TOUS_GROUPES = {
    'A3': [f'GROUPE {i}' for i in range(1,11)],
    'A4': [f'GROUPE {i}' for i in range(1,9)],
    'B3': [f'GROUPE {i}' for i in range(1,4)],
}

print(f"\n  Modules distincts dans les EDT : {len(all_modules_edt)}")
for mod in sorted(all_modules_edt):
    print(f"    • {mod}")

print(f"\n  Couverture par grade/groupe :")
manquants = []
for grade, groupes in sorted(TOUS_GROUPES.items()):
    for groupe in groupes:
        mods = couverture[grade].get(groupe, set())
        status = '✅' if mods else '❌ MANQUANT'
        print(f"    {grade} {groupe:12s}: {status} ({len(mods)} modules)")
        if not mods:
            manquants.append(f"{grade} {groupe}")

print(f"\n  ❌ {len(manquants)} groupe(s) sans EDT :")
for m in manquants:
    print(f"     • {m}")

print(f"\n{'─'*65}")
print("SYNTHÈSE FINALE")
print(f"{'─'*65}")
print(f"  Fichier FAB COMPLET : 21 groupes, 1246 participants")
print(f"  EDT couverts        : {21-len(manquants)}/21 groupes")
print(f"  EDT manquants       : {len(manquants)} groupes → {', '.join(manquants)}")
print(f"  Modules EDT         : {len(all_modules_edt)}")
