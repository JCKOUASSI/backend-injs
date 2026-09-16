"""Export PDF / Excel du relevé de notes auditeur avec décision finale."""
import io

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    OPENPYXL_AVAILABLE = True
except Exception:
    OPENPYXL_AVAILABLE = False


DECISION_LABELS = {
    'ADMIS': 'Admis',
    'AJOURNE': 'Ajourné',
    'EXCLUSION': 'Exclusion',
    'EN_ATTENTE': 'En attente',
}

MENTION_LABELS = {
    'TRES_BIEN': 'Très bien',
    'BIEN': 'Bien',
    'ASSEZ_BIEN': 'Assez bien',
    'PASSABLE': 'Passable',
    'INSUFFISANT': 'Insuffisant',
}


def _fmt_num(val, suffix=''):
    if val is None or val == '':
        return '—'
    return f'{val}{suffix}'


def _decision_label(code):
    return DECISION_LABELS.get(code or '', code or '—')


def _mention_label(code):
    return MENTION_LABELS.get(code or '', code or '')


def build_releve_notes_export_data(participant, inscriptions, notes_payload):
    """Structure les données pour export (cours + synthèse par formation)."""
    modules_by_id = {ins.module_id: ins.module for ins in inscriptions}
    module_rows = []
    for item in notes_payload.get('modules', []):
        mod = modules_by_id.get(item['module_id'])
        module_rows.append({
            'module_id': item['module_id'],
            'formation_id': item['formation_id'],
            'module': mod.intitule if mod else '',
            'formation': (
                mod.formation.formation
                if mod and mod.formation_id else ''
            ),
            'moyenne': item.get('moyenne'),
            'taux_presence': item.get('taux_presence'),
            'heures_presence': item.get('heures_presence'),
            'heures_prevues': item.get('heures_prevues'),
            'mention': item.get('mention') or '',
        })

    formations = []
    for f in notes_payload.get('formations', []):
        dec = f.get('decision') or {}
        formations.append({
            'formation_id': f['formation_id'],
            'formation_libelle': f.get('formation_libelle') or '',
            'criteres': f.get('criteres') or {},
            'moyenne_generale': dec.get('moyenne_generale'),
            'taux_presence': dec.get('taux_presence'),
            'total_heures_presence': dec.get('total_heures_presence'),
            'total_heures_prevues': dec.get('total_heures_prevues'),
            'decision': dec.get('decision') or 'EN_ATTENTE',
            'mention': dec.get('mention') or '',
        })

    return {
        'participant': {
            'nom': participant.nom,
            'prenom': participant.prenom,
            'matricule': participant.matricule or '',
        },
        'modules': module_rows,
        'formations': formations,
    }


def export_releve_notes_pdf(data):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('reportlab non disponible')

    p = data['participant']
    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    elems = []

    elems.append(Paragraph(
        f"Relevé de notes — {p['nom']} {p['prenom']}",
        styles['Title'],
    ))
    elems.append(Spacer(1, 8))
    elems.append(Paragraph(
        f"N° d'inscription : {p['matricule'] or '—'}",
        styles['Normal'],
    ))
    elems.append(Spacer(1, 14))

    multi_formations = len(data['formations']) > 1
    modules_by_formation = {}
    for row in data['modules']:
        modules_by_formation.setdefault(row['formation_id'], []).append(row)

    if not data['formations']:
        formations_order = [{'formation_id': None, 'formation_libelle': ''}]
        modules_by_formation[None] = data['modules']
    else:
        formations_order = data['formations']

    for formation in formations_order:
        fid = formation['formation_id']
        rows = modules_by_formation.get(fid, [])
        if multi_formations and formation.get('formation_libelle'):
            elems.append(Paragraph(formation['formation_libelle'], styles['Heading2']))
            elems.append(Spacer(1, 6))

        table_data = [[
            'Cours',
            'Moyenne /20',
            'Temps effectué %',
            'Heures effectuées / prévues',
        ]]
        for row in rows:
            hp = row.get('heures_presence')
            hpr = row.get('heures_prevues')
            heures = (
                f"{hp}h / {hpr}h"
                if hp is not None and hpr is not None
                else '—'
            )
            table_data.append([
                row.get('module') or '—',
                _fmt_num(row.get('moyenne')),
                _fmt_num(row.get('taux_presence'), '%'),
                heures,
            ])

        if len(table_data) > 1:
            tbl = Table(
                table_data,
                colWidths=[8 * cm, 2.5 * cm, 3 * cm, 4 * cm],
                repeatRows=1,
            )
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
            ]))
            elems.append(tbl)
        else:
            elems.append(Paragraph('Aucun cours inscrit.', styles['Normal']))

        elems.append(Spacer(1, 12))
        elems.append(Paragraph('Synthèse et décision finale', styles['Heading3']))
        elems.append(Spacer(1, 6))

        mg = formation.get('moyenne_generale')
        tp = formation.get('taux_presence')
        thp = formation.get('total_heures_presence')
        thpr = formation.get('total_heures_prevues')
        criteres = formation.get('criteres') or {}
        seuil_note = criteres.get('seuil_admission', 12)
        seuil_taux = criteres.get('taux_presence_min', 80)

        synth_lines = [
            f"<b>Moyenne générale :</b> {_fmt_num(mg, '/20')} (seuil : {seuil_note}/20)",
            f"<b>Temps de cours effectué :</b> {_fmt_num(tp, '%')}"
            + (
                f" ({thp}h / {thpr}h)"
                if thp is not None and thpr is not None else ''
            )
            + f" — seuil : {seuil_taux}%",
            f"<b>Décision finale :</b> {_decision_label(formation.get('decision'))}",
        ]
        mention = _mention_label(formation.get('mention'))
        if mention:
            synth_lines.append(f"<b>Mention :</b> {mention}")

        for line in synth_lines:
            elems.append(Paragraph(line, styles['Normal']))
        elems.append(Spacer(1, 16))

    doc.build(elems)
    buffer.seek(0)
    return buffer


def export_releve_notes_excel(data):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('openpyxl non disponible')

    p = data['participant']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Relevé de notes'

    ws['A1'] = f"Relevé de notes — {p['nom']} {p['prenom']}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"N° d'inscription : {p['matricule'] or '—'}"
    row_idx = 4

    multi_formations = len(data['formations']) > 1
    modules_by_formation = {}
    for row in data['modules']:
        modules_by_formation.setdefault(row['formation_id'], []).append(row)

    formations_order = data['formations'] or [{'formation_id': None, 'formation_libelle': '', 'criteres': {}}]
    if not data['formations']:
        modules_by_formation[None] = data['modules']

    headers = ['Cours', 'Moyenne /20', 'Temps effectué %', 'Heures effectuées / prévues']

    for formation in formations_order:
        fid = formation['formation_id']
        if multi_formations and formation.get('formation_libelle'):
            ws.cell(row=row_idx, column=1, value=formation['formation_libelle']).font = Font(bold=True, size=12)
            row_idx += 1

        for col, h in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col, value=h).font = Font(bold=True)
        row_idx += 1

        for mod_row in modules_by_formation.get(fid, []):
            hp = mod_row.get('heures_presence')
            hpr = mod_row.get('heures_prevues')
            heures = (
                f"{hp}h / {hpr}h"
                if hp is not None and hpr is not None
                else '—'
            )
            ws.cell(row=row_idx, column=1, value=mod_row.get('module') or '—')
            ws.cell(
                row=row_idx, column=2,
                value=mod_row.get('moyenne') if mod_row.get('moyenne') is not None else '—',
            )
            ws.cell(
                row=row_idx, column=3,
                value=f"{mod_row['taux_presence']}%" if mod_row.get('taux_presence') is not None else '—',
            )
            ws.cell(row=row_idx, column=4, value=heures)
            row_idx += 1

        row_idx += 1
        ws.cell(row=row_idx, column=1, value='Synthèse et décision finale').font = Font(bold=True)
        row_idx += 1

        criteres = formation.get('criteres') or {}
        mg = formation.get('moyenne_generale')
        tp = formation.get('taux_presence')
        thp = formation.get('total_heures_presence')
        thpr = formation.get('total_heures_prevues')

        ws.cell(row=row_idx, column=1, value='Moyenne générale')
        ws.cell(row=row_idx, column=2, value=f"{mg}/20" if mg is not None else '—')
        row_idx += 1
        ws.cell(row=row_idx, column=1, value='Temps de cours effectué')
        taux_txt = f"{tp}%" if tp is not None else '—'
        if thp is not None and thpr is not None:
            taux_txt += f" ({thp}h / {thpr}h)"
        ws.cell(row=row_idx, column=2, value=taux_txt)
        row_idx += 1
        ws.cell(row=row_idx, column=1, value='Décision finale')
        ws.cell(row=row_idx, column=2, value=_decision_label(formation.get('decision')))
        row_idx += 1
        mention = _mention_label(formation.get('mention'))
        if mention:
            ws.cell(row=row_idx, column=1, value='Mention')
            ws.cell(row=row_idx, column=2, value=mention)
            row_idx += 1
        ws.cell(
            row=row_idx, column=1,
            value=(
                f"Critères : moyenne ≥ {criteres.get('seuil_admission', 12)}/20"
                f" et temps effectué ≥ {criteres.get('taux_presence_min', 80)}%"
            ),
        )
        row_idx += 2

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
