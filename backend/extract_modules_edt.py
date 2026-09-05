"""Extrait tous les modules/matières distincts des fichiers EDT."""
import openpyxl, os, re

SRC = '/Users/tobidesis/Downloads/données'

JOURS = ('LUNDI','MARDI','MERCREDI','JEUDI','VENDREDI','SAMEDI','DIMANCHE')
MOIS  = ('JANVIER','FEVRIER','MARS','AVRIL','MAI','JUIN','JUILLET',
          'AOUT','SEPTEMBRE','OCTOBRE','NOVEMBRE','DECEMBRE')

def is_skip(v):
    v = str(v).upper().strip().replace('\xa0','')
    if not v or len(v) < 3: return True
    if any(j in v for j in JOURS): return True
    if any(m in v for m in MOIS): return True
    if re.match(r'^\d+\s*H', v): return True
    if 'MINISTERE' in v or 'DIRECTION' in v or 'CENTRE' in v: return True
    if 'FORMATION EN' in v or 'EMPLOI' in v or 'MATIERES' in v: return True
    if 'DU MODULE' in v or 'RESTANT' in v or 'FORMATEUR' in v: return True
    if 'SITE' in v or 'GROUPE' in v or 'CPFAE' in v: return True
    if 'REPUBLIQUE' in v or 'UNION' in v: return True
    return False

modules = {}  # module -> set(grade)
edt_files = [f for f in os.listdir(SRC) if 'Emploi' in f and f.endswith('.xlsx')]

for fname in sorted(edt_files):
    path = os.path.join(SRC, fname)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for shname in wb.sheetnames:
        ws = wb[shname]
        all_rows = list(ws.iter_rows(values_only=True))

        # Trouver grade
        grade = ''
        header_idx = None
        for i, row in enumerate(all_rows):
            vals = [str(v).strip().replace('\xa0',' ') if v else '' for v in row]
            full = ' '.join(vals).upper()
            if 'GRADE' in full and 'EMPLOI' in full:
                m = re.search(r'GRADE\s+([A-Z]\d)', full)
                if m: grade = m.group(1)
            if 'MATIERES' in full:
                header_idx = i
                break

        if header_idx is None:
            continue

        # Skip ligne "DU MODULE | RESTANT..."
        data_start = header_idx + 1
        if data_start < len(all_rows):
            nv = [str(v).upper() if v else '' for v in all_rows[data_start]]
            if any('DU MODULE' in v or 'RESTANT' in v for v in nv):
                data_start += 1

        current = ''
        for row in all_rows[data_start:]:
            first = str(row[0]).strip().replace('\xa0','') if row[0] else ''
            if not first or is_skip(first): continue
            if not re.match(r'^\d+\s*H?\s*$', first.upper()):
                current = first.strip()
                if current and not is_skip(current):
                    if current not in modules:
                        modules[current] = set()
                    modules[current].add(grade)
    wb.close()

print(f"\n{len(modules)} MODULES DISTINCTS DANS LES EDT :\n")
for mod in sorted(modules.keys()):
    print(f"  {mod:<50} grades: {sorted(modules[mod])}")
