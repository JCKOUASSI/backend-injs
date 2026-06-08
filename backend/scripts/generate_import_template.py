#!/usr/bin/env python
"""
Génère les fichiers Excel modèles pour le chargement de la base de données.

Fichiers générés :
  1. import_formations.xlsx / modele_formations.csv
  2. import_formateurs.xlsx / modele_formateurs.csv
  3. import_participants.xlsx / modele_participants.csv
  4. import_seances.xlsx / modele_seances.csv

Usage :
  python scripts/generate_import_template.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Styles ──────────────────────────────────────────────
header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill_green = PatternFill(start_color='388E3C', end_color='388E3C', fill_type='solid')
header_fill_orange = PatternFill(start_color='F57C00', end_color='F57C00', fill_type='solid')
header_fill_blue = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
example_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
note_font = Font(italic=True, color='888888', size=9)
thin_border = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)


def _write_sheet(ws, headers, examples, notes=None, header_fill=None):
    """Écrit les en-têtes, les exemples et les notes dans une feuille."""
    fill = header_fill or header_fill_green
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    for row_idx, example_row in enumerate(examples, 2):
        for col, val in enumerate(example_row, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = example_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    if notes:
        note_row = len(examples) + 3
        for col, note in enumerate(notes, 1):
            if note:
                cell = ws.cell(row=note_row, column=col, value=note)
                cell.font = note_font

    # Auto-width
    for col in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col).value or ''))
            for r in range(1, len(examples) + 3)
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 4, 15), 45)


# ═════════════════════════════════════════════════════════
# 1. FORMATIONS (modules uniquement)
# ═════════════════════════════════════════════════════════
FORMATIONS_HEADERS = [
    'N°', 'Formation', 'Module (titre)', 'Site', 'Bâtiment', 'Salle',
    'Date début', 'Date fin', 'Volume horaire (h)', 'Catégorie', 'Grade', 'Groupe', 'Vague',
]

FORMATIONS_EXAMPLES = [
    [1, 'FORMATION EN ADMINISTRATION DE BASE', 'Déontologie de la Fonction Publique',
     'CPFAE', '', 'SALLE A', '2026-05-05 08:00', '2026-05-09 17:00', 40, 'FAB A', 'A4', 'GROUPE 1', 'SESSION 2026'],
    [2, 'FORMATION EN ADMINISTRATION DE BASE', 'Protocole et Savoir-vivre',
     'CPFAE', '', 'SALLE A', '2026-05-12 08:00', '2026-05-13 17:00', 16, 'FAB A', 'A4', 'GROUPE 1', 'SESSION 2026'],
    [3, 'FORMATION EN ADMINISTRATION DE BASE', 'Finances Publiques',
     'CPFAE', '', 'SALLE A', '2026-05-19 08:00', '2026-05-22 17:00', 30, 'FAB A', 'A4', 'GROUPE 1', 'SESSION 2026'],
    [4, 'FORMATION EN ADMINISTRATION DE BASE', 'Déontologie de la Fonction Publique',
     'CPFAE', '', 'SALLE B', '2026-05-05 08:00', '2026-05-09 17:00', 40, 'FAB A', 'A4', 'GROUPE 2', 'SESSION 2026'],
]

FORMATEURS_HEADERS = [
    'Numéro', 'Nom', 'Prénom', 'E-mail', 'Téléphone', 'Spécialité', 'Organisation',
]

FORMATEURS_EXAMPLES = [
    ['F0001', 'CHRAIBI', 'Nadia', 'nadia.chraibi@expert.ci', '0700000001', 'Rédaction administrative', 'CPFAE'],
    ['F0002', 'BERRADA', 'Karim', 'karim.berrada@expert.ci', '0700000002', 'Droit administratif', 'ENA'],
    ['F0003', 'HAJJI', 'Leila', 'leila.hajji@univ.ci', '0700000003', 'Finances publiques', 'Université'],
]

PARTICIPANTS_HEADERS = [
    "N° d'inscription", 'Nom', 'Prénoms', 'Genre', 'Date de naissance', 'Lieu de naissance',
    'E-mail', 'Téléphone 1', 'Téléphone 2', 'Type concours', 'Libellé concours',
    'Catégorie', 'Grade', 'Groupe', 'Grade-Groupe', 'Vague', 'Formation(s)',
]

PARTICIPANTS_EXAMPLES = [
    ['FNCP26-001', 'KOUAME', 'Jean-Marc', 'MASCULIN', '15/03/1990', 'Abidjan',
     'jm.kouame@gouv.ci', '0701001001', '', 'Concours direct', 'Administrateur Civil',
     'FAB A', 'A4', 'GROUPE 1', 'A4/GROUPE 1', 'SESSION 2026', 'FORMATION EN ADMINISTRATION DE BASE'],
    ['FNCP26-002', 'DIALLO', 'Mariam', 'FEMININ', '22/07/1992', 'Bouaké',
     'diallo.m@gouv.ci', '0702002002', '', 'Concours direct', 'Administrateur Civil',
     'FAB A', 'A4', 'GROUPE 1', 'A4/GROUPE 1', 'SESSION 2026', 'FORMATION EN ADMINISTRATION DE BASE'],
    ['FNCP26-003', 'BAMBA', 'Oumar', 'MASCULIN', '05/11/1988', 'Korhogo',
     '', '0703003003', '', 'Concours direct', 'Administrateur Civil',
     'FAB A', 'A4', 'GROUPE 2', 'A4/GROUPE 2', 'SESSION 2026', 'FORMATION EN ADMINISTRATION DE BASE'],
]

SEANCES_HEADERS = [
    'module_titre', 'grade', 'groupe', 'vague',
    'date_journee', 'numero', 'intitule', 'heure_debut', 'heure_fin',
]

SEANCES_EXAMPLES = [
    ['Déontologie de la Fonction Publique', 'A4', 'GROUPE 1', 'SESSION 2026', '05/05/2026', 1, 'Matin', '08:30', '12:00'],
    ['Déontologie de la Fonction Publique', 'A4', 'GROUPE 1', 'SESSION 2026', '05/05/2026', 2, 'Après-midi', '14:00', '17:00'],
    ['Protocole et Savoir-vivre', 'A4', 'GROUPE 1', 'SESSION 2026', '12/05/2026', 1, 'Matin', '08:30', '12:00'],
]


def _write_csv(filepath, headers, examples):
    import csv
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f, delimiter=';')
        writer.writerow(headers)
        writer.writerows(examples)


def generate_formations():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Formations'
    _write_sheet(ws,
        headers=FORMATIONS_HEADERS,
        header_fill=header_fill_green,
        examples=FORMATIONS_EXAMPLES,
        notes=[
            'N° du module dans le cycle',
            'Titre du cycle — Obligatoire',
            'Intitulé du module — Obligatoire',
            'Centre (Lieu accepté)',
            'Optionnel',
            'Optionnel',
            'AAAA-MM-JJ HH:MM — Obligatoire',
            'AAAA-MM-JJ HH:MM — Obligatoire',
            'Nombre (heures)',
            'FAB A, FAB B ou FAB C',
            'A4, A3, B3…',
            'GROUPE 1, GROUPE 2…',
            'SESSION 2026… — Obligatoire pour les séances',
        ],
    )
    return wb


# ═════════════════════════════════════════════════════════
# 2. FORMATEURS (uniquement)
# ═════════════════════════════════════════════════════════
def generate_formateurs():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Formateurs'
    _write_sheet(ws,
        headers=FORMATEURS_HEADERS,
        header_fill=header_fill_blue,
        examples=FORMATEURS_EXAMPLES,
        notes=[
            'Numéro badge — Obligatoire',
            'Obligatoire',
            'Obligatoire',
            'Optionnel',
            'Optionnel',
            'Optionnel',
            'Optionnel',
        ],
    )
    return wb


# ═════════════════════════════════════════════════════════
# 3. PARTICIPANTS + libellé de leurs formations
# ═════════════════════════════════════════════════════════
def generate_participants():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Participants'
    _write_sheet(ws,
        headers=PARTICIPANTS_HEADERS,
        header_fill=header_fill_orange,
        examples=PARTICIPANTS_EXAMPLES,
        notes=[
            "N° d'inscription UNIQUE — Obligatoire",
            'Obligatoire',
            'Obligatoire',
            'MASCULIN / FEMININ',
            'JJ/MM/AAAA',
            'Optionnel',
            'Optionnel',
            'Optionnel',
            'Optionnel',
            'Optionnel',
            'Optionnel',
            'FAB A, FAB B ou FAB C',
            'A4, A3, B3…',
            'GROUPE 1, GROUPE 2…',
            'Optionnel (ex: A4/GROUPE 1)',
            'SESSION 2026…',
            'Titre(s) séparés par |',
        ],
    )
    return wb


def generate_seances():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Séances'
    _write_sheet(ws,
        headers=SEANCES_HEADERS,
        header_fill=PatternFill(start_color='5C6BC0', end_color='5C6BC0', fill_type='solid'),
        examples=SEANCES_EXAMPLES,
        notes=[
            'Intitulé du module — identique à l\'import Cours',
            'A4, B3… — Obligatoire',
            'GROUPE 1… — Obligatoire',
            'SESSION 2026… — Obligatoire',
            'JJ/MM/AAAA',
            '1, 2, 3…',
            'Matin, Après-midi…',
            'HH:MM',
            'HH:MM',
        ],
    )
    return wb


def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    outputs = []

    wb = generate_formations()
    p1 = os.path.join(base_dir, 'import_formations.xlsx')
    wb.save(p1)
    outputs.append((p1, 'Cours [Excel]'))

    c1 = os.path.join(base_dir, 'modele_formations.csv')
    _write_csv(c1, FORMATIONS_HEADERS, FORMATIONS_EXAMPLES)
    outputs.append((c1, 'Cours [CSV]'))

    wb = generate_formateurs()
    p2 = os.path.join(base_dir, 'import_formateurs.xlsx')
    wb.save(p2)
    outputs.append((p2, 'Formateurs [Excel]'))

    c2 = os.path.join(base_dir, 'modele_formateurs.csv')
    _write_csv(c2, FORMATEURS_HEADERS, FORMATEURS_EXAMPLES)
    outputs.append((c2, 'Formateurs [CSV]'))

    wb = generate_participants()
    p3 = os.path.join(base_dir, 'import_participants.xlsx')
    wb.save(p3)
    outputs.append((p3, 'Auditeurs [Excel]'))

    c3 = os.path.join(base_dir, 'modele_participants.csv')
    _write_csv(c3, PARTICIPANTS_HEADERS, PARTICIPANTS_EXAMPLES)
    outputs.append((c3, 'Auditeurs [CSV]'))
    # Alias historique
    c3b = os.path.join(base_dir, 'import_participants.csv')
    _write_csv(c3b, PARTICIPANTS_HEADERS, PARTICIPANTS_EXAMPLES)
    outputs.append((c3b, 'Auditeurs [CSV alias]'))

    wb = generate_seances()
    p4 = os.path.join(base_dir, 'import_seances.xlsx')
    wb.save(p4)
    outputs.append((p4, 'Séances [Excel]'))

    c4 = os.path.join(base_dir, 'modele_seances.csv')
    _write_csv(c4, SEANCES_HEADERS, SEANCES_EXAMPLES)
    outputs.append((c4, 'Séances [CSV]'))

    for path, label in outputs:
        print(f'✅ {path}')
        print(f'   → {label}')

    print()
    print("Ordre d'import :")
    print('  1. python manage.py import_excel import_formations.xlsx')
    print('  2. python manage.py import_excel import_formateurs.xlsx')
    print('  3. python manage.py import_excel import_participants.xlsx')
    print('  4. python manage.py import_excel import_seances.xlsx')


if __name__ == '__main__':
    main()
