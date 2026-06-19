"""Tests pour import_excel — résolution RefSite et endpoint API."""
from datetime import datetime
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework import status
from rest_framework.test import APIClient

from authentication.models import User
from formations.management.commands.import_excel import Command
from formations.models import (
    Formation, Module, RefSite, RefTypeSecretariat, Secretariat, SessionModule,
)


FORMATION_HEADERS = (
    'Formation',
    'Module (titre)',
    'Site',
    'Bâtiment',
    'Salle',
    'Date début',
    'Date fin',
    'Volume horaire (h)',
    'Catégorie',
    'Grade',
    'Groupe',
)


def _formation_defaults(**overrides):
    row = {
        'Formation': 'Cycle import test',
        'Module (titre)': 'Module site test',
        'Site': 'CPFAE',
        'Bâtiment': 'Bâtiment Principal',
        'Salle': 'Salle Polyvalente',
        'Date début': '2026-05-05 08:00',
        'Date fin': '2026-05-09 17:00',
        'Volume horaire (h)': 40,
        'Catégorie': 'FAB A',
        'Grade': 'A4',
        'Groupe': 'GROUPE 1',
    }
    row.update(overrides)
    return row


class _MemorySheet:
    def __init__(self, rows):
        self._rows = rows

    def iter_rows(self, values_only=True):
        return iter(self._rows)


class _SilentStdout:
    def write(self, msg):
        pass


def build_formations_sheet(*data_rows):
    rows = [FORMATION_HEADERS]
    for overrides in data_rows:
        merged = _formation_defaults(**overrides)
        rows.append(tuple(merged[h] for h in FORMATION_HEADERS))
    return _MemorySheet(rows)


def build_formations_xlsx(*data_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Formations'
    ws.append(list(FORMATION_HEADERS))
    for overrides in data_rows:
        merged = _formation_defaults(**overrides)
        ws.append([merged[h] for h in FORMATION_HEADERS])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ImportFormationsRefSiteTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        ref_type = RefTypeSecretariat.objects.create(libelle='FAB A')
        Secretariat.objects.create(nom='Secrétariat FAB A', type=ref_type)

    def setUp(self):
        self.cmd = Command()
        self.cmd.stdout = _SilentStdout()

    def test_import_creates_ref_site_fk(self):
        errors = []
        created, updated = self.cmd._import_formations(
            build_formations_sheet({}), errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 1)
        self.assertEqual(updated, 0)
        self.assertEqual(RefSite.objects.filter(nom='CPFAE').count(), 1)

        mod = Module.objects.get(intitule='Module site test')
        self.assertEqual(mod.site.nom, 'CPFAE')
        self.assertEqual(mod.site_legacy, 'CPFAE')
        self.assertEqual(mod.batiment, 'Bâtiment Principal')
        self.assertEqual(mod.salle, 'Salle Polyvalente')

    def test_import_reuses_existing_ref_site(self):
        existing = RefSite.objects.create(nom='CPFAE', actif=True)
        errors = []
        self.cmd._import_formations(build_formations_sheet({}), errors)
        self.assertEqual(RefSite.objects.filter(nom='CPFAE').count(), 1)
        mod = Module.objects.get(intitule='Module site test')
        self.assertEqual(mod.site_id, existing.id)

    def test_import_without_site_leaves_fk_null(self):
        errors = []
        created, _ = self.cmd._import_formations(
            build_formations_sheet({'Site': ''}), errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 1)
        mod = Module.objects.get(intitule='Module site test')
        self.assertIsNone(mod.site_id)
        self.assertEqual(mod.site_legacy, '')

    def test_import_update_preserves_site_fk(self):
        errors = []
        self.cmd._import_formations(build_formations_sheet({}), errors)
        created, updated = self.cmd._import_formations(
            build_formations_sheet({'Volume horaire (h)': 32}), errors,
        )
        self.assertEqual(created, 0)
        self.assertEqual(updated, 1)
        mod = Module.objects.get(intitule='Module site test')
        self.assertIsInstance(mod.site, RefSite)
        self.assertEqual(float(mod.duree_prevue_heures), 32)

    def test_parse_datetime_malformed_excel_dates(self):
        cases = [
            (' 15/06/ 2026', datetime(2026, 6, 15)),
            ('03/072026', datetime(2026, 7, 3)),
            ('23/072026', datetime(2026, 7, 23)),
            ('26/072026', datetime(2026, 7, 26)),
        ]
        for raw, expected in cases:
            parsed = self.cmd._parse_datetime(raw)
            self.assertIsNotNone(parsed, msg=raw)
            self.assertEqual(parsed.date(), expected.date(), msg=raw)

    def test_import_real_workbook_malformed_dates(self):
        from pathlib import Path
        from openpyxl import load_workbook
        path = Path(__file__).resolve().parent.parent / 'importverif' / 'MODELE FORMATION A4 GROUPE 4-15 excel.xlsx'
        if not path.exists():
            self.skipTest('fichier de vérification absent')
        wb = load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        errors = []
        created, updated = self.cmd._import_formations(ws, errors)
        wb.close()
        date_errors = [e for e in errors if 'date_debut ou date_fin invalide' in e]
        self.assertEqual(date_errors, [], msg=date_errors)
        self.assertGreater(created + updated, 0)


class ImportExcelAPITest(TestCase):
    @classmethod
    def setUpTestData(cls):
        ref_type = RefTypeSecretariat.objects.create(libelle='FAB A')
        Secretariat.objects.create(nom='Secrétariat FAB A', type=ref_type)

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='import_api', password='pass', role='CPFAE_ADMIN',
        )
        self.client.force_authenticate(self.user)

    def test_api_import_formations_resolves_site(self):
        xlsx = build_formations_xlsx({})
        res = self.client.post(
            '/api/formations/import-excel/',
            {
                'file': SimpleUploadedFile(
                    'formations.xlsx',
                    xlsx.read(),
                    content_type=(
                        'application/vnd.openxmlformats-officedocument'
                        '.spreadsheetml.sheet'
                    ),
                ),
                'type': 'formations',
            },
            format='multipart',
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertEqual(res.data['created'], 1)
        self.assertEqual(res.data['errors'], [])
        mod = Module.objects.get(intitule='Module site test')
        self.assertEqual(mod.site.nom, 'CPFAE')


SEANCE_HEADERS = (
    'module_titre', 'grade', 'groupe', 'vague',
    'date_journee', 'numero', 'intitule', 'heure_debut', 'heure_fin',
)


def _seance_defaults(**overrides):
    row = {
        'module_titre': 'Module site test',
        'grade': 'A4',
        'groupe': 'GROUPE 1',
        'vague': 'SESSION 2026',
        'date_journee': '05/05/2026',
        'numero': 1,
        'intitule': 'Matin',
        'heure_debut': '08:30',
        'heure_fin': '12:00',
    }
    row.update(overrides)
    return row


def build_seances_sheet(*data_rows):
    rows = [SEANCE_HEADERS]
    for overrides in data_rows:
        merged = _seance_defaults(**overrides)
        rows.append(tuple(merged[h] for h in SEANCE_HEADERS))
    return _MemorySheet(rows)


class ImportSeancesMatchTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        ref_type = RefTypeSecretariat.objects.create(libelle='FAB A')
        Secretariat.objects.create(nom='Secrétariat FAB A', type=ref_type)
        formation = Formation.objects.create(formation='Cycle import test')
        cls.module = Module.objects.create(
            formation=formation,
            intitule='Module site test',
            grade='A4',
            groupe='GROUPE 1',
            vague='SESSION 2026',
            cycle='Cycle import test',
        )

    def setUp(self):
        self.cmd = Command()
        self.cmd.stdout = _SilentStdout()

    def test_import_seance_matches_module_by_grade_groupe_vague(self):
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet({}), errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 1)
        self.assertEqual(updated, 0)
        self.assertTrue(
            SessionModule.objects.filter(
                module=self.module, date_journee='2026-05-05', numero=1,
            ).exists()
        )

    def test_import_seance_rejects_missing_grade_groupe_vague(self):
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet({'grade': '', 'groupe': '', 'vague': ''}), errors,
        )
        self.assertEqual(created, 0)
        self.assertEqual(updated, 0)
        self.assertEqual(len(errors), 1)
        self.assertIn('grade', errors[0])
        self.assertIn('groupe', errors[0])
        self.assertIn('vague', errors[0])

    def test_import_seance_rejects_wrong_groupe(self):
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet({'groupe': 'GROUPE 99'}), errors,
        )
        self.assertEqual(created, 0)
        self.assertEqual(len(errors), 1)
        self.assertIn('introuvable', errors[0])
        self.assertIn('GROUPE 99', errors[0])

    def test_import_seance_carries_forward_missing_grade(self):
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet(
                {},
                {'grade': '', 'numero': 2, 'date_journee': '06/05/2026'},
            ),
            errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 2)
        self.assertEqual(updated, 0)

    def test_import_seance_ignores_wrong_grade_when_unique(self):
        self.module.grade = 'A5'
        self.module.save(update_fields=['grade'])
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet({'grade': 'A4'}), errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 1)

    def test_import_seance_fuzzy_title_typo(self):
        formation = self.module.formation
        Module.objects.create(
            formation=formation,
            intitule='DROIT ADMINISTRATIF',
            grade='A4',
            groupe='GROUPE 13',
            vague='SESSION 2026',
            cycle='Cycle import test',
        )
        errors = []
        created, updated = self.cmd._import_seances(
            build_seances_sheet({
                'module_titre': 'DROIT ADMINISTRATIVF',
                'groupe': 'GROUPE 13',
                'vague': 'SESSION 2026',
            }),
            errors,
        )
        self.assertEqual(errors, [])
        self.assertEqual(created, 1)

    def test_import_real_seances_workbook(self):
        from pathlib import Path
        from openpyxl import load_workbook

        path = Path(__file__).resolve().parent.parent / 'importverif' / 'MODELE DE SEANCE A4 G9-15.xlsx'
        form_path = Path(__file__).resolve().parent.parent / 'importverif' / 'MODELE FORMATION A4 GROUPE 4-15 excel.xlsx'
        if not path.exists() or not form_path.exists():
            self.skipTest('fichiers de vérification absents')

        form_wb = load_workbook(form_path, read_only=True)
        form_errors = []
        self.cmd._import_formations(form_wb[form_wb.sheetnames[0]], form_errors)
        form_wb.close()
        self.assertEqual(form_errors, [], msg=form_errors)

        wb = load_workbook(path, read_only=True)
        errors = []
        created, updated = self.cmd._import_seances(wb[wb.sheetnames[0]], errors)
        wb.close()
        self.assertEqual(errors, [], msg=errors)
        self.assertGreater(created + updated, 0)


PARTICIPANT_HEADERS = (
    "N° d'inscription", 'Nom', 'Prénoms', 'Genre', 'Date de naissance', 'Lieu de naissance',
    'E-mail', 'Téléphone 1', 'Téléphone 2', 'Type concours', 'Libellé concours',
    'Catégorie', 'Grade', 'Groupe', 'Grade-Groupe', 'Vague', 'Formation(s)',
)


def build_participants_sheet(*data_rows):
    defaults = {
        "N° d'inscription": 'FNCE25-0001',
        'Nom': 'DUPONT',
        'Prénoms': 'Alice',
        'Genre': 'FEMININ',
        'Date de naissance': '01/01/1990',
        'Lieu de naissance': 'Abidjan',
        'E-mail': 'alice@example.com',
        'Téléphone 1': '0700000000',
        'Téléphone 2': '',
        'Type concours': 'RECRUTEMENT',
        'Libellé concours': 'Test',
        'Catégorie': 'A',
        'Grade': 'A4',
        'Groupe': 'GROUPE 9',
        'Grade-Groupe': 'A4-GROUPE 9',
        'Vague': 'SESSION 2026 VAGUE 2',
        'Formation(s)': '',
    }
    rows = [PARTICIPANT_HEADERS]
    for overrides in data_rows:
        merged = {**defaults, **overrides}
        rows.append(tuple(merged[h] for h in PARTICIPANT_HEADERS))
    return _MemorySheet(rows)


class ImportParticipantsAutoMatchTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        ref_type = RefTypeSecretariat.objects.create(libelle='FAB A')
        Secretariat.objects.create(nom='Secrétariat FAB A', type=ref_type)
        formation = Formation.objects.create(formation='FORMATION EN ADMINISTRATION DE BASE')
        for intitule in ('DEONTOLOGIE', 'CULTURE CIVIQUE', 'REDACTION ADMINISTRATIVE'):
            Module.objects.create(
                formation=formation,
                intitule=intitule,
                grade='A4',
                groupe='GROUPE 9',
                vague='SESSION 2026 VAGUE 2',
                cycle='FORMATION EN ADMINISTRATION DE BASE',
            )

    def setUp(self):
        self.cmd = Command()
        self.cmd.stdout = _SilentStdout()

    def test_auto_inscription_all_modules_of_group(self):
        from formations.models import ModuleParticipant
        errors = []
        created = self.cmd._import_participants(build_participants_sheet({}), errors)
        self.assertEqual(errors, [], msg=errors)
        self.assertEqual(created, 1)
        participant = ModuleParticipant.objects.filter(
            participant__matricule='FNCE25-0001',
        ).count()
        self.assertEqual(participant, 3)

    def test_inscription_with_formation_column(self):
        from formations.models import ModuleParticipant
        errors = []
        created = self.cmd._import_participants(
            build_participants_sheet({
                'Formation(s)': 'FORMATION EN ADMINISTRATION DE BASE',
            }),
            errors,
        )
        self.assertEqual(errors, [], msg=errors)
        self.assertEqual(created, 1)
        self.assertEqual(
            ModuleParticipant.objects.filter(participant__matricule='FNCE25-0001').count(),
            3,
        )


class ImportCoherenceWorkbookTest(TestCase):
    """Cohérence inter-feuilles du fichier de test officiel."""

    def test_test_import_workbook_is_coherent(self):
        from pathlib import Path
        from openpyxl import load_workbook

        from formations.import_coherence import check_workbook

        path = Path(__file__).resolve().parent.parent / 'test_import_formations_seances.xlsx'
        self.assertTrue(path.exists(), f'Fichier de test manquant : {path}')
        wb = load_workbook(path, read_only=True, data_only=True)
        report = check_workbook(wb)
        wb.close()
        if report.errors:
            self.fail('Incohérences détectées :\n' + '\n'.join(f'  • {e}' for e in report.errors))
        self.assertGreater(report.stats.get('modules_distincts', 0), 0)
        self.assertGreater(report.stats.get('seances_lignes', 0), 0)
        self.assertEqual(report.stats.get('seances_orphelines', -1), 0)


class ImportPipelineCoherenceTest(TestCase):
    """Import complet formations → participants → séances sans erreur."""

    @classmethod
    def setUpTestData(cls):
        for libelle in ('FAB A', 'FAB B'):
            ref = RefTypeSecretariat.objects.create(libelle=libelle)
            Secretariat.objects.create(nom=f'Secrétariat {libelle}', type=ref)
        Secretariat.objects.create(nom='FAC')
        Secretariat.objects.create(nom='FAB')

    def setUp(self):
        self.cmd = Command()
        self.cmd.stdout = _SilentStdout()

    def test_full_pipeline_dry_run(self):
        from pathlib import Path
        from openpyxl import load_workbook

        path = Path(__file__).resolve().parent.parent / 'test_import_formations_seances.xlsx'
        wb = load_workbook(path, read_only=True)
        errors = []
        try:
            self.cmd._import_formations(wb['Formations'], errors)
            self.cmd._import_formateurs(wb['Formateurs'], errors)
            self.cmd._import_participants(wb['Participants'], errors)
            self.cmd._import_seances(wb['Séances'], errors)
        finally:
            wb.close()
        self.assertEqual(errors, [], '\n'.join(errors))
