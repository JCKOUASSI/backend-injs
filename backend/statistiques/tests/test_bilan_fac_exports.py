"""Tests exports bilan FAC."""
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from authentication.models import User
from formations.models import Formation, Module
from statistiques.bilan_fac_exports import export_excel_fac, _apply_fac_meta
from statistiques.bilans import compute_bilan_fac


class BilanFACExportsTest(TestCase):
    def setUp(self):
        self.formation = Formation.objects.create(formation='Formation Export FAC')
        Module.objects.create(
            formation=self.formation,
            intitule='Module test',
            grade='A4',
            groupe='G1',
        )
        self.user = User.objects.create_user(
            username='fac_export', password='test', role='CPFAE_ADMIN',
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_export_excel_fac_workbook_sheets(self):
        data = compute_bilan_fac(formation_id=self.formation.id)
        self.assertIsNotNone(data)
        buffer = export_excel_fac(data)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        self.assertIn('Point_global', wb.sheetnames)
        self.assertIn('VH_par_groupe', wb.sheetnames)
        self.assertIn('Absents_notoires', wb.sheetnames)
        self.assertIn('Modules', wb.sheetnames)

    def test_build_export_injects_meta(self):
        data = compute_bilan_fac(formation_id=self.formation.id)
        meta = {'justificatifs': {'A4': 'Maladie'}, 'difficultes': {'A4': 'Salle indisponible'}}
        patched = _apply_fac_meta(data, meta)
        ligne = patched['point_global']['lignes'][0]
        self.assertEqual(ligne['justificatifs'], 'Maladie')
        self.assertEqual(ligne['difficultes'], 'Salle indisponible')

    def test_api_bilan_fac_export_xlsx(self):
        res = self.client.get(
            '/api/statistiques/bilan-fac-export/',
            {'export': 'xlsx', 'formation_id': self.formation.id},
        )
        self.assertEqual(res.status_code, 200)
        self.assertIn(
            'spreadsheetml',
            res.get('Content-Type', ''),
        )
        self.assertGreater(len(res.content), 200)
