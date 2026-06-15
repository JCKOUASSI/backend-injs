"""
Exports Point Journalier — reproduction du format Excel CPFAE.
"""
import io
import re
from datetime import datetime

from django.http import HttpResponse

from .point_journalier import compute_point_journalier, MOIS_FR

# Couleurs extraites du modèle CPFAE « POINT JOURNALIER CAT A » (Office theme accent2 + custom)
FILL_HEADER = 'FFED7D31'       # Orange en-têtes (theme accent2)
FILL_GROUPE = 'FFFBFDBB'       # Jaune clair — ligne GROUPES
FILL_TOTAL = 'FFF7FA82'        # Jaune — colonne TOTAL
FILL_SIDEBAR = 'FFFFFF00'      # Jaune vif — bande SECONDE VAGUE
FILL_ABSENCE = 'FFED7D31'      # Orange — cellules taux d'absence
FILL_PRESENCE_JOUR = 'FFED7D31'
FILL_ABSENCE_JOUR = 'FFBDD7EE' # Bleu clair — taux absence du jour
FONT_ABSENCE = 'FFFFFFFF'      # Blanc — pourcentages d'absence (fond orange)
FMT_PCT = '0.00%'

# Grille fixe modèle CPFAE (17 colonnes A→Q)
COL_LABEL = 1
COL_GROUP_START = 2
COL_TOTAL = 14
COL_SIDEBAR = 15
COL_END = 17
MAX_GROUP_COLS = 12
ROW_FOOTER = 22
ROW_FOOTER_END = 24
PDF_GROUPES_PAR_TABLE = 14


def _chunk_total_groupes(groupes):
    effectif = sum(g['effectif'] for g in groupes)
    presents = sum(g['presents'] for g in groupes)
    absents = sum(g['absents'] for g in groupes)
    denom = presents + absents
    return {
        'effectif': effectif,
        'presents': presents,
        'absents': absents,
        'taux_presence': (presents / denom) if denom else 0.0,
        'taux_absence': (absents / denom) if denom else 0.0,
    }


def _iter_groupe_chunks(groupes, chunk_size=PDF_GROUPES_PAR_TABLE):
    """Découpe les groupes en blocs affichables sur une page PDF."""
    if not groupes:
        return
    n_chunks = (len(groupes) + chunk_size - 1) // chunk_size
    for i in range(0, len(groupes), chunk_size):
        chunk = groupes[i:i + chunk_size]
        yield chunk, _chunk_total_groupes(chunk), (i // chunk_size) + 1, n_chunks


def _pdf_font_size(n_groupes):
    if n_groupes <= 8:
        return 7
    if n_groupes <= 12:
        return 6
    if n_groupes <= 16:
        return 5.5
    return 5


def _pdf_col_widths(n_groupes, page_width, cm):
    """Largeurs de colonnes pour tenir dans la page paysage."""
    label_w = 2.6 * cm
    spacer_w = 0.55 * cm
    total_w = 1.15 * cm
    avail = page_width - label_w - spacer_w - total_w
    if n_groupes <= 0:
        return [label_w, spacer_w, total_w]
    gw = max(0.42 * cm, avail / n_groupes)
    return [label_w, spacer_w] + [gw] * n_groupes + [total_w]


def _pdf_bloc_tables(bloc, page_width, cm, colors, Paragraph, ParagraphStyle, styles, TA_CENTER):
    """Construit une ou plusieurs tables ReportLab pour un créneau MATIN/SOIR."""
    from reportlab.platypus import Table, TableStyle

    st_cell = ParagraphStyle(
        'PJCell', parent=styles['Normal'], fontSize=6, leading=7, alignment=TA_CENTER,
    )
    st_label = ParagraphStyle(
        'PJLabel', parent=styles['Normal'], fontSize=6, leading=7, alignment=TA_CENTER,
    )
    st_hdr = ParagraphStyle(
        'PJHdr', parent=styles['Normal'], fontSize=6, leading=7, alignment=TA_CENTER,
        fontName='Helvetica-Bold',
    )

    def _p(text, hdr=False):
        s = st_hdr if hdr else st_cell
        safe = str(text).replace('&', '&amp;').replace('<', '&lt;')
        return Paragraph(safe, s)

    tables = []
    groupes = bloc['groupes'] or []
    for chunk, chunk_total, part, n_parts in _iter_groupe_chunks(groupes):
        n = len(chunk)
        font_size = _pdf_font_size(n)
        col_widths = _pdf_col_widths(n, page_width, cm)

        hdr = [_p('', hdr=True), _p('GROUPES', hdr=True)] + [_p(g['label'], hdr=True) for g in chunk]
        if n_parts > 1:
            hdr.append(_p(f'TOT.{part}', hdr=True))
        else:
            hdr.append(_p('TOTAL', hdr=True))

        rows = [
            hdr,
            [_p(''), _p('SALLES')] + [_p(g['salle'] or '—') for g in chunk] + [_p('')],
        ]
        total_row = chunk_total if n_parts > 1 else bloc['total']
        for lbl, key, pct in [
            ('ÉFFECTIF', 'effectif', False),
            ('PRÉSENTS', 'presents', False),
            ('ABSENTS', 'absents', False),
            ('TAUX PRÉS.', 'taux_presence', True),
            ('TAUX ABS.', 'taux_absence', True),
        ]:
            row = [_p(lbl), _p('')] + [
                _p(f"{g[key]:.2%}" if pct else str(g[key])) for g in chunk
            ]
            tv = total_row[key]
            row.append(_p(f"{tv:.2%}" if pct else str(tv)))
            rows.append(row)

        t = Table(rows, colWidths=col_widths, repeatRows=2)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FBFBDB')),
            ('BACKGROUND', (-1, 0), (-1, -1), colors.HexColor('#F7FA82')),
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F8FAFC')),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('FONTSIZE', (0, 0), (-1, -1), font_size),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        tables.append(t)
    return tables


def _safe_filename(text):
    return re.sub(r'[^\w\-]+', '_', text, flags=re.UNICODE).strip('_')[:80]


def _filter_tableaux(data, formation_id=None, categorie=None, mois=None, jour=None):
    tableaux = data['tableaux']
    if formation_id:
        tableaux = [t for t in tableaux if t['formation_id'] == formation_id]
    if categorie:
        tableaux = [t for t in tableaux if t['categorie'].upper() == categorie.upper()]
    if mois:
        tableaux = [t for t in tableaux if t['mois'] == mois]
    if jour:
        tableaux = [t for t in tableaux if t['date'] == jour]
    return tableaux


def _write_bloc_creneau(ws, start_row, bloc, label_creneau, styles):
    """Écrit MATIN ou SOIR — format CPFAE (colonnes B→M groupes, N total)."""
    from openpyxl.utils import get_column_letter

    groupes = (bloc.get('groupes') or [])[:MAX_GROUP_COLS]
    n_groups = len(groupes)
    last_group_col = COL_GROUP_START + max(n_groups, 1) - 1
    data_end_col = COL_TOTAL - 1
    groupes_row = start_row + 2
    salles_row = start_row + 3
    first_data_row = start_row + 4
    last_data_row = start_row + 8

    # Ligne créneau (ex. MATIN: 08H00-12H00)
    ws.merge_cells(
        start_row=start_row, start_column=COL_LABEL,
        end_row=start_row, end_column=COL_TOTAL,
    )
    c = ws.cell(row=start_row, column=COL_LABEL, value=f"{label_creneau}: {bloc['horaire']}")
    c.font = styles['font_creneau']
    c.fill = styles['fill_header']
    c.alignment = styles['align_center']
    c.border = styles['thin']
    start_row += 1

    # Ligne vide (spacer) + fusion colonne TOTAL (N4:N6 équivalent)
    ws.merge_cells(
        start_row=start_row, start_column=COL_LABEL,
        end_row=start_row, end_column=data_end_col,
    )
    ws.merge_cells(
        start_row=start_row, start_column=COL_TOTAL,
        end_row=start_row + 2, end_column=COL_TOTAL,
    )
    for col in range(COL_LABEL, COL_TOTAL + 1):
        ws.cell(row=start_row, column=col).border = styles['thin']
    start_row += 1

    # GROUPES
    ws.cell(row=start_row, column=COL_LABEL, value='GROUPES').font = styles['font_label']
    ws.cell(row=start_row, column=COL_LABEL).border = styles['thin']
    for i, g in enumerate(groupes):
        col = COL_GROUP_START + i
        cell = ws.cell(row=start_row, column=col, value=g['label'])
        cell.font = styles['font_groupe']
        cell.fill = styles['fill_groupe']
        cell.alignment = styles['align_center']
        cell.border = styles['thin']
    if n_groups < MAX_GROUP_COLS and last_group_col < data_end_col:
        ws.merge_cells(
            start_row=groupes_row, start_column=last_group_col + 1,
            end_row=last_data_row, end_column=data_end_col,
        )
    ws.cell(row=start_row, column=COL_TOTAL).fill = styles['fill_total']
    ws.cell(row=start_row, column=COL_TOTAL).border = styles['thin']
    start_row += 1

    # SALLES
    ws.cell(row=start_row, column=COL_LABEL, value='SALLES').font = styles['font_label']
    ws.cell(row=start_row, column=COL_LABEL).border = styles['thin']
    for i, g in enumerate(groupes):
        col = COL_GROUP_START + i
        cell = ws.cell(row=start_row, column=col, value=g['salle'])
        cell.font = styles['font_salle']
        cell.alignment = styles['align_center']
        cell.border = styles['thin']
    ws.cell(row=start_row, column=COL_TOTAL).fill = styles['fill_total']
    ws.cell(row=start_row, column=COL_TOTAL).border = styles['thin']
    start_row += 1

    rows_def = [
        ('ÉFFECTIF', 'effectif', False, False),
        ('PRÉSENTS', 'presents', False, False),
        ('ABSENTS', 'absents', True, False),
        ('TAUX DE PRÉSENCE', 'taux_presence', True, True),
        ("TAUX D'ABSENCE", 'taux_absence', True, True),
    ]

    for row_label, key, bold, is_pct in rows_def:
        label_cell = ws.cell(row=start_row, column=COL_LABEL, value=row_label)
        label_cell.font = styles['font_bold'] if bold else styles['font_label']
        label_cell.border = styles['thin']
        is_absence_row = key == 'taux_absence'
        for i, g in enumerate(groupes):
            col = COL_GROUP_START + i
            val = g[key]
            cell = ws.cell(row=start_row, column=col, value=val)
            cell.alignment = styles['align_center']
            cell.border = styles['thin']
            if is_pct:
                cell.number_format = FMT_PCT
                cell.font = styles['font_bold']
                if is_absence_row:
                    cell.font = styles['font_absence']
                    cell.fill = styles['fill_absence']
            elif bold:
                cell.font = styles['font_bold']

        tot = bloc['total']
        tval = tot[key]
        tcell = ws.cell(row=start_row, column=COL_TOTAL, value=tval)
        tcell.fill = styles['fill_total'] if key != 'taux_absence' else styles['fill_absence']
        tcell.font = styles['font_total']
        tcell.alignment = styles['align_center']
        tcell.border = styles['thin']
        if is_pct:
            tcell.number_format = FMT_PCT
            if is_absence_row:
                tcell.font = styles['font_absence_bold']
        start_row += 1

    # Bordures zone données (libellés + colonnes groupes + total)
    for r in range(groupes_row, start_row):
        for col in range(COL_LABEL, COL_TOTAL + 1):
            ws.cell(row=r, column=col).border = styles['thin']

    return start_row + 1


def _write_sidebar(ws, row_start, row_end, text, styles):
    from openpyxl.styles import Alignment

    ws.merge_cells(
        start_row=row_start, start_column=COL_SIDEBAR,
        end_row=row_end, end_column=COL_END,
    )
    cell = ws.cell(row=row_start, column=COL_SIDEBAR, value=(text or 'SECONDE VAGUE').strip())
    cell.fill = styles['fill_sidebar']
    cell.font = styles['font_sidebar']
    cell.alignment = Alignment(
        horizontal='center', vertical='center', text_rotation=90, wrap_text=True,
    )


def _write_sheet(ws, tb, styles):
    """Remplit une feuille au format modèle CPFAE (17 colonnes, bande vague)."""
    from openpyxl.utils import get_column_letter

    # Ligne 1 — titre
    ws.merge_cells(start_row=1, start_column=COL_LABEL, end_row=1, end_column=COL_TOTAL)
    c1 = ws.cell(row=1, column=COL_LABEL, value=tb['titre_ligne1'])
    c1.font = styles['font_title']
    c1.fill = styles['fill_header']
    c1.alignment = styles['align_center']
    c1.border = styles['thin']

    # Ligne 2 — DATE
    ws.merge_cells(start_row=2, start_column=COL_LABEL, end_row=2, end_column=3)
    for col, val in [(COL_LABEL, 'DATE'), (4, tb['jour']), (5, tb['mois_libelle']), (6, tb['annee'])]:
        cell = ws.cell(row=2, column=col, value=val)
        cell.font = styles['font_date']
        cell.border = styles['thin']
    ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=COL_TOTAL)
    org = ws.cell(row=2, column=7, value=tb['organisme'])
    org.font = styles['font_date']
    org.alignment = styles['align_center']
    org.border = styles['thin']

    row = 3
    matin_start = row
    row = _write_bloc_creneau(ws, row, tb['matin'], 'MATIN', styles)
    row = _write_bloc_creneau(ws, row, tb['soir'], 'SOIR', styles)

    # Pied — taux du jour
    ws.merge_cells(start_row=ROW_FOOTER, start_column=COL_LABEL, end_row=ROW_FOOTER_END, end_column=2)
    ws.cell(row=ROW_FOOTER, column=COL_LABEL, value='Taux de présence du jour').font = styles['font_bold']
    ws.cell(row=ROW_FOOTER, column=COL_LABEL).alignment = styles['align_center_wrap']

    ws.merge_cells(start_row=ROW_FOOTER, start_column=3, end_row=ROW_FOOTER_END, end_column=6)
    cp = ws.cell(row=ROW_FOOTER, column=3, value=tb['taux_presence_jour'])
    cp.number_format = FMT_PCT
    cp.font = styles['font_taux_jour']
    cp.fill = styles['fill_presence_jour']
    cp.alignment = styles['align_center']

    ws.merge_cells(start_row=ROW_FOOTER, start_column=7, end_row=ROW_FOOTER_END, end_column=8)
    ws.cell(row=ROW_FOOTER, column=7, value="Taux d'absence du jour").font = styles['font_bold']
    ws.cell(row=ROW_FOOTER, column=7).alignment = styles['align_center_wrap']

    ws.merge_cells(start_row=ROW_FOOTER, start_column=9, end_row=ROW_FOOTER_END, end_column=COL_TOTAL)
    ca = ws.cell(row=ROW_FOOTER, column=9, value=tb['taux_absence_jour'])
    ca.number_format = FMT_PCT
    ca.font = styles['font_taux_jour']
    ca.fill = styles['fill_absence_jour']
    ca.alignment = styles['align_center']

    _write_sidebar(ws, 1, ROW_FOOTER_END, tb.get('vague_sidebar', 'SECONDE VAGUE'), styles)

    # Largeurs colonnes (modèle CPFAE)
    widths = {
        'A': 19.37, 'B': 35.51, 'C': 21.93, 'D': 38.47, 'E': 35.51,
        'F': 29.05, 'G': 15.47, 'N': 15.47, 'O': 13.0, 'P': 13.0, 'Q': 13.0,
    }
    for i in range(1, COL_END + 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(letter, 13.0)


def _make_styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    thin = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000'),
    )
    return {
        'thin': thin,
        'fill_header': PatternFill(start_color=FILL_HEADER, end_color=FILL_HEADER, fill_type='solid'),
        'fill_groupe': PatternFill(start_color=FILL_GROUPE, end_color=FILL_GROUPE, fill_type='solid'),
        'fill_total': PatternFill(start_color=FILL_TOTAL, end_color=FILL_TOTAL, fill_type='solid'),
        'fill_sidebar': PatternFill(start_color=FILL_SIDEBAR, end_color=FILL_SIDEBAR, fill_type='solid'),
        'fill_absence': PatternFill(start_color=FILL_ABSENCE, end_color=FILL_ABSENCE, fill_type='solid'),
        'fill_presence_jour': PatternFill(start_color=FILL_PRESENCE_JOUR, end_color=FILL_PRESENCE_JOUR, fill_type='solid'),
        'fill_absence_jour': PatternFill(start_color=FILL_ABSENCE_JOUR, end_color=FILL_ABSENCE_JOUR, fill_type='solid'),
        'font_title': Font(bold=True, size=16),
        'font_date': Font(bold=True, size=16),
        'font_creneau': Font(bold=True, size=16),
        'font_groupe': Font(bold=True, size=16),
        'font_salle': Font(size=9),
        'font_label': Font(size=11),
        'font_bold': Font(bold=True, size=11),
        'font_total': Font(bold=True, size=16),
        'font_absence': Font(bold=True, size=16, color=FONT_ABSENCE),
        'font_absence_bold': Font(bold=True, size=16, color=FONT_ABSENCE),
        'font_taux_jour': Font(bold=True, size=20),
        'font_sidebar': Font(bold=True, size=22),
        'align_center': Alignment(horizontal='center', vertical='center'),
        'align_center_wrap': Alignment(horizontal='center', vertical='center', wrap_text=True),
    }


def export_excel(tableaux, annee):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    styles = _make_styles()

    if not tableaux:
        ws = wb.create_sheet('Vide')
        ws['A1'] = 'Aucun point journalier pour cette période.'
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    used = set()
    for tb in tableaux:
        base = _safe_filename(tb.get('grade') or tb['categorie'] or tb['date'])[:28]
        name = base
        i = 1
        while name in used:
            name = f"{base[:24]}_{i}"
            i += 1
        used.add(name)
        ws = wb.create_sheet(name)
        _write_sheet(ws, tb, styles)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_pdf(tableaux, annee):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    page_size = landscape(A4)
    margin = 0.6 * cm
    page_width = page_size[0] - 2 * margin

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=page_size,
        leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin,
    )
    styles = getSampleStyleSheet()
    st_t = ParagraphStyle('PJTitle', parent=styles['Normal'], fontSize=9, alignment=TA_CENTER)
    st_sub = ParagraphStyle('PJSub', parent=styles['Normal'], fontSize=8, alignment=TA_CENTER)

    elements = []
    for idx, tb in enumerate(tableaux or []):
        if idx:
            elements.append(PageBreak())
        elements.append(Paragraph(f"<b>{tb['titre_ligne1']}</b>", st_t))
        elements.append(Paragraph(
            f"DATE {tb['jour']} {tb['mois_libelle']} {tb['annee']} — {tb['organisme']}",
            st_sub,
        ))
        elements.append(Spacer(1, 0.15 * cm))

        for bloc_name, bloc in [('MATIN', tb['matin']), ('SOIR', tb['soir'])]:
            n_g = len(bloc.get('groupes') or [])
            suffix = ''
            if n_g > PDF_GROUPES_PAR_TABLE:
                suffix = f" ({n_g} groupes — {((n_g - 1) // PDF_GROUPES_PAR_TABLE) + 1} tableaux)"
            elements.append(Paragraph(f"<b>{bloc_name}: {bloc['horaire']}</b>{suffix}", st_sub))
            for ti, tbl in enumerate(_pdf_bloc_tables(
                bloc, page_width, cm, colors, Paragraph, ParagraphStyle, styles, TA_CENTER,
            )):
                if ti:
                    elements.append(Spacer(1, 0.1 * cm))
                elements.append(tbl)
            elements.append(Spacer(1, 0.2 * cm))

        elements.append(Paragraph(
            f"Taux présence du jour : {tb['taux_presence_jour']:.2%} — "
            f"Taux absence du jour : {tb['taux_absence_jour']:.2%}",
            st_sub,
        ))

    if not tableaux:
        elements.append(Paragraph('Aucun tableau.', st_t))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _word_add_bloc(doc, bloc_name, bloc, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT):
    """Ajoute un créneau Word en une ou plusieurs tables (découpage si large)."""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH as WAP

    n_g = len(bloc.get('groupes') or [])
    suffix = f" ({n_g} groupes)" if n_g > PDF_GROUPES_PAR_TABLE else ''
    p = doc.add_paragraph(f"{bloc_name}: {bloc['horaire']}{suffix}")
    p.runs[0].bold = True

    for chunk, chunk_total, part, n_parts in _iter_groupe_chunks(bloc.get('groupes') or []):
        if n_parts > 1:
            doc.add_paragraph(f"Suite des groupes ({part}/{n_parts})").alignment = WAP.CENTER
        cols = 2 + len(chunk) + 1
        table = doc.add_table(rows=7, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        hdr[0].text = ''
        hdr[1].text = 'GROUPES'
        for i, g in enumerate(chunk):
            hdr[2 + i].text = g['label']
            for run in hdr[2 + i].paragraphs[0].runs:
                run.font.size = Pt(7)
        hdr[-1].text = f'TOT.{part}' if n_parts > 1 else 'TOTAL'

        salles = table.rows[1].cells
        salles[0].text = ''
        salles[1].text = 'SALLES'
        for i, g in enumerate(chunk):
            salles[2 + i].text = g['salle'] or '—'

        total_row = chunk_total if n_parts > 1 else bloc['total']
        for ri, (lbl, key, pct) in enumerate([
            ('ÉFFECTIF', 'effectif', False),
            ('PRÉSENTS', 'presents', False),
            ('ABSENTS', 'absents', False),
            ('TAUX DE PRÉSENCE', 'taux_presence', True),
            ("TAUX D'ABSENCE", 'taux_absence', True),
        ], start=2):
            cells = table.rows[ri].cells
            cells[0].text = lbl
            cells[1].text = ''
            for i, g in enumerate(chunk):
                v = g[key]
                cells[2 + i].text = f"{v:.2%}" if pct else str(v)
            tv = total_row[key]
            cells[-1].text = f"{tv:.2%}" if pct else str(tv)


def export_word(tableaux, annee):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.section import WD_ORIENT

        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

        for idx, tb in enumerate(tableaux or []):
            if idx:
                doc.add_page_break()
            p = doc.add_paragraph(tb['titre_ligne1'])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].bold = True

            doc.add_paragraph(
                f"DATE {tb['jour']} {tb['mois_libelle']} {tb['annee']} — {tb['organisme']}"
            ).alignment = WD_ALIGN_PARAGRAPH.CENTER

            for bloc_name, bloc in [('MATIN', tb['matin']), ('SOIR', tb['soir'])]:
                _word_add_bloc(doc, bloc_name, bloc, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT)

            doc.add_paragraph(
                f"Taux présence du jour : {tb['taux_presence_jour']:.2%} | "
                f"Taux absence : {tb['taux_absence_jour']:.2%}"
            )

        if not tableaux:
            doc.add_paragraph('Aucun point journalier.')

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer, 'docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    except ImportError:
        parts = ['<html><head><meta charset="utf-8"></head><body>']
        for tb in tableaux or []:
            parts.append(f"<h2>{tb['titre_ligne1']}</h2>")
            parts.append(f"<p>DATE {tb['jour']} {tb['mois_libelle']} {tb['annee']}</p>")
            for name, bloc in [('MATIN', tb['matin']), ('SOIR', tb['soir'])]:
                parts.append(f"<h3>{name}: {bloc['horaire']}</h3><table border='1'>")
                parts.append('<tr><td></td><td>GROUPES</td>' +
                             ''.join(f"<td>{g['label']}</td>" for g in bloc['groupes']) +
                             '<td>TOTAL</td></tr>')
                for lbl, key, pct in [
                    ('ÉFFECTIF', 'effectif', False), ('PRÉSENTS', 'presents', False),
                    ('ABSENTS', 'absents', False), ('TAUX PRÉS.', 'taux_presence', True),
                ]:
                    parts.append(f"<tr><td>{lbl}</td><td></td>")
                    for g in bloc['groupes']:
                        v = g[key]
                        parts.append(f"<td>{v:.2%}</td>" if pct else f"<td>{v}</td>")
                    tv = bloc['total'][key]
                    parts.append(f"<td>{tv:.2%}</td>" if pct else f"<td>{tv}</td>")
                    parts.append('</tr>')
                parts.append('</table>')
        parts.append('</body></html>')
        buffer = io.BytesIO(''.join(parts).encode('utf-8'))
        buffer.seek(0)
        return buffer, 'doc', 'application/msword'


def build_export_response(fmt, annee, mois=None, categorie=None, formation_id=None,
                          secretariat_id=None, module_ids=None, jour=None):
    data = compute_point_journalier(
        annee=annee, mois=mois, categorie=categorie,
        formation_id=formation_id, secretariat_id=secretariat_id,
        module_ids=module_ids, jour=jour,
    )
    tableaux = _filter_tableaux(data, formation_id, categorie, mois, jour)

    if fmt == 'xlsx':
        buffer = export_excel(tableaux, annee)
        ext, ctype = 'xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif fmt == 'pdf':
        buffer = export_pdf(tableaux, annee)
        ext, ctype = 'pdf', 'application/pdf'
    elif fmt in ('docx', 'doc', 'word'):
        buffer, ext, ctype = export_word(tableaux, annee)
    else:
        raise ValueError(f'Format inconnu : {fmt}')

    parts = ['POINT_JOURNALIER', str(annee)]
    if categorie:
        parts.append(f'CAT{categorie.upper()}')
    if jour:
        parts.append(jour.replace('-', ''))
    elif mois:
        parts.append(f'M{int(mois):02d}')
    if tableaux and len(tableaux) == 1:
        tb0 = tableaux[0]
        parts = ['POINT_JOURNALIER', f"CAT{tb0['categorie']}", tb0['date'].replace('-', '')]
    fname = _safe_filename('_'.join(parts)) + f'.{ext}'

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=ctype)
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    return response
