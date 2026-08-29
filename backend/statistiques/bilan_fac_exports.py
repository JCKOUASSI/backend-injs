"""
Exports Bilan FAC — Excel / PDF / Word.
"""
import copy
import io
import json

from django.http import HttpResponse

from .bilans import compute_bilan_fac
from .bilans_exports import _fmt_nb, _safe_filename, _style_cell, _xlsx_styles


def _fmt_pct_fac(n):
    if n is None:
        return '—'
    v = float(n or 0)
    if 0 <= v <= 1:
        v *= 100
    return f'{v:.2f}'.replace('.', ',') + ' %'


def _fmt_vh(n):
    if n is None:
        return '—'
    v = float(n or 0)
    return str(int(v)) if v == int(v) else f'{v:.1f}'


def _parse_meta(raw):
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}


def _apply_fac_meta(data, meta):
    if not meta:
        return data
    data = copy.deepcopy(data)
    justifs = meta.get('justificatifs') or {}
    diffs = meta.get('difficultes') or {}
    for ligne in data.get('point_global', {}).get('lignes', []):
        grade = ligne.get('grade')
        if grade in justifs and str(justifs[grade]).strip():
            ligne['justificatifs'] = str(justifs[grade]).strip()
        if grade in diffs and str(diffs[grade]).strip():
            ligne['difficultes'] = str(diffs[grade]).strip()
    return data


def _write_point_global_xlsx(ws, data, styles):
    pg = data.get('point_global') or {}
    lignes = pg.get('lignes') or []
    totaux = pg.get('totaux') or {}

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=16)
    c = ws.cell(row=1, column=1, value=f"POINT GLOBAL — {data.get('formation', '')}")
    c.font = styles['font_title']
    c.alignment = styles['align']

    headers = [
        'CATÉGORIE / GRADE',
        'EFFECTIF SECRÉTARIAT',
        'EFFECTIF ENCADREURS',
        'NOMBRE DE GROUPES',
        'EFFECTIF AUDITEURS',
        'ABSENTS NOTOIRES',
        'GROUPES TERMINÉS',
        'JUSTIFICATIFS ABSENCES NOTOIRES',
        'TAUX PARTICIPATION',
        'TAUX ABSENTS NOTOIRES',
        'TOTAL VH',
        'VH ÉPUISÉ',
        "TAUX EXÉCUTION VH",
        'TAUX PRÉSENCE COURS',
        'TAUX ABSENCE COURS',
        'DIFFICULTÉS RENCONTRÉES',
    ]
    row = 3
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        fill = 'header_dark' if col == 5 else 'header'
        if col in (11, 12, 13):
            fill = 'gray'
        _style_cell(cell, styles, fill=fill, bold=True)
    row += 1

    for l in lignes:
        vals = [
            l.get('grade'),
            l.get('effectif_secretariat'),
            l.get('nb_encadrants'),
            l.get('nb_groupes'),
            l.get('effectif_auditeurs'),
            l.get('absents_notoires'),
            l.get('groupes_termines'),
            l.get('justificatifs') or '',
            _fmt_pct_fac(l.get('taux_participation')),
            _fmt_pct_fac(l.get('taux_absents_notoires')),
            _fmt_vh(l.get('vh_total')),
            _fmt_vh(l.get('vh_epuise')),
            _fmt_pct_fac(l.get('taux_exec_vh')),
            _fmt_pct_fac(l.get('taux_presence_cours')),
            _fmt_pct_fac(l.get('taux_absence_cours')),
            l.get('difficultes') or '',
        ]
        fills = ['td_data'] * 16
        fills[4] = 'td_blue'
        fills[10] = fills[11] = fills[12] = 'gray'
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            align = styles['align_left'] if col in (8, 16) else styles['align']
            _style_cell(cell, styles, fill=fills[col - 1], bold=col == 1, align=align)
        row += 1

    vals = [
        'TOTAL',
        totaux.get('effectif_secretariat'),
        totaux.get('nb_encadrants'),
        totaux.get('nb_groupes'),
        totaux.get('effectif_auditeurs'),
        totaux.get('absents_notoires'),
        totaux.get('groupes_termines'),
        '—', '—', '—',
        _fmt_vh(totaux.get('vh_total')),
        _fmt_vh(totaux.get('vh_epuise')),
        _fmt_pct_fac(totaux.get('taux_exec_vh')),
        '—', '—', '—',
    ]
    for col, val in enumerate(vals, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        _style_cell(cell, styles, fill='td_total', bold=True)
    return row + 2


def _write_vh_par_grade_xlsx(ws, data, styles):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    c = ws.cell(row=1, column=1, value=f"VH PAR GROUPE — {data.get('formation', '')}")
    c.font = styles['font_title']
    c.alignment = styles['align']
    row = 3

    for grade_block in data.get('vh_par_grade') or []:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f"GRADE {grade_block.get('grade', '')}")
        _style_cell(cell, styles, fill='header_dark', bold=True)
        row += 1

        groupes = grade_block.get('groupes') or []
        headers = ['INDICATEUR'] + [g.get('groupe', '') for g in groupes] + ['RÉCAP']
        for col, label in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=label)
            _style_cell(cell, styles, fill='header', bold=True)
        row += 1

        rows_def = [
            ('vh_prevu', 'VOLUME HORAIRE DU CYCLE', 'td_data'),
            ('vh_epuise', 'VOLUME HORAIRE ÉPUISÉ', 'td_data'),
            ('taux_execution', "TAUX D'EXÉCUTION (%)", 'green'),
            ('vh_restant', 'VOLUME HORAIRE RESTANT', 'td_data'),
            ('taux_restant', 'TAUX VH RESTANT (%)', 'td_data'),
        ]
        recap = grade_block.get('recap') or {}
        for key, label, fill in rows_def:
            cell = ws.cell(row=row, column=1, value=label)
            _style_cell(cell, styles, fill='gray', bold=True)
            for i, g in enumerate(groupes, start=2):
                val = g.get(key)
                if key.startswith('taux'):
                    val = _fmt_pct_fac(val / 100 if val and val > 1 else val)
                else:
                    val = _fmt_vh(val)
                cell = ws.cell(row=row, column=i, value=val)
                _style_cell(cell, styles, fill=fill)
            recap_val = recap.get(key)
            if key.startswith('taux'):
                recap_val = _fmt_pct_fac(recap_val / 100 if recap_val and recap_val > 1 else recap_val)
            else:
                recap_val = _fmt_vh(recap_val)
            cell = ws.cell(row=row, column=len(groupes) + 2, value=recap_val)
            _style_cell(cell, styles, fill='td_total', bold=True)
            row += 1
        row += 1
    return row


def _write_absents_xlsx(ws, data, styles):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    c = ws.cell(row=1, column=1, value='ÉTAT DES AUDITEURS ABSENTS NOTOIRES')
    c.font = styles['font_title']
    c.alignment = styles['align']
    row = 3
    headers = ['N°', 'MATRICULE', 'NOM', 'PRÉNOM', 'LIBELLÉ CONCOURS', 'CONTACTS', 'GROUPE', 'GRADE', 'OBSERVATIONS']
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        _style_cell(cell, styles, fill='header', bold=True)
    row += 1
    for a in data.get('absents_notoires') or []:
        vals = [
            a.get('numero'), a.get('matricule'), a.get('nom'), a.get('prenom'),
            a.get('libelle_concours'), a.get('contacts'), a.get('groupe'),
            a.get('grade'), a.get('observations'),
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            align = styles['align_left'] if col == 9 else styles['align']
            _style_cell(cell, styles, fill='td_data', align=align)
        row += 1
    if row == 4:
        ws.cell(row=4, column=1, value='Aucun absent notoire.')
    return row + 1


def _write_modules_xlsx(ws, data, styles):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    c = ws.cell(row=1, column=1, value="ÉTAT D'AVANCEMENT DES MODULES")
    c.font = styles['font_title']
    c.alignment = styles['align']
    row = 3
    headers = ['INTITULÉ', 'GRADE', 'GROUPE', 'STATUT', 'DÉBUT', 'FIN', 'VH PRÉVU', 'VH RÉALISÉ', 'VH RESTANT']
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        _style_cell(cell, styles, fill='header', bold=True)
    row += 1
    for m in data.get('modules_statuts') or []:
        vals = [
            m.get('intitule'), m.get('grade'), m.get('groupe'), m.get('statut'),
            m.get('date_debut') or '—', m.get('date_fin') or '—',
            _fmt_vh(m.get('vh_prevu')), _fmt_vh(m.get('vh_realise')), _fmt_vh(m.get('vh_restant')),
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            align = styles['align_left'] if col == 1 else styles['align']
            _style_cell(cell, styles, fill='td_data', align=align)
        row += 1
    return row + 1


def export_excel_fac(data):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    styles = _xlsx_styles()

    sheets = [
        ('Point_global', _write_point_global_xlsx),
        ('VH_par_groupe', _write_vh_par_grade_xlsx),
        ('Absents_notoires', _write_absents_xlsx),
        ('Modules', _write_modules_xlsx),
    ]
    for name, writer in sheets:
        ws = wb.create_sheet(name[:31])
        writer(ws, data, styles)
        for col in range(1, 18):
            ws.column_dimensions[get_column_letter(col)].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_pdf_fac(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    title_style = ParagraphStyle('Title', fontName='Helvetica-Bold', fontSize=10, alignment=1)
    elements = [Paragraph(f"<b>{data.get('titre', 'BILAN FAC')}</b>", title_style), Spacer(1, 12)]

    pg = data.get('point_global') or {}
    hdr = ['GRADE', 'AUDITEURS', 'ABS. NOT.', 'VH TOTAL', 'VH ÉPUISÉ', 'TAUX EXÉC.', 'DIFFICULTÉS']
    rows = [hdr]
    for l in pg.get('lignes') or []:
        rows.append([
            l.get('grade', ''),
            _fmt_nb(l.get('effectif_auditeurs')),
            _fmt_nb(l.get('absents_notoires')),
            _fmt_vh(l.get('vh_total')),
            _fmt_vh(l.get('vh_epuise')),
            _fmt_pct_fac(l.get('taux_exec_vh')),
            (l.get('difficultes') or '')[:80],
        ])
    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FCE7B4')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(Paragraph('<b>Point global</b>', title_style))
    elements.append(t)
    elements.append(Spacer(1, 16))

    abs_rows = [['N°', 'MATRICULE', 'NOM', 'GROUPE', 'GRADE', 'OBSERVATIONS']]
    for a in (data.get('absents_notoires') or [])[:50]:
        abs_rows.append([
            str(a.get('numero', '')),
            a.get('matricule', ''),
            f"{a.get('nom', '')} {a.get('prenom', '')}".strip(),
            a.get('groupe', ''),
            a.get('grade', ''),
            (a.get('observations') or '')[:60],
        ])
    if len(abs_rows) > 1:
        t2 = Table(abs_rows, repeatRows=1)
        t2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FCE7B4')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
        ]))
        elements.append(Paragraph('<b>Absents notoires</b>', title_style))
        elements.append(t2)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def export_word_fac(data):
    parts = [
        '<html><head><meta charset="utf-8"/></head><body>',
        f"<h2>{data.get('titre', 'BILAN FAC')}</h2>",
        '<h3>Point global</h3><table border="1" cellspacing="0" cellpadding="4">',
        '<tr><th>GRADE</th><th>AUDITEURS</th><th>ABS. NOT.</th><th>VH TOTAL</th>'
        '<th>VH ÉPUISÉ</th><th>TAUX EXÉC.</th><th>JUSTIFICATIFS</th><th>DIFFICULTÉS</th></tr>',
    ]
    for l in (data.get('point_global') or {}).get('lignes') or []:
        parts.append(
            f"<tr><td>{l.get('grade', '')}</td>"
            f"<td>{_fmt_nb(l.get('effectif_auditeurs'))}</td>"
            f"<td>{_fmt_nb(l.get('absents_notoires'))}</td>"
            f"<td>{_fmt_vh(l.get('vh_total'))}</td>"
            f"<td>{_fmt_vh(l.get('vh_epuise'))}</td>"
            f"<td>{_fmt_pct_fac(l.get('taux_exec_vh'))}</td>"
            f"<td>{(l.get('justificatifs') or '').replace('<', '&lt;')}</td>"
            f"<td>{(l.get('difficultes') or '').replace('<', '&lt;')}</td></tr>"
        )
    parts.append('</table>')

    parts.append('<h3>Absents notoires</h3><table border="1" cellspacing="0" cellpadding="4">')
    parts.append('<tr><th>N°</th><th>Matricule</th><th>Nom</th><th>Groupe</th><th>Grade</th><th>Observations</th></tr>')
    for a in data.get('absents_notoires') or []:
        parts.append(
            f"<tr><td>{a.get('numero', '')}</td><td>{a.get('matricule', '')}</td>"
            f"<td>{a.get('nom', '')} {a.get('prenom', '')}</td>"
            f"<td>{a.get('groupe', '')}</td><td>{a.get('grade', '')}</td>"
            f"<td>{(a.get('observations') or '').replace('<', '&lt;')}</td></tr>"
        )
    parts.append('</table></body></html>')
    buffer = io.BytesIO(''.join(parts).encode('utf-8'))
    buffer.seek(0)
    return buffer, 'doc', 'application/msword'


def _export_filename(data):
    formation = _safe_filename(data.get('formation') or 'FAC')
    annee = data.get('annee') or ''
    return _safe_filename(f'BILAN_FAC_{formation}_{annee}')


def build_bilan_fac_export_response(
    fmt,
    formation_id,
    annee=None,
    categorie=None,
    secretariat_id=None,
    calendrier=None,
    module_ids=None,
    grades_filter=None,
    groupes_filter=None,
    meta=None,
):
    data = compute_bilan_fac(
        formation_id=formation_id,
        annee=annee,
        categorie=categorie,
        secretariat_id=secretariat_id,
        calendrier=calendrier,
        module_ids=module_ids,
        grades_filter=grades_filter,
        groupes_filter=groupes_filter,
    )
    if not data:
        raise ValueError('Formation introuvable.')
    data = _apply_fac_meta(data, meta)

    if fmt == 'xlsx':
        buffer = export_excel_fac(data)
        ext, ctype = 'xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif fmt == 'pdf':
        buffer = export_pdf_fac(data)
        ext, ctype = 'pdf', 'application/pdf'
    elif fmt in ('docx', 'doc', 'word'):
        buffer, ext, ctype = export_word_fac(data)
    else:
        raise ValueError(f'Format inconnu : {fmt}')

    fname = _export_filename(data) + f'.{ext}'
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=ctype)
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    return response
