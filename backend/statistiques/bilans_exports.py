"""
Exports Bilans & Rapports — Excel / PDF / Word (modèles CPFAE).
"""
import io
import re

from django.http import HttpResponse

from .bilans import (
    compute_bilans,
    compute_bilan_effectifs_module,
    compute_bilan_effectifs_matiere,
    compute_bilan_effectifs_categorie,
    compute_bilan_periode_formation,
)

# Couleurs CPFAE (aRGB openpyxl)
C_HEADER = 'FFFCD5B4'
C_HEADER_DARK = 'FFF4B084'
C_GRAY = 'FFD9D9D9'
C_BLUE = 'FF99CCFF'
C_PINK = 'FFCCFF'
C_GREEN = 'FFA9D08E'
C_TD_BLUE = 'FFDDEBF7'
C_TD_TOTAL = 'FFFFD966'
C_TD_DATA = 'FFFFFBF5'
C_BLACK = 'FF000000'


def _safe_filename(text):
    return re.sub(r'[^\w\-]+', '_', text, flags=re.UNICODE).strip('_')[:80]


def _fmt_nb(n):
    return f'{int(n or 0):,}'.replace(',', '\u00a0')


def _fmt_pct(n):
    return f'{(n or 0):.2f}'.replace('.', ',') + '%'


def _pad2(n):
    return str(int(n or 0)).zfill(2)


def _inscrits_full(t):
    pct = t.get('pct_inscrits')
    if pct is None and t.get('inscrits_reference'):
        pct = round((t.get('inscrits_actifs', 0) or 0) / t['inscrits_reference'] * 100, 2)
    return f"{_fmt_nb(t.get('inscrits_actifs', 0))} / {_fmt_nb(t.get('inscrits_reference', 0))}\nSoit {_fmt_pct(pct or 0)}"


def _absents_full(t):
    if t.get('inscrits_reference', 0) > 0:
        pct = t.get('pct_absents')
        if pct is None:
            pct = round((t.get('absents_notoires', 0) or 0) / t['inscrits_reference'] * 100, 2)
        return f"{_fmt_nb(t.get('absents_notoires', 0))}\nSoit {_fmt_pct(pct)} de l'effectif"
    return _fmt_nb(t.get('absents_notoires', 0))


def collect_bilans_tableaux(
    dimension,
    annee,
    mois=None,
    categorie=None,
    module_id=None,
    module_ids=None,
    formation_id=None,
    secretariat_id=None,
    periode=None,
    calendrier=None,
    ref_module_id=None,
):
    """Construit la liste des tableaux détaillés à exporter."""
    data = compute_bilans(
        annee=annee, mois=mois, categorie=categorie, module_id=module_id,
        module_ids=module_ids,
        formation_id=formation_id, secretariat_id=secretariat_id,
        periode=periode, calendrier=calendrier, dimension=dimension,
        ref_module_id=ref_module_id,
    )
    bilans = data.get('bilans') or []

    if module_id:
        bilans = [b for b in bilans if b.get('module_id') == module_id]
    elif formation_id:
        bilans = [b for b in bilans if b.get('formation_id') == formation_id]
    elif categorie:
        bilans = [b for b in bilans if (b.get('categorie') or '').upper() == categorie.upper()]

    tableaux = []
    for b in bilans:
        t = None
        if b['dimension'] == 'module' and b.get('module_id'):
            t = compute_bilan_effectifs_module(
                b['module_id'], categorie=categorie or b.get('categorie'),
                annee=annee, mois=mois, calendrier=calendrier, periode=periode,
            )
        elif b['dimension'] == 'matiere' and b.get('formation_id'):
            t = compute_bilan_effectifs_matiere(
                b['formation_id'],
                ref_module_id=b.get('ref_module_id'),
                matiere_intitule=b.get('matiere_intitule'),
                categorie=categorie or b.get('categorie'),
                secretariat_id=secretariat_id,
                annee=annee, mois=mois, calendrier=calendrier, periode=periode,
                module_ids=module_ids,
            )
        elif b['dimension'] == 'formation' and b.get('formation_id'):
            t = compute_bilan_periode_formation(
                b['formation_id'], annee=annee, mois=mois, calendrier=calendrier,
                periode=periode, secretariat_id=secretariat_id,
                categorie_filter=categorie or b.get('categorie'),
                module_ids=module_ids,
            )
        elif b['dimension'] == 'categorie' and b.get('categorie') and b['categorie'] != '—':
            t = compute_bilan_effectifs_categorie(
                b['categorie'], formation_id=formation_id, secretariat_id=secretariat_id,
                annee=annee, mois=mois, calendrier=calendrier, periode=periode,
                module_ids=module_ids,
            )
        if t:
            tableaux.append(t)
    return tableaux


def _xlsx_styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000'),
    )
    return {
        'border': border,
        'align': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'align_left': Alignment(horizontal='left', vertical='top', wrap_text=True),
        'font_bold': Font(bold=True, size=10),
        'font_title': Font(bold=True, size=11, underline='single'),
        'font_hdr': Font(bold=True, size=8),
        'fills': {
            'header': PatternFill(start_color=C_HEADER, end_color=C_HEADER, fill_type='solid'),
            'header_dark': PatternFill(start_color=C_HEADER_DARK, end_color=C_HEADER_DARK, fill_type='solid'),
            'gray': PatternFill(start_color=C_GRAY, end_color=C_GRAY, fill_type='solid'),
            'blue': PatternFill(start_color=C_BLUE, end_color=C_BLUE, fill_type='solid'),
            'pink': PatternFill(start_color=C_PINK, end_color=C_PINK, fill_type='solid'),
            'green': PatternFill(start_color=C_GREEN, end_color=C_GREEN, fill_type='solid'),
            'td_blue': PatternFill(start_color=C_TD_BLUE, end_color=C_TD_BLUE, fill_type='solid'),
            'td_total': PatternFill(start_color=C_TD_TOTAL, end_color=C_TD_TOTAL, fill_type='solid'),
            'td_data': PatternFill(start_color=C_TD_DATA, end_color=C_TD_DATA, fill_type='solid'),
        },
    }


def _style_cell(cell, styles, fill='td_data', bold=False, align=None):
    cell.border = styles['border']
    cell.alignment = align or styles['align']
    cell.font = styles['font_bold'] if bold else styles['font_hdr']
    cell.fill = styles['fills'].get(fill, styles['fills']['td_data'])


def _write_effectifs_xlsx(ws, start_row, data, styles):
    """Tableau effectifs module / catégorie."""
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=5)
    c = ws.cell(row=start_row, column=1, value=data['titre'])
    c.font = styles['font_title']
    c.alignment = styles['align']
    start_row += 2

    hdr1 = [
        ('EFFECTIFS\nDES AUDITEURS', 'header', 1, 1),
        ('EFFECTIFS\nPRESENTS', 'header', 1, 1),
        ('EFFECTIFS PRESENTS PAR GENRE', 'gray', 1, 2),
        ('ABSENTS', 'green', 1, 1),
    ]
    col = 1
    for label, fill, rs, cs in hdr1:
        if cs > 1:
            ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row + rs - 1, end_column=col + cs - 1)
        else:
            ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row + rs - 1, end_column=col)
        cell = ws.cell(row=start_row, column=col, value=label)
        _style_cell(cell, styles, fill=fill, bold=True)
        col += cs

    ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=2)
    for i, (label, fill) in enumerate([('MASCULIN', 'blue'), ('FEMININ', 'pink')], start=3):
        cell = ws.cell(row=start_row + 1, column=i, value=label)
        _style_cell(cell, styles, fill=fill, bold=True)
    ws.merge_cells(start_row=start_row + 1, start_column=5, end_row=start_row + 1, end_column=5)

    row = start_row + 2
    cells = [
        (_fmt_nb(data['effectifs_auditeurs']), 'td_data'),
        (f"{_fmt_nb(data['effectifs_presents'])}\nsoit {_fmt_pct(data['pct_presents_total'])} de l'effectif total", 'td_data'),
        (f"{_fmt_nb(data['masculin'])}\nsoit {_fmt_pct(data['pct_masculin_presents'])} de l'effectif des présents", 'td_data'),
        (f"{_fmt_nb(data['feminin'])}\nsoit {_fmt_pct(data['pct_feminin_presents'])} de l'effectif des présents", 'td_data'),
        (f"{_fmt_nb(data['absents'])}\nsoit {_fmt_pct(data['pct_absents_total'])} de l'effectif total", 'td_data'),
    ]
    for i, (val, fill) in enumerate(cells, start=1):
        cell = ws.cell(row=row, column=i, value=val)
        _style_cell(cell, styles, fill=fill, bold=True)
    return row + 3


def _write_formation_xlsx(ws, start_row, data, styles):
    """Tableau bilan période formation."""
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=9)
    c = ws.cell(row=start_row, column=1, value=data['titre'])
    c.font = styles['font_title']
    c.alignment = styles['align']
    start_row += 2

    headers = [
        ('CATEGORIE / GRADE', 'header', 2),
        ('CATEGORIE / GRADE', 'header', 2),
        ('NBRE DE GRPES', 'header', 2),
        ('NBRE D\'ENCADRANTS', 'header', 2),
        ('EFFECTIF MEMBRE DE SECRETARIAT', 'header', 2),
        (f"EFFECTIF DES INSCRITS AU {data['date_inscrits']}", 'header_dark', 2),
        ('EFFECTIF DES AUDITEURS SUR LES LISTES DE CLASSES', 'header', 2),
        ('EFFECTIF DES AUDITEURS ABSENTS NOTOIRES', 'header', 2),
        ('JUSTIFICATIFS DES ABSENCES NOTOIRES', 'header', 2),
    ]
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
    for col, (label, fill, _) in enumerate(headers, start=1):
        if col == 2:
            continue
        if col == 1:
            cell = ws.cell(row=start_row, column=1, value='CATEGORIE / GRADE')
        else:
            cell = ws.cell(row=start_row, column=col, value=label)
        _style_cell(cell, styles, fill=fill, bold=True)
        if col >= 3:
            ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row + 1, end_column=col)

    row = start_row + 2
    justif_text = '\n'.join(f'• {j}' for j in data.get('justificatifs', []))
    justif_row_start = row
    justif_placed = False

    for ligne in data.get('lignes', []):
        if ligne.get('multi_grade') and ligne.get('sous_lignes'):
            n_sub = len(ligne['sous_lignes'])
            cat_start = row
            for idx, sl in enumerate(ligne['sous_lignes']):
                vals = [
                    ligne['categorie'] if idx == 0 else None,
                    sl['grade'],
                    _pad2(sl['nb_groupes']),
                    ligne['totaux']['nb_encadrants'] if idx == 0 else None,
                    ligne['totaux']['effectif_secretariat'] if idx == 0 else None,
                    _fmt_nb(sl['inscrits_actifs']),
                    _fmt_nb(sl['auditeurs_listes']),
                    _pad2(sl['absents_notoires']),
                    justif_text if not justif_placed else None,
                ]
                fills = ['td_data', 'td_data', 'td_data', 'td_data', 'td_data', 'td_data', 'td_blue', 'td_data', 'td_data']
                for col, val in enumerate(vals, start=1):
                    if val is None:
                        continue
                    cell = ws.cell(row=row, column=col, value=val)
                    fill = fills[col - 1]
                    align = styles['align_left'] if col == 9 else styles['align']
                    _style_cell(cell, styles, fill=fill, bold=True, align=align)
                if not justif_placed:
                    justif_placed = True
                row += 1
            ws.merge_cells(start_row=cat_start, start_column=1, end_row=row - 1, end_column=1)
            ws.merge_cells(start_row=cat_start, start_column=4, end_row=row - 1, end_column=4)
            ws.merge_cells(start_row=cat_start, start_column=5, end_row=row - 1, end_column=5)

            tot = ligne['totaux']
            for col, val in enumerate(['Σ', _pad2(tot['nb_groupes']), _inscrits_full(tot), _fmt_nb(tot['auditeurs_listes']), _absents_full(tot)], start=2):
                cell = ws.cell(row=row, column=col, value=val)
                _style_cell(cell, styles, fill='td_blue' if col == 5 else 'td_data', bold=True)
            row += 1
        else:
            tot = ligne['totaux']
            vals = [
                ligne['categorie'], '',
                _pad2(tot['nb_groupes']),
                _pad2(tot['nb_encadrants']),
                tot['effectif_secretariat'],
                _inscrits_full(tot),
                _fmt_nb(tot['auditeurs_listes']),
                _absents_full(tot),
                justif_text if not justif_placed else None,
            ]
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            fills = ['td_data', 'td_data', 'td_data', 'td_data', 'td_data', 'td_data', 'td_blue', 'td_data', 'td_data']
            for col, val in enumerate(vals, start=1):
                if val is None or val == '':
                    continue
                cell = ws.cell(row=row, column=col, value=val)
                align = styles['align_left'] if col == 9 else styles['align']
                _style_cell(cell, styles, fill=fills[col - 1], bold=True, align=align)
            if not justif_placed:
                justif_placed = True
            row += 1

    if justif_placed:
        ws.merge_cells(start_row=justif_row_start, start_column=9, end_row=row - 1, end_column=9)

    total = data.get('total', {})
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=1, value='TOTAL')
    _style_cell(cell, styles, fill='td_total', bold=True)
    for col, val in enumerate(
        [_pad2(total.get('nb_groupes', 0)), _pad2(total.get('nb_encadrants', 0)),
         total.get('effectif_secretariat', 0), _inscrits_full(total),
         _fmt_nb(total.get('auditeurs_listes', 0)), _fmt_nb(total.get('absents_notoires', 0))],
        start=3,
    ):
        cell = ws.cell(row=row, column=col, value=val)
        fill = 'td_blue' if col == 5 else 'td_total'
        _style_cell(cell, styles, fill=fill, bold=True)
    return row + 3


def export_excel(tableaux):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    styles = _xlsx_styles()

    if not tableaux:
        ws = wb.create_sheet('Vide')
        ws['A1'] = 'Aucun bilan pour cette sélection.'
    else:
        for i, tb in enumerate(tableaux):
            name = _safe_filename(tb.get('titre', f'Bilan_{i+1}'))[:28] or f'Bilan_{i+1}'
            base, n = name, 1
            while name in wb.sheetnames:
                name = f'{base[:24]}_{n}'
                n += 1
            ws = wb.create_sheet(name)
            if tb['type'] in ('effectifs_module', 'effectifs_categorie'):
                _write_effectifs_xlsx(ws, 1, tb, styles)
            else:
                _write_formation_xlsx(ws, 1, tb, styles)
            for col in range(1, 10):
                ws.column_dimensions[get_column_letter(col)].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _pdf_effectifs_table(data, title_style, Table, TableStyle, colors, Paragraph):
    rows = [
        ['EFFECTIFS DES AUDITEURS', 'EFFECTIFS PRESENTS', 'MASCULIN', 'FEMININ', 'ABSENTS'],
        [
            _fmt_nb(data['effectifs_auditeurs']),
            f"{_fmt_nb(data['effectifs_presents'])}\nsoit {_fmt_pct(data['pct_presents_total'])} de l'effectif total",
            f"{_fmt_nb(data['masculin'])}\nsoit {_fmt_pct(data['pct_masculin_presents'])} de l'effectif des présents",
            f"{_fmt_nb(data['feminin'])}\nsoit {_fmt_pct(data['pct_feminin_presents'])} de l'effectif des présents",
            f"{_fmt_nb(data['absents'])}\nsoit {_fmt_pct(data['pct_absents_total'])} de l'effectif total",
        ],
    ]
    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#FCD5B4')),
        ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#99CCFF')),
        ('BACKGROUND', (3, 0), (3, 0), colors.HexColor('#FFCCFF')),
        ('BACKGROUND', (4, 0), (4, 0), colors.HexColor('#A9D08E')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ]))
    return [Paragraph(f"<b>{data['titre']}</b>", title_style), t]


def _pdf_formation_table(data, title_style, Table, TableStyle, colors, Paragraph):
    hdr = [
        'CAT/GRADE', 'GRADE', 'GRPES', 'ENC.', 'SEC.',
        f"INSCRITS {data['date_inscrits']}", 'LISTES', 'ABS.', 'JUSTIFICATIFS',
    ]
    rows = [hdr]
    justifs = '\n'.join(f'• {j}' for j in data.get('justificatifs', []))
    first = True
    for ligne in data.get('lignes', []):
        if ligne.get('multi_grade') and ligne.get('sous_lignes'):
            for sl in ligne['sous_lignes']:
                rows.append([
                    ligne['categorie'], sl['grade'], _pad2(sl['nb_groupes']),
                    _pad2(ligne['totaux']['nb_encadrants']),
                    str(ligne['totaux']['effectif_secretariat']),
                    _fmt_nb(sl['inscrits_actifs']), _fmt_nb(sl['auditeurs_listes']),
                    _pad2(sl['absents_notoires']),
                    justifs if first else '',
                ])
                first = False
            tot = ligne['totaux']
            rows.append([
                '', 'Σ', _pad2(tot['nb_groupes']), '', '',
                _inscrits_full(tot).replace('\n', ' '), _fmt_nb(tot['auditeurs_listes']),
                _absents_full(tot).replace('\n', ' '), '',
            ])
        else:
            tot = ligne['totaux']
            rows.append([
                ligne['categorie'], '—', _pad2(tot['nb_groupes']),
                _pad2(tot['nb_encadrants']), str(tot['effectif_secretariat']),
                _inscrits_full(tot).replace('\n', ' '), _fmt_nb(tot['auditeurs_listes']),
                _absents_full(tot).replace('\n', ' '), justifs if first else '',
            ])
            first = False

    total = data.get('total', {})
    rows.append([
        'TOTAL', '', _pad2(total.get('nb_groupes', 0)), _pad2(total.get('nb_encadrants', 0)),
        str(total.get('effectif_secretariat', 0)),
        _inscrits_full(total).replace('\n', ' '), _fmt_nb(total.get('auditeurs_listes', 0)),
        _fmt_nb(total.get('absents_notoires', 0)), '',
    ])

    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FCD5B4')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FFD966')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 0), (-1, -1), 6),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ]))
    return [Paragraph(f"<b>{data['titre']}</b>", title_style), t]


def export_pdf(tableaux):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('BilanTitle', parent=styles['Heading2'], alignment=TA_CENTER, spaceAfter=8)

    story = []
    if not tableaux:
        story.append(Paragraph('Aucun bilan pour cette sélection.', styles['Normal']))
    else:
        from reportlab.platypus import Table, TableStyle
        for i, tb in enumerate(tableaux):
            if i:
                story.append(PageBreak())
            if tb['type'] in ('effectifs_module', 'effectifs_categorie'):
                story.extend(_pdf_effectifs_table(tb, title_style, Table, TableStyle, colors, Paragraph))
            else:
                story.extend(_pdf_formation_table(tb, title_style, Table, TableStyle, colors, Paragraph))
            story.append(Spacer(1, 12))

    doc.build(story)
    buffer.seek(0)
    return buffer


def export_word(tableaux):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt, Cm
    except ImportError:
        return _export_word_html(tableaux)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width

    if not tableaux:
        doc.add_paragraph('Aucun bilan pour cette sélection.')
    else:
        for i, tb in enumerate(tableaux):
            if i:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(tb.get('titre', 'BILAN'))
            run.bold = True
            run.font.size = Pt(11)
            run.underline = True

            if tb['type'] in ('effectifs_module', 'effectifs_categorie'):
                table = doc.add_table(rows=2, cols=5)
                table.style = 'Table Grid'
                hdrs = ['EFFECTIFS DES AUDITEURS', 'EFFECTIFS PRESENTS', 'MASCULIN', 'FEMININ', 'ABSENTS']
                for j, h in enumerate(hdrs):
                    table.rows[0].cells[j].text = h
                vals = [
                    _fmt_nb(tb['effectifs_auditeurs']),
                    f"{_fmt_nb(tb['effectifs_presents'])} — soit {_fmt_pct(tb['pct_presents_total'])} de l'effectif total",
                    f"{_fmt_nb(tb['masculin'])} — soit {_fmt_pct(tb['pct_masculin_presents'])} des présents",
                    f"{_fmt_nb(tb['feminin'])} — soit {_fmt_pct(tb['pct_feminin_presents'])} des présents",
                    f"{_fmt_nb(tb['absents'])} — soit {_fmt_pct(tb['pct_absents_total'])} de l'effectif total",
                ]
                for j, v in enumerate(vals):
                    table.rows[1].cells[j].text = v
            else:
                _add_formation_word_table(doc, tb)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer, 'docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _add_formation_word_table(doc, data):
    rows_data = [
        ['CAT/GRADE', 'GRADE', 'GRPES', 'ENC.', 'SEC.', f"INSCRITS {data['date_inscrits']}", 'LISTES', 'ABS.', 'JUSTIFICATIFS'],
    ]
    justifs = '\n'.join(f'• {j}' for j in data.get('justificatifs', []))
    first = True
    for ligne in data.get('lignes', []):
        if ligne.get('multi_grade') and ligne.get('sous_lignes'):
            for sl in ligne['sous_lignes']:
                rows_data.append([
                    ligne['categorie'], sl['grade'], _pad2(sl['nb_groupes']),
                    _pad2(ligne['totaux']['nb_encadrants']),
                    str(ligne['totaux']['effectif_secretariat']),
                    _fmt_nb(sl['inscrits_actifs']), _fmt_nb(sl['auditeurs_listes']),
                    _pad2(sl['absents_notoires']), justifs if first else '',
                ])
                first = False
            tot = ligne['totaux']
            rows_data.append(['', 'Σ', _pad2(tot['nb_groupes']), '', '', _inscrits_full(tot), _fmt_nb(tot['auditeurs_listes']), _absents_full(tot), ''])
        else:
            tot = ligne['totaux']
            rows_data.append([
                ligne['categorie'], '—', _pad2(tot['nb_groupes']),
                _pad2(tot['nb_encadrants']), str(tot['effectif_secretariat']),
                _inscrits_full(tot), _fmt_nb(tot['auditeurs_listes']), _absents_full(tot),
                justifs if first else '',
            ])
            first = False
    total = data.get('total', {})
    rows_data.append([
        'TOTAL', '', _pad2(total.get('nb_groupes', 0)), _pad2(total.get('nb_encadrants', 0)),
        str(total.get('effectif_secretariat', 0)), _inscrits_full(total),
        _fmt_nb(total.get('auditeurs_listes', 0)), _fmt_nb(total.get('absents_notoires', 0)), '',
    ])

    table = doc.add_table(rows=len(rows_data), cols=9)
    table.style = 'Table Grid'
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = str(val)


def _export_word_html(tableaux):
    parts = ['<html><head><meta charset="utf-8"></head><body>']
    if not tableaux:
        parts.append('<p>Aucun bilan pour cette sélection.</p>')
    else:
        for tb in tableaux:
            parts.append(f"<h3 style=\"text-align:center\">{tb.get('titre', 'BILAN')}</h3>")
            if tb['type'] in ('effectifs_module', 'effectifs_categorie'):
                parts.append('<table border="1" cellpadding="4"><tr>')
                for h in ['Auditeurs', 'Présents', 'Masculin', 'Féminin', 'Absents']:
                    parts.append(f'<th>{h}</th>')
                parts.append('</tr><tr>')
                for k in ['effectifs_auditeurs', 'effectifs_presents', 'masculin', 'feminin', 'absents']:
                    parts.append(f'<td>{_fmt_nb(tb[k])}</td>')
                parts.append('</tr></table>')
            else:
                parts.append('<table border="1" cellpadding="3"><tr>')
                for h in ['Cat', 'Grade', 'Grpes', 'Enc.', 'Sec.', 'Inscrits', 'Listes', 'Abs.']:
                    parts.append(f'<th>{h}</th>')
                parts.append('</tr>')
                for ligne in tb.get('lignes', []):
                    if ligne.get('multi_grade'):
                        for sl in ligne['sous_lignes']:
                            parts.append(f"<tr><td>{ligne['categorie']}</td><td>{sl['grade']}</td>"
                                         f"<td>{_pad2(sl['nb_groupes'])}</td><td></td><td></td>"
                                         f"<td>{_fmt_nb(sl['inscrits_actifs'])}</td>"
                                         f"<td>{_fmt_nb(sl['auditeurs_listes'])}</td>"
                                         f"<td>{_pad2(sl['absents_notoires'])}</td></tr>")
                    else:
                        t = ligne['totaux']
                        parts.append(f"<tr><td>{ligne['categorie']}</td><td>—</td>"
                                     f"<td>{_pad2(t['nb_groupes'])}</td><td>{_pad2(t['nb_encadrants'])}</td>"
                                     f"<td>{t['effectif_secretariat']}</td>"
                                     f"<td>{_inscrits_full(t)}</td>"
                                     f"<td>{_fmt_nb(t['auditeurs_listes'])}</td>"
                                     f"<td>{_absents_full(t)}</td></tr>")
                parts.append('</table>')
            parts.append('<br/>')
    parts.append('</body></html>')
    html = ''.join(parts)
    buffer = io.BytesIO(html.encode('utf-8'))
    return buffer, 'doc', 'application/msword'


def _export_filename(tableaux, dimension, annee):
    if len(tableaux) == 1:
        return _safe_filename(tableaux[0].get('titre', f'BILAN_{annee}'))
    return _safe_filename(f'BILANS_{dimension.upper()}_{annee}')


def build_bilans_export_response(
    fmt,
    dimension,
    annee,
    mois=None,
    categorie=None,
    module_id=None,
    module_ids=None,
    formation_id=None,
    secretariat_id=None,
    periode=None,
    calendrier=None,
    ref_module_id=None,
):
    tableaux = collect_bilans_tableaux(
        dimension=dimension, annee=annee, mois=mois, categorie=categorie,
        module_id=module_id, module_ids=module_ids,
        formation_id=formation_id, secretariat_id=secretariat_id,
        periode=periode, calendrier=calendrier, ref_module_id=ref_module_id,
    )

    if fmt == 'xlsx':
        buffer = export_excel(tableaux)
        ext, ctype = 'xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif fmt == 'pdf':
        buffer = export_pdf(tableaux)
        ext, ctype = 'pdf', 'application/pdf'
    elif fmt in ('docx', 'doc', 'word'):
        result = export_word(tableaux)
        if isinstance(result, tuple):
            buffer, ext, ctype = result
        else:
            buffer, ext, ctype = result, 'docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        raise ValueError(f'Format inconnu : {fmt}')

    fname = _export_filename(tableaux, dimension, annee) + f'.{ext}'
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=ctype)
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    return response
