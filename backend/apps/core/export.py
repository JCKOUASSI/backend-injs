"""Export utilities: PDF, Word, Excel."""
import io
from datetime import datetime
from typing import Any

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


INJS_HEADER = "INSTITUT NATIONAL DE LA JEUNESSE ET DES SPORTS - Côte d'Ivoire"


def export_to_excel(headers: list[str], rows: list[list[Any]], filename: str) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Export INJS'
    ws.append([INJS_HEADER])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws['A1'].font = Font(bold=True, size=12)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def export_to_pdf(title: str, headers: list[str], rows: list[list[Any]], filename: str) -> HttpResponse:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('INJSTitle', parent=styles['Heading1'], textColor=colors.HexColor('#0D47A1'))
    elements = [
        Paragraph(INJS_HEADER, title_style),
        Spacer(1, 12),
        Paragraph(title, styles['Heading2']),
        Spacer(1, 12),
    ]
    data = [headers] + [[str(c) for c in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D47A1')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E3F2FD')]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
        styles['Normal'],
    ))
    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
    return response


def export_to_word(title: str, headers: list[str], rows: list[list[Any]], filename: str) -> HttpResponse:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    header = document.add_heading(INJS_HEADER, level=1)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading(title, level=2)
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, val in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = str(val)
    document.add_paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.docx"'
    return response
