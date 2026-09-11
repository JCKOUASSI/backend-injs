"""Export PDF / Excel de l'analyse d'évaluation par groupe."""
import io

from django.utils import timezone

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OPENPYXL_AVAILABLE = True
except Exception:
    OPENPYXL_AVAILABLE = False

from .access import _normalize_groupe_value
from .resultats_utils import build_resultats_payload


def _safe_sheet_name(name, used):
    base = (name or 'Groupe')[:31]
    candidate = base
    n = 1
    while candidate in used:
        suffix = f'_{n}'
        candidate = f'{base[:31 - len(suffix)]}{suffix}'
        n += 1
    used.add(candidate)
    return candidate


def _analyse_payload(questionnaire, groupe_norm):
    """Données d'analyse pour un groupe (réutilise resultats + tendance)."""
    from collections import defaultdict
    from django.db.models import Avg, Count

    payload = build_resultats_payload(questionnaire, groupe_norm)
    block = payload['par_groupe'].get(groupe_norm, {})
    questions = block.get('questions') or []

    soumissions = [
        s for s in questionnaire.reponses.select_related('participant')
        if _normalize_groupe_value(s.participant.groupe) == groupe_norm
    ]
    tendance = defaultdict(int)
    for s in soumissions:
        if s.soumis_le:
            tendance[s.soumis_le.date().isoformat()] += 1

    notes_globales = [
        q['moyenne'] for q in questions
        if q.get('type_question') == 'NOTE' and q.get('moyenne') is not None
    ]
    score_global = (
        round(sum(notes_globales) / len(notes_globales), 2)
        if notes_globales else None
    )

    distribution_notes = defaultdict(int)
    for q in questions:
        if q.get('type_question') == 'NOTE':
            for k, v in (q.get('distribution') or {}).items():
                distribution_notes[k] += v

    return {
        'groupe': groupe_norm,
        'titres': questionnaire.titres or [],
        'nb_soumissions': block.get('nb_soumissions', 0),
        'score_global': score_global,
        'questions': questions,
        'tendance': [{'date': d, 'nb': n} for d, n in sorted(tendance.items())],
        'distribution_notes': dict(distribution_notes),
        'genere_le': timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M'),
    }


def _question_rows(questions):
    rows = []
    for idx, q in enumerate(questions, start=1):
        if q.get('type_question') == 'NOTE':
            dist = q.get('distribution') or {}
            dist_txt = ' · '.join(
                f'{n}★:{dist.get(str(n), 0)}' for n in range(1, 6) if dist.get(str(n), 0)
            )
            rows.append([
                f'Q{idx}',
                q.get('intitule', ''),
                'Note /5',
                str(q.get('moyenne') or '—'),
                str(q.get('total_reponses') or 0),
                dist_txt or '—',
            ])
        elif q.get('type_question') in ('CHOIX_UN', 'CHOIX_MUL'):
            for c in q.get('choix_stats') or []:
                rows.append([
                    f'Q{idx}',
                    q.get('intitule', ''),
                    'Choix',
                    c.get('libelle', ''),
                    str(c.get('nb_reponses') or 0),
                    '',
                ])
        elif q.get('type_question') == 'TEXTE':
            rows.append([
                f'Q{idx}',
                q.get('intitule', ''),
                'Texte libre',
                '—',
                str(q.get('nb_reponses_texte') or 0),
                '',
            ])
    return rows


def export_questionnaire_groupe_pdf(questionnaire, groupe_norm):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('reportlab non disponible')

    data = _analyse_payload(questionnaire, groupe_norm)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title2', parent=styles['Title'], fontSize=14, spaceAfter=8)
    h_style = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=11, spaceAfter=6)
    normal = styles['Normal']
    elems = []

    titres = ' · '.join(data['titres']) or 'Questionnaire'
    elems.append(Paragraph(f'Analyse d&apos;évaluation — {titres}', title_style))
    elems.append(Paragraph(f'<b>Groupe :</b> {data["groupe"]}', normal))
    elems.append(Paragraph(
        f'Soumissions : {data["nb_soumissions"]} · '
        f'Score moyen : {data["score_global"] or "—"}/5 · '
        f'Généré le {data["genere_le"]}',
        normal,
    ))
    elems.append(Spacer(1, 12))

    if data['tendance']:
        elems.append(Paragraph('Tendance des réponses (par jour)', h_style))
        tend_rows = [['Date', 'Nb réponses']] + [
            [t['date'], str(t['nb'])] for t in data['tendance']
        ]
        tbl = Table(tend_rows, colWidths=[5 * cm, 3 * cm])
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ]))
        elems.append(tbl)
        elems.append(Spacer(1, 12))

    if data['distribution_notes']:
        elems.append(Paragraph('Répartition des notes (1 à 5)', h_style))
        dist_rows = [['Note', 'Effectif']]
        for n in range(1, 6):
            nb = data['distribution_notes'].get(str(n), 0)
            if nb:
                dist_rows.append([str(n), str(nb)])
        tbl = Table(dist_rows, colWidths=[3 * cm, 3 * cm])
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        elems.append(tbl)
        elems.append(Spacer(1, 12))

    elems.append(Paragraph('Détail par question', h_style))
    q_rows = [['#', 'Question', 'Type', 'Résultat', 'Nb rép.', 'Distribution notes']]
    q_rows += _question_rows(data['questions'])
    tbl = Table(q_rows, colWidths=[1 * cm, 7 * cm, 2 * cm, 3 * cm, 1.5 * cm, 3 * cm])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
    ]))
    elems.append(tbl)

    doc.build(elems)
    buffer.seek(0)
    return buffer


def _write_groupe_sheet(ws, questionnaire, groupe_norm):
    data = _analyse_payload(questionnaire, groupe_norm)
    header_fill = PatternFill('solid', fgColor='E65100')
    header_font = Font(bold=True, color='FFFFFF')

    ws['A1'] = f"Analyse — {' · '.join(data['titres']) or 'Questionnaire'}"
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = f"Groupe : {data['groupe']}"
    ws['A3'] = (
        f"Soumissions : {data['nb_soumissions']} · "
        f"Score moyen : {data['score_global'] or '—'}/5 · "
        f"Généré le {data['genere_le']}"
    )

    row = 5
    ws.cell(row=row, column=1, value='Tendance (date · nb réponses)').font = Font(bold=True)
    row += 1
    if data['tendance']:
        ws.cell(row=row, column=1, value='Date')
        ws.cell(row=row, column=2, value='Nb réponses')
        for c in (1, 2):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
        row += 1
        for t in data['tendance']:
            ws.cell(row=row, column=1, value=t['date'])
            ws.cell(row=row, column=2, value=t['nb'])
            row += 1
    else:
        ws.cell(row=row, column=1, value='Aucune réponse')
        row += 1

    row += 1
    ws.cell(row=row, column=1, value='Répartition notes 1-5').font = Font(bold=True)
    row += 1
    for n in range(1, 6):
        nb = data['distribution_notes'].get(str(n), 0)
        if nb:
            ws.cell(row=row, column=1, value=f'Note {n}')
            ws.cell(row=row, column=2, value=nb)
            row += 1

    row += 1
    headers = ['#', 'Question', 'Type', 'Résultat', 'Nb rép.', 'Distribution']
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    row += 1
    for qrow in _question_rows(data['questions']):
        for col, val in enumerate(qrow, start=1):
            ws.cell(row=row, column=col, value=val)
        row += 1

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 48
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 24
    return data


def export_questionnaire_groupe_excel(questionnaire, groupe_norm):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('openpyxl non disponible')

    wb = Workbook()
    ws = wb.active
    used = set()
    ws.title = _safe_sheet_name(groupe_norm, used)
    _write_groupe_sheet(ws, questionnaire, groupe_norm)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def export_questionnaire_tous_groupes_excel(questionnaire):
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError('openpyxl non disponible')

    payload = build_resultats_payload(questionnaire)
    groupes = list(payload.get('par_groupe') or {})
    if not groupes:
        raise ValueError('Aucun groupe à exporter')

    wb = Workbook()
    used = set()
    first = True
    for groupe in groupes:
        if first:
            ws = wb.active
            ws.title = _safe_sheet_name(groupe, used)
            first = False
        else:
            ws = wb.create_sheet(title=_safe_sheet_name(groupe, used))
        _write_groupe_sheet(ws, questionnaire, groupe)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
