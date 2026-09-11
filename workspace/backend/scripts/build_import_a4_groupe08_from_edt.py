#!/usr/bin/env python3
"""
Lit l'EDT papier-style (Emploi du Temps A4 GROUPE 8MAJ.xlsx) et génère
import_formations / import_seances compatibles import_excel.

Usage :
  python backend/scripts/build_import_a4_groupe08_from_edt.py \\
    "/chemin/vers/Emploi du Temps A4 GROUPE 8MAJ.xlsx"

Sortie (backend/) — A4 groupe 8, deux jeux de noms (Formations avec Site, Bâtiment, Salle) :
  import_formations_A4_GROUPE_08_MAJ.xlsx  (+ copie import_formations_A4_GROUPE_08.xlsx)
  import_seances_A4_GROUPE_08_MAJ.xlsx     (+ copie import_seances_A4_GROUPE_08.xlsx)
  import_A4_GROUPE_08_MAJ_complet.xlsx     (+ copie import_A4_GROUPE_08_complet.xlsx)
"""
from __future__ import annotations

import re
import shutil
import sys
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook

# Répertoire projet (…/backend)
BACKEND = Path(__file__).resolve().parents[1]

MOIS_FR = {
    "janvier": 1,
    "février": 2,
    "fevrier": 2,
    "mars": 3,
    "avril": 4,
    "mai": 5,
    "juin": 6,
    "juillet": 7,
    "août": 8,
    "aout": 8,
    "septembre": 9,
    "octobre": 10,
    "novembre": 11,
    "décembre": 12,
    "decembre": 12,
}

FORMATION_CYCLE = "FORMATION EN ADMINISTRATION DE BASE"
GRADE = "A4"
# Aligné sur import_formations_reel (sans zéro devant) pour update_or_create des modules.
GROUPE = "GROUPE 8"
CATEGORIE = "FAB A"
SITE = "CPFAE-AGC"
VAGUE = "SESSION 2026"


def _norm(s: str) -> str:
    return (
        s.lower()
        .replace("é", "e")
        .replace("è", "e")
        .replace("ê", "e")
        .replace("à", "a")
        .replace("ô", "o")
        .replace("’", "'")
        .strip()
    )


def parse_french_date(s: str) -> date | None:
    if not s:
        return None
    s = _norm(s)
    # retirer jour de la semaine : "samedi 25 avril 2026"
    tokens = s.split()
    if len(tokens) < 4:
        return None
    try:
        jour = int(tokens[-3])
        mois = MOIS_FR.get(tokens[-2], 0)
        annee = int(tokens[-1])
        if not mois:
            return None
        return date(annee, mois, jour)
    except (ValueError, IndexError):
        return None


def parse_horaire(s: str) -> tuple[time | None, time | None]:
    if not s:
        return None, None
    s = s.upper().replace(" ", "")
    m = re.match(r"(\d{1,2})H-(\d{1,2})H", s)
    if not m:
        return None, None
    h1, h2 = int(m.group(1)), int(m.group(2))
    return time(h1, 0), time(h2, 0)


def title_case_module(name: str) -> str:
    """Aligné sur import_formations_reel (Title Case par mot)."""
    return " ".join(w.capitalize() for w in name.split())


def raw_to_canonical_module(raw: str) -> str:
    r = raw.upper().strip()
    if r.startswith("SIGFAE"):
        return "Sigfae Et Outils Collaboratifs Et Teletravail"
    mapping = {
        "CULTURE CIVIQUE": "Culture Civique",
        "PROTOCOLE ET SAVOIR-VIVRE": "Protocole Et Savoir-Vivre",
        "MANAGEMENT DES ADMINISTRATIONS PUBLIQUES": "Management Des Administrations Publiques",
        "GESTION DU BUDGET FAMILIAL": "Gestion Du Budget Familial",
        "DROIT ADMINISTRATIF": "Droit Administratif",
        "DEONTOLOGIE": "Deontologie",
        "ETHIQUE PUBLIQUE ET LUTTE CONTRE LA CORRUPTION": "Ethique Publique Et Lutte Contre La Corruption",
        "REDACTION ADMINISTRATIVE": "Redaction Administrative",
        "FINANCES PUBLIQUES": "Finances Publiques",
    }
    if r in mapping:
        return mapping[r]
    return title_case_module(raw)


def extract_rows(ws) -> list[tuple[int, list[str]]]:
    out = []
    for row in range(1, ws.max_row + 1):
        vals = []
        for col in range(5, 13):
            v = ws.cell(row=row, column=col).value
            vals.append(str(v).strip() if v is not None else "")
        if any(vals):
            out.append((row, vals))
    return out


def parse_schedule(rows: list[tuple[int, list[str]]]):
    """
    Retourne :
      sessions : liste (module_canon, date, heure_debut, heure_fin, intitule_creneau)
      volumes : module -> volume h (float) depuis ligne '16 H' en col6 après entête module
      module_batiment, module_salle : premiers libellés non vides (colonnes EDT bâtiment / salle)
    """
    sessions: list[tuple[str, date, time, time, str]] = []
    module_volumes: dict[str, float] = {}
    module_batiment: dict[str, str] = {}
    module_salle: dict[str, str] = {}
    current_module: str | None = None

    i = 0
    while i < len(rows):
        _, cells = rows[i]
        c5, c6, c7, c8, c9, c10 = cells[0], cells[1], cells[2], cells[3], cells[4], cells[5]

        if c5 == "MATIERES" and "PERIODE" in c9.upper().replace("É", "E"):
            i += 1
            continue

        # Ligne suite « TELETRAVAIL » : même module SIGFAE que la ligne précédente
        if c5.upper().strip() == "TELETRAVAIL" and not c6:
            pass  # ne pas écraser current_module
        elif c5 and len(c5) > 2 and c5 not in (
            "MATIERES",
            "FORMATION EN ADMINISTRATION DE BASE SESSION 2026",
        ):
            if "EMPLOI DU TEMPS" in c5.upper():
                i += 1
                continue
            raw_mod = c5
            current_module = raw_to_canonical_module(raw_mod)
            # volume total module souvent en col6 sur une ligne suivante
            if c6 and re.match(r"^\d+\s*H$", c6.replace(" ", "")):
                try:
                    if current_module:
                        module_volumes[current_module] = float(
                            c6.replace(" ", "").replace("H", "")
                        )
                except ValueError:
                    pass

        # ligne volume seule "16 H", "30 H" — rattacher au module courant
        if not c5 and c6 and re.match(r"^\d+\s*H$", c6.replace(" ", "")) and current_module:
            try:
                module_volumes[current_module] = float(c6.replace(" ", "").replace("H", ""))
            except ValueError:
                pass

        dt = parse_french_date(c9)
        t0, t1 = parse_horaire(c10)
        if dt and t0 and t1 and current_module:
            label = f"{t0.strftime('%H:%M')}-{t1.strftime('%H:%M')}"
            sessions.append((current_module, dt, t0, t1, label))
            # Colonnes 7–8 de l’EDT (bâtiment / salle) — une valeur par module (première renseignée)
            bat = cells[2].strip() if len(cells) > 2 else ""
            sal = cells[3].strip() if len(cells) > 3 else ""
            if bat and current_module not in module_batiment:
                module_batiment[current_module] = bat
            if sal and current_module not in module_salle:
                module_salle[current_module] = sal

        i += 1

    return sessions, module_volumes, module_batiment, module_salle


def module_date_ranges(sessions):
    by_mod: dict[str, list[date]] = defaultdict(list)
    for mod, d, *_ in sessions:
        by_mod[mod].append(d)
    ranges = {}
    for mod, dates in by_mod.items():
        ranges[mod] = (min(dates), max(dates))
    return ranges


def write_formations(
    wb: Workbook,
    ranges: dict[str, tuple[date, date]],
    volumes: dict[str, float],
    module_batiment: dict[str, str],
    module_salle: dict[str, str],
):
    ws = wb.create_sheet("Formations")
    headers = (
        "N°",
        "Formation",
        "Module (titre)",
        "Site",
        "Bâtiment",
        "Salle",
        "Date début",
        "Date fin",
        "Volume horaire (h)",
        "Catégorie",
        "Grade",
        "Groupe",
        "Vague",
    )
    ws.append(headers)
    num = 0
    # ordre identique à l'EDT source
    order = sorted(ranges.keys(), key=lambda m: ranges[m][0])
    for mod in order:
        num += 1
        d0, d1 = ranges[mod]
        vol = volumes.get(mod, 0)
        ws.append(
            [
                num,
                FORMATION_CYCLE,
                mod,
                SITE,
                module_batiment.get(mod, ""),
                module_salle.get(mod, ""),
                datetime.combine(d0, time(8, 0)),
                datetime.combine(d1, time(17, 0)),
                vol,
                CATEGORIE,
                GRADE,
                GROUPE,
                VAGUE,
            ]
        )


def write_seances(wb: Workbook, sessions: list[tuple[str, date, time, time, str]]):
    ws = wb.create_sheet("Séances")
    ws.append(
        [
            "module_titre",
            "grade",
            "groupe",
            "date_journee",
            "numero",
            "intitule",
            "heure_debut",
            "heure_fin",
        ]
    )
    # numéro de séance : par (module, date) incrémental
    per_day: dict[tuple[str, date], int] = defaultdict(int)
    for mod, d, h0, h1, label in sessions:
        per_day[(mod, d)] += 1
        n = per_day[(mod, d)]
        ws.append(
            [
                mod,
                GRADE,
                GROUPE,
                datetime.combine(d, time(0, 0)),
                n,
                label,
                h0,
                h1,
            ]
        )


def main():
    src = Path(
        sys.argv[1]
        if len(sys.argv) > 1
        else "/Users/tobidesis/Downloads/données/Emploi du Temps A4 GROUPE 8MAJ.xlsx"
    )
    if not src.exists():
        print("Fichier introuvable:", src, file=sys.stderr)
        sys.exit(1)

    wb_in = load_workbook(src, data_only=True)
    ws_in = wb_in.active
    rows = extract_rows(ws_in)
    wb_in.close()

    sessions, volumes, module_batiment, module_salle = parse_schedule(rows)
    if not sessions:
        print("Aucune séance extraite — vérifier le fichier source.", file=sys.stderr)
        sys.exit(2)

    ranges = module_date_ranges(sessions)

    # --- Fichier formations seul ---
    wf = Workbook()
    wf.remove(wf.active)
    write_formations(wf, ranges, volumes, module_batiment, module_salle)
    out_f = BACKEND / "import_formations_A4_GROUPE_08_MAJ.xlsx"
    wf.save(out_f)
    print("Écrit:", out_f)
    shutil.copyfile(out_f, BACKEND / "import_formations_A4_GROUPE_08.xlsx")
    print("Écrit:", BACKEND / "import_formations_A4_GROUPE_08.xlsx")

    # --- Fichier séances seul ---
    ws_only = Workbook()
    ws_only.remove(ws_only.active)
    write_seances(ws_only, sessions)
    out_s = BACKEND / "import_seances_A4_GROUPE_08_MAJ.xlsx"
    ws_only.save(out_s)
    print("Écrit:", out_s)
    shutil.copyfile(out_s, BACKEND / "import_seances_A4_GROUPE_08.xlsx")
    print("Écrit:", BACKEND / "import_seances_A4_GROUPE_08.xlsx")

    # --- Complet ---
    wc = Workbook()
    wc.remove(wc.active)
    write_formations(wc, ranges, volumes, module_batiment, module_salle)
    write_seances(wc, sessions)
    out_c = BACKEND / "import_A4_GROUPE_08_MAJ_complet.xlsx"
    wc.save(out_c)
    print("Écrit:", out_c)
    shutil.copyfile(out_c, BACKEND / "import_A4_GROUPE_08_complet.xlsx")
    print("Écrit:", BACKEND / "import_A4_GROUPE_08_complet.xlsx")
    print(f"Séances générées : {len(sessions)}, modules : {len(ranges)}")


if __name__ == "__main__":
    main()
