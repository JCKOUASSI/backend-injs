#!/usr/bin/env python3
"""
Analyse et corrige les fichiers Excel d'import (Cours + Séances) d'un dossier.

Usage :
  python scripts/fix_import_folder.py "/chemin/vers/dossier"
  → écrit les fichiers corrigés dans <dossier>/corrige/

Corrections appliquées :
  - Dates mal saisies (espaces, slash manquant)
  - Catégorie A/B/C → FAB A/B/C
  - Groupe normalisé (GROUPE 04 → GROUPE 4)
  - Suppression des lignes résiduelles en bas des feuilles
  - Séances : propagation grade/groupe/vague (cellules fusionnées Excel)
  - Séances : alignement grade/vague sur le fichier Cours
"""
from __future__ import annotations

import ast
import os
import re
import sys
from collections import Counter
from copy import copy
from datetime import datetime

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')

import django
django.setup()

from openpyxl import load_workbook
from formations.management.commands.import_excel import Command


def _is_formations_file(wb, filename):
    name = filename.lower()
    if 'formation' in name or 'cours' in name or 'modele formation' in name:
        return True
    if 'seance' in name or 'séance' in name or 'edt' in name:
        return False
    return 'Formations' in wb.sheetnames


def _is_seances_file(wb, filename):
    name = filename.lower()
    if 'seance' in name or 'séance' in name or 'edt' in name:
        return True
    if 'formation' in name or 'cours' in name:
        return False
    return any(s in wb.sheetnames for s in ('Séances', 'Seances'))


def _sheet(wb, prefer):
    for name in prefer:
        if name in wb.sheetnames:
            return wb[name]
    return wb.active


MODULE_TYPO_FIXES = {
    'ETHIQUE PUBLIQUE ET MUTTE CONTRE LA CORRUPTION': 'ETHIQUE PUBLIQUE ET LUTTE CONTRE LA CORRUPTION',
    'DROIT ADMINISTRATIVF': 'DROIT ADMINISTRATIF',
}


def _normalize_module_title(val):
    s = (val or '').strip().upper()
    return MODULE_TYPO_FIXES.get(s, (val or '').strip())


def _normalize_groupe(val):
    s = (val or '').strip().upper()
    m = re.fullmatch(r'GROUPE\s+0*(\d+)', s)
    if m:
        return f'GROUPE {m.group(1)}'
    return s if s else ''


def _cell_str(val):
    if val is None:
        return ''
    return str(val).strip()


def _is_junk_formation_row(data):
    return (
        not _cell_str(data.get('formation'))
        and not _cell_str(data.get('module'))
        and not data.get('date_debut')
        and not data.get('date_fin')
    )


def _format_date_for_excel(cmd, val):
    """Retourne une datetime naïve pour openpyxl ou None."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.replace(tzinfo=None) if val.tzinfo else val
    dt = cmd._parse_formation_datetime(val)
    if dt:
        return dt.replace(tzinfo=None)
    return None


def _build_modules_index(cmd, ws):
    """Index des modules du fichier Cours : clés de rapprochement séances."""
    index = {}
    combos = Counter()
    for _row_idx, data in cmd._rows(ws, sheet_type='formation'):
        if _is_junk_formation_row(data):
            continue
        module = _cell_str(data.get('module')) or _cell_str(data.get('formation'))
        grade = _cell_str(data.get('grade'))
        groupe = _normalize_groupe(data.get('groupe'))
        vague = _cell_str(data.get('vague'))
        if not module or not groupe:
            continue
        key = (module.upper(), groupe)
        index[key] = {'module': module, 'grade': grade, 'groupe': groupe, 'vague': vague}
        combos[(module.upper(), groupe, vague)] += 1
    return index, combos


def _extract_header_keys(ws, sheet_type='formation'):
    """Retourne les clés internes d'en-tête (même logique que import_excel._rows)."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    class _HeaderOnly:
        def __init__(self, header):
            self._data = [header]

        def iter_rows(self, values_only=True):
            return iter(self._data)

    logs = []

    class _LogCmd(Command):
        def __init__(self):
            super().__init__()
            self.stdout = type('L', (), {'write': lambda s, m: logs.append(str(m))})()

    lc = _LogCmd()
    list(lc._rows(_HeaderOnly(rows[0]), sheet_type=sheet_type))
    for log in logs:
        if '[DEBUG] Headers normalisés' in log:
            return ast.literal_eval(log.split(':', 1)[1].strip())
    return []


def _fix_formations_ws(cmd, ws, fixes_log):
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return 0
    header_display = list(all_rows[0])
    header_keys = _extract_header_keys(ws, 'formation')
    if not header_keys:
        return 0

    col_map = {k: i for i, k in enumerate(header_keys) if k}
    rows_out = [header_display]

    for row_idx, data in cmd._rows(ws, sheet_type='formation'):
        if _is_junk_formation_row(data):
            fixes_log.append(f'  L{row_idx}: ligne résiduelle supprimée')
            continue

        values = list(all_rows[row_idx - 1]) if row_idx - 1 < len(all_rows) else []
        while len(values) < len(header_keys):
            values.append(None)

        if 'categorie' in col_map and data.get('categorie'):
            new_cat = cmd._normalize_categorie(_cell_str(data['categorie']))
            if new_cat != _cell_str(data['categorie']):
                values[col_map['categorie']] = new_cat
                fixes_log.append(f'  L{row_idx}: catégorie → {new_cat}')

        if 'groupe' in col_map and data.get('groupe'):
            new_g = _normalize_groupe(data['groupe'])
            if new_g != _cell_str(data['groupe']):
                values[col_map['groupe']] = new_g
                fixes_log.append(f'  L{row_idx}: groupe → {new_g}')

        for date_key in ('date_debut', 'date_fin'):
            if date_key not in col_map:
                continue
            raw = data.get(date_key)
            parsed = _format_date_for_excel(cmd, raw)
            if parsed:
                if raw != parsed:
                    values[col_map[date_key]] = parsed
                    fixes_log.append(f'  L{row_idx}: {date_key} corrigée ({raw!r})')
            elif raw:
                normalized = cmd._normalize_date_string(raw)
                if normalized != _cell_str(raw):
                    values[col_map[date_key]] = normalized
                    fixes_log.append(f'  L{row_idx}: {date_key} normalisée ({raw!r} → {normalized})')

        rows_out.append(values)

    ws.delete_rows(1, ws.max_row)
    for r in rows_out:
        ws.append(list(r))
    return len(rows_out) - 1


def _fix_seances_ws(cmd, ws, modules_index, fixes_log):
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return 0
    header_keys = _extract_header_keys(ws, 'participant')
    if not header_keys:
        return 0
    col_map = {k: i for i, k in enumerate(header_keys) if k}
    rows_out = [list(all_rows[0])]
    last_grade = ''
    last_groupe = ''
    last_vague = ''

    for row_idx, data in cmd._rows(ws):
        values = list(all_rows[row_idx - 1]) if row_idx - 1 < len(all_rows) else []
        while len(values) < len(header_keys):
            values.append(None)
        module = _normalize_module_title(
            data.get('module_titre') or data.get('module') or data.get('formation')
        )
        if not module:
            fixes_log.append(f'  L{row_idx}: ligne sans module_titre supprimée')
            continue

        grade = _cell_str(data.get('grade'))
        groupe = _normalize_groupe(data.get('groupe'))
        vague = _cell_str(data.get('vague'))

        if not grade and last_grade:
            grade = last_grade
            if 'grade' in col_map:
                values[col_map['grade']] = grade
            fixes_log.append(f'  L{row_idx}: grade propagé → {grade}')
        if not groupe and last_groupe:
            groupe = last_groupe
            if 'groupe' in col_map:
                values[col_map['groupe']] = groupe
            fixes_log.append(f'  L{row_idx}: groupe propagé → {groupe}')
        if not vague and last_vague:
            vague = last_vague
            if 'vague' in col_map:
                values[col_map['vague']] = vague
            fixes_log.append(f'  L{row_idx}: vague propagée → {vague}')

        ref = modules_index.get((module.upper(), groupe))
        if ref:
            if ref['vague'] and vague != ref['vague']:
                fixes_log.append(
                    f'  L{row_idx}: vague {vague!r} → {ref["vague"]!r} (aligné sur Cours)'
                )
                vague = ref['vague']
                if 'vague' in col_map:
                    values[col_map['vague']] = vague
            if ref['grade'] and grade and grade != ref['grade']:
                fixes_log.append(
                    f'  L{row_idx}: grade {grade!r} → {ref["grade"]!r} (aligné sur Cours)'
                )
                grade = ref['grade']
                if 'grade' in col_map:
                    values[col_map['grade']] = grade
            elif ref['grade'] and not grade:
                grade = ref['grade']
                if 'grade' in col_map:
                    values[col_map['grade']] = grade
                fixes_log.append(f'  L{row_idx}: grade aligné sur Cours → {grade}')

        if grade:
            last_grade = grade
        if groupe:
            last_groupe = groupe
        if vague:
            last_vague = vague

        if 'module_titre' in col_map:
            values[col_map['module_titre']] = module
        if 'groupe' in col_map:
            values[col_map['groupe']] = groupe
        if 'date_journee' in col_map and data.get('date_journee'):
            parsed = cmd._parse_date(data.get('date_journee'))
            if parsed:
                values[col_map['date_journee']] = parsed.strftime('%d/%m/%Y')

        rows_out.append(values)

    ws.delete_rows(1, ws.max_row)
    for r in rows_out:
        ws.append(list(r))
    return len(rows_out) - 1


def _simulate_import(cmd, wb_path, file_type):
    wb = load_workbook(wb_path, read_only=True)
    cmd.stdout = type('S', (), {'write': lambda self, m: None})()
    errors = []
    if file_type == 'formations':
        ws = _sheet(wb, ['Formations'])
        created, updated = cmd._import_formations(ws, errors)
        wb.close()
        return errors, {'created': created, 'updated': updated}
    if file_type == 'seances':
        ws = _sheet(wb, ['Séances', 'Seances'])
        created, updated = cmd._import_seances(ws, errors)
        wb.close()
        return errors, {'created': created, 'updated': updated}
    wb.close()
    return [], {}


def process_folder(folder):
    folder = os.path.abspath(folder)
    if not os.path.isdir(folder):
        print(f'❌ Dossier introuvable : {folder}')
        sys.exit(1)

    out_dir = os.path.join(folder, 'corrige')
    os.makedirs(out_dir, exist_ok=True)
    cmd = Command()
    cmd.stdout = type('S', (), {'write': lambda self, m: None})()

    xlsx_files = sorted(
        f for f in os.listdir(folder)
        if f.lower().endswith('.xlsx') and not f.startswith('~$')
    )
    if not xlsx_files:
        print(f'❌ Aucun fichier .xlsx dans {folder}')
        sys.exit(1)

    print('=' * 60)
    print('CORRECTION FICHIERS IMPORT')
    print('=' * 60)
    print(f'Dossier  : {folder}')
    print(f'Fichiers : {len(xlsx_files)}')
    print(f'Sortie   : {out_dir}/\n')

    modules_index = {}
    formations_path = None

    # 1) Corriger d'abord les fichiers Cours
    for fname in xlsx_files:
        path = os.path.join(folder, fname)
        wb = load_workbook(path)
        if not _is_formations_file(wb, fname):
            wb.close()
            continue

        print(f'📗 COURS : {fname}')
        ws = _sheet(wb, ['Formations'])
        fixes = []
        n = _fix_formations_ws(cmd, ws, fixes)
        out_path = os.path.join(out_dir, fname.replace('.xlsx', '_corrige.xlsx'))
        wb.save(out_path)
        wb.close()

        wb_idx = load_workbook(out_path, read_only=True)
        idx, _ = _build_modules_index(cmd, _sheet(wb_idx, ['Formations']))
        wb_idx.close()
        modules_index.update(idx)
        formations_path = out_path

        for line in fixes[:30]:
            print(line)
        if len(fixes) > 30:
            print(f'  … et {len(fixes) - 30} autres corrections')
        print(f'  → {n} lignes conservées → {out_path}\n')

    if modules_index:
        print(f'Index modules Cours : {len(modules_index)} entrées (module+groupe)\n')

    # 2) Corriger les fichiers Séances
    total_errors = 0
    for fname in xlsx_files:
        path = os.path.join(folder, fname)
        wb = load_workbook(path)
        if not _is_seances_file(wb, fname):
            wb.close()
            continue

        print(f'📘 SÉANCES : {fname}')
        ws = _sheet(wb, ['Séances', 'Seances'])
        fixes = []
        n = _fix_seances_ws(cmd, ws, modules_index, fixes)
        out_path = os.path.join(out_dir, fname.replace('.xlsx', '_corrige.xlsx'))
        wb.save(out_path)
        wb.close()

        for line in fixes[:30]:
            print(line)
        if len(fixes) > 30:
            print(f'  … et {len(fixes) - 30} autres corrections')
        print(f'  → {n} lignes conservées → {out_path}')

        errors, stats = _simulate_import(cmd, out_path, 'seances')
        print(f'  Simulation import : {stats}, {len(errors)} erreur(s)')
        for e in errors[:10]:
            print(f'    • {e}')
        total_errors += len(errors)
        print()

    # 3) Vérifier formations corrigées
    if formations_path:
        errors, stats = _simulate_import(cmd, formations_path, 'formations')
        print(f'📗 Vérif. Cours corrigé : {stats}, {len(errors)} erreur(s)')
        for e in errors[:10]:
            print(f'  • {e}')
        total_errors += len(errors)

    print('=' * 60)
    if total_errors == 0:
        print('✅ Tous les fichiers corrigés sont prêts pour import.')
        print(f'   Importez les fichiers du dossier : {out_dir}/')
    else:
        print(f'⚠️  {total_errors} erreur(s) restantes — voir détails ci-dessus.')
    print('=' * 60)
    sys.exit(1 if total_errors else 0)


def main():
    if len(sys.argv) < 2:
        print('Usage: python scripts/fix_import_folder.py <dossier>')
        sys.exit(1)
    process_folder(sys.argv[1])


if __name__ == '__main__':
    main()
