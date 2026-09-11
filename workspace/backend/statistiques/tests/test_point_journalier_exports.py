"""Tests exports point journalier — alignement données / Excel."""
from datetime import time
from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from formations.models import Formation, Module, ModuleParticipant, Participant, SessionModule
from statistiques.point_journalier import compute_point_journalier
from statistiques.point_journalier_exports import export_excel, export_pdf, _filter_tableaux


class PointJournalierExportsTest(TestCase):
    def setUp(self):
        self.formation = Formation.objects.create(formation='Formation en Administration de Base')
        self.jour = timezone.localdate()
        self.participant = Participant.objects.create(
            matricule='EXP-PJ-001', nom='Export', prenom='Test', categorie='A',
        )

        self.mod_cours = Module.objects.create(
            formation=self.formation, intitule='DEONTOLOGIE', grade='A4',
            groupe='GROUPE 10', salle='SALLE A',
        )
        self.mod_sans = Module.objects.create(
            formation=self.formation, intitule='CULTURE CIVIQUE', grade='A4',
            groupe='GROUPE 11', salle='SALLE B',
        )
        for mod in (self.mod_cours, self.mod_sans):
            ModuleParticipant.objects.create(module=mod, participant=self.participant)
        SessionModule.objects.create(
            module=self.mod_cours,
            date_journee=self.jour,
            numero=1,
            intitule='Matin',
            heure_debut_prevue=time(8, 0),
            heure_fin_prevue=time(12, 0),
        )

    def _tableaux_export(self):
        data = compute_point_journalier(
            annee=self.jour.year,
            mois=self.jour.month,
            categorie='A',
            formation_id=self.formation.id,
            jour=self.jour.isoformat(),
        )
        return _filter_tableaux(
            data,
            formation_id=self.formation.id,
            categorie='A',
            jour=self.jour.isoformat(),
        )

    def test_export_excel_uniquement_groupes_avec_cours(self):
        tableaux = self._tableaux_export()
        self.assertEqual(len(tableaux), 1)

        buffer = export_excel(tableaux, self.jour.year)
        wb = load_workbook(BytesIO(buffer.getvalue()))
        ws = wb.active

        groupes_row = None
        for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
            if row and row[0] == 'GROUPES':
                groupes_row = [c for c in row[1:] if c]
                break

        self.assertIsNotNone(groupes_row)
        labels = [str(c) for c in groupes_row if c and str(c) != 'TOTAL']
        self.assertIn('G10', labels)
        self.assertNotIn('G11', labels)
        self.assertEqual(len(labels), 1)

    def test_export_pdf_ne_leve_pas_avec_creneau_soir_vide(self):
        tableaux = self._tableaux_export()
        buffer = export_pdf(tableaux, self.jour.year)
        self.assertGreater(len(buffer.getvalue()), 100)

    def test_build_export_meme_donnees_que_api_detail(self):
        tableaux = self._tableaux_export()
        tb = tableaux[0]
        self.assertEqual(len(tb['matin']['groupes']), 1)
        self.assertEqual(len(tb['soir']['groupes']), 0)
        self.assertEqual(tb['matin']['groupes'][0]['label'], 'G10')
