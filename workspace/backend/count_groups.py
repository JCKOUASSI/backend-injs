import openpyxl, os

SRC = '/Users/tobidesis/Downloads/données'
fname = 'FAB 2026 PHASE 1 FICHIER COMPLET.xlsx'
wb = openpyxl.load_workbook(os.path.join(SRC, fname), read_only=True, data_only=True)

groups = {}
for shname in wb.sheetnames:
    if shname in ('TCD', 'RECAP'):
        continue
    ws = wb[shname]
    all_rows = list(ws.iter_rows(values_only=True, max_row=3))
    if len(all_rows) < 2:
        continue
    headers = [str(h).upper().strip().replace('\n',' ') if h else '' for h in all_rows[0]]
    row = all_rows[1]

    def gcol(pat):
        for i, h in enumerate(headers):
            if pat in h:
                return i
        return None

    ig = gcol('GRADE')
    igr = gcol('GROUPE')
    grade  = str(row[ig]).strip()  if ig  is not None and ig  < len(row) and row[ig]  else '?'
    groupe = str(row[igr]).strip() if igr is not None and igr < len(row) and row[igr] else '?'
    nb = sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
             if r and any(v for v in r))
    groups.setdefault(grade, []).append((groupe, nb, shname))

wb.close()

print("RÉSUMÉ PAR GRADE :")
total_groupes = 0
total_parts = 0
for grade in sorted(groups.keys()):
    glist = sorted(groups[grade])
    sous_total = sum(n for _, n, _ in glist)
    total_parts += sous_total
    total_groupes += len(glist)
    print(f"\n  Grade {grade} — {len(glist)} groupes — {sous_total} participants")
    for groupe, nb, sh in glist:
        print(f"    {groupe:15s} ({sh:8s}) : {nb:4d} participants")

print(f"\n{'='*50}")
print(f"  TOTAL : {total_groupes} groupes — {total_parts} participants")
