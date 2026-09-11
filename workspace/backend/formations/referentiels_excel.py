"""Import et export Excel des référentiels administratifs."""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import (
    RefBatiment, RefCategorie, RefFormation, RefGrade, RefModule,
    RefModuleVolumeHoraire, RefSalle, RefSite, RefTypeSecretariat, RefVague,
)


class ReferentielExcelError(ValueError):
    pass


SPECS = {
    'formations': ('Formations', ['intitule', 'prix_heure_realisee', 'actif']),
    'modules': ('Modules', ['intitule', 'formation', 'categorie', 'volume_horaire', 'actif']),
    'categories': ('Categories', ['libelle', 'actif']),
    'grades': ('Grades', ['libelle', 'categorie', 'actif']),
    'vagues': ('Vagues', ['libelle', 'ordre', 'actif']),
    'sites': ('Sites', ['nom', 'geofence_latitude', 'geofence_longitude', 'geofence_rayon_m', 'actif']),
    'batiments': ('Batiments', ['nom', 'site', 'actif']),
    'salles': ('Salles', ['nom', 'site', 'batiment', 'type_lieu', 'capacite', 'equipements', 'actif']),
    'types_secretariat': ('Types secretariat', ['libelle', 'actif']),
}


def _text(value):
    return '' if value is None else str(value).strip()


def _active(value, default=True):
    if value in (None, ''):
        return default
    normalized = _text(value).lower()
    if normalized in ('1', 'oui', 'o', 'true', 'vrai', 'actif', 'active'):
        return True
    if normalized in ('0', 'non', 'n', 'false', 'faux', 'inactif', 'inactive'):
        return False
    raise ReferentielExcelError('La colonne « actif » accepte oui/non ou vrai/faux.')


def _number(value, label, *, integer=False, nullable=True):
    if value in (None, ''):
        if nullable:
            return None
        raise ReferentielExcelError(f'La colonne « {label} » est obligatoire.')
    try:
        parsed = Decimal(str(value).replace(',', '.'))
    except (InvalidOperation, ValueError):
        raise ReferentielExcelError(f'La colonne « {label} » doit être numérique.')
    if parsed < 0:
        raise ReferentielExcelError(f'La colonne « {label} » ne peut pas être négative.')
    return int(parsed) if integer else parsed


def _rows_from_workbook(uploaded, kind):
    try:
        workbook = load_workbook(uploaded, data_only=True)
    except Exception as exc:
        raise ReferentielExcelError('Fichier Excel (.xlsx) invalide.') from exc
    sheet = workbook.active
    expected = SPECS[kind][1]
    headers = [_text(cell.value).lower() for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())]
    missing = [column for column in expected if column not in headers]
    if missing:
        raise ReferentielExcelError(f'Colonnes manquantes : {", ".join(missing)}.')
    indices = {header: headers.index(header) for header in expected}
    rows = []
    for line, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {column: cells[index] if index < len(cells) else None for column, index in indices.items()}
        if not any(value not in (None, '') for value in row.values()):
            continue
        row['_line'] = line
        rows.append(row)
    workbook.close()
    return rows


def _required(row, column):
    value = _text(row.get(column))
    if not value:
        raise ReferentielExcelError(f'Ligne {row["_line"]} : « {column} » est obligatoire.')
    return value


def _by_label(model, field, value, row, label):
    obj = model.objects.filter(**{f'{field}__iexact': value}).first()
    if not obj:
        raise ReferentielExcelError(f'Ligne {row["_line"]} : {label} « {value} » introuvable.')
    return obj


def import_referentiel_xlsx(kind, uploaded):
    if kind not in SPECS:
        raise ReferentielExcelError('Référentiel inconnu.')
    rows = _rows_from_workbook(uploaded, kind)
    created = updated = 0
    with transaction.atomic():
        for row in rows:
            try:
                was_created = _import_row(kind, row)
            except ReferentielExcelError:
                raise
            except Exception as exc:
                raise ReferentielExcelError(f'Ligne {row["_line"]} : {exc}') from exc
            if was_created:
                created += 1
            else:
                updated += 1
    return {'created': created, 'updated': updated, 'processed': len(rows)}


def _import_row(kind, row):
    actif = _active(row.get('actif'))
    if kind == 'formations':
        intitule = _required(row, 'intitule')
        obj, created = RefFormation.objects.get_or_create(intitule=intitule)
        obj.actif = actif
        obj.prix_heure_realisee = _number(row.get('prix_heure_realisee'), 'prix_heure_realisee')
        obj.save()
        return created
    if kind == 'categories':
        obj, created = RefCategorie.objects.get_or_create(libelle=_required(row, 'libelle'))
        obj.actif = actif
        obj.save()
        return created
    if kind == 'types_secretariat':
        obj, created = RefTypeSecretariat.objects.get_or_create(libelle=_required(row, 'libelle'))
        obj.actif = actif
        obj.save()
        return created
    if kind == 'vagues':
        obj, created = RefVague.objects.get_or_create(libelle=_required(row, 'libelle'))
        obj.actif = actif
        obj.ordre = _number(row.get('ordre'), 'ordre', integer=True, nullable=False)
        obj.save()
        return created
    if kind == 'sites':
        obj, created = RefSite.objects.get_or_create(nom=_required(row, 'nom'))
        obj.actif = actif
        obj.geofence_latitude = _number(row.get('geofence_latitude'), 'geofence_latitude')
        obj.geofence_longitude = _number(row.get('geofence_longitude'), 'geofence_longitude')
        obj.geofence_rayon_m = _number(row.get('geofence_rayon_m'), 'geofence_rayon_m', integer=True) or 200
        obj.save()
        return created
    if kind == 'batiments':
        site = _by_label(RefSite, 'nom', _required(row, 'site'), row, 'Site')
        obj, created = RefBatiment.objects.get_or_create(site=site, nom=_required(row, 'nom'))
        obj.actif = actif
        obj.save()
        return created
    if kind == 'grades':
        categorie_label = _text(row.get('categorie'))
        categorie = _by_label(RefCategorie, 'libelle', categorie_label, row, 'Catégorie') if categorie_label else None
        obj, created = RefGrade.objects.get_or_create(categorie=categorie, libelle=_required(row, 'libelle'))
        obj.actif = actif
        obj.save()
        return created
    if kind == 'salles':
        site = _by_label(RefSite, 'nom', _required(row, 'site'), row, 'Site')
        batiment_label = _text(row.get('batiment'))
        batiment = None
        if batiment_label:
            batiment = RefBatiment.objects.filter(site=site, nom__iexact=batiment_label).first()
            if not batiment:
                raise ReferentielExcelError(f'Ligne {row["_line"]} : bâtiment « {batiment_label} » introuvable sur ce site.')
        nom = _required(row, 'nom')
        obj, created = RefSalle.objects.get_or_create(site=site, batiment=batiment, nom=nom)
        type_lieu = _text(row.get('type_lieu')).upper() or RefSalle.TypeLieu.SALLE
        if type_lieu not in RefSalle.TypeLieu.values:
            raise ReferentielExcelError(f'Ligne {row["_line"]} : type_lieu invalide.')
        obj.type_lieu = type_lieu
        obj.capacite = _number(row.get('capacite'), 'capacite', integer=True)
        obj.equipements = _text(row.get('equipements'))[:255]
        obj.actif = actif
        obj.save()
        return created
    # Modules : une ligne par association formation × catégorie.
    intitule = RefModule.normalize_intitule(_required(row, 'intitule'))
    formation = _by_label(RefFormation, 'intitule', _required(row, 'formation'), row, 'Formation')
    obj, created = RefModule.objects.get_or_create(intitule=intitule)
    obj.actif = actif
    obj.save()
    obj.formations.add(formation)
    categorie_label = _text(row.get('categorie'))
    volume = _number(row.get('volume_horaire'), 'volume_horaire')
    if bool(categorie_label) != (volume is not None):
        raise ReferentielExcelError(f'Ligne {row["_line"]} : catégorie et volume_horaire doivent être renseignés ensemble.')
    if categorie_label:
        categorie = _by_label(RefCategorie, 'libelle', categorie_label, row, 'Catégorie')
        RefModuleVolumeHoraire.objects.update_or_create(
            module=obj, formation=formation, categorie=categorie,
            defaults={'volume_horaire': volume},
        )
    return created


def export_referentiel_xlsx(kind):
    if kind not in SPECS:
        raise ReferentielExcelError('Référentiel inconnu.')
    title, headers = SPECS[kind]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='163D6A')
    for row in _export_rows(kind):
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(len(_text(cell.value)) for cell in column) + 2, 45)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _export_rows(kind):
    if kind == 'formations':
        return ((o.intitule, o.prix_heure_realisee, o.actif) for o in RefFormation.objects.order_by('intitule'))
    if kind == 'categories':
        return ((o.libelle, o.actif) for o in RefCategorie.objects.order_by('libelle'))
    if kind == 'types_secretariat':
        return ((o.libelle, o.actif) for o in RefTypeSecretariat.objects.order_by('libelle'))
    if kind == 'vagues':
        return ((o.libelle, o.ordre, o.actif) for o in RefVague.objects.order_by('ordre', 'libelle'))
    if kind == 'sites':
        return ((o.nom, o.geofence_latitude, o.geofence_longitude, o.geofence_rayon_m, o.actif) for o in RefSite.objects.order_by('nom'))
    if kind == 'batiments':
        return ((o.nom, o.site.nom, o.actif) for o in RefBatiment.objects.select_related('site').order_by('nom'))
    if kind == 'grades':
        return ((o.libelle, o.categorie.libelle if o.categorie else '', o.actif) for o in RefGrade.objects.select_related('categorie').order_by('libelle'))
    if kind == 'salles':
        return ((o.nom, o.site.nom, o.batiment.nom if o.batiment else '', o.type_lieu, o.capacite, o.equipements, o.actif)
                for o in RefSalle.objects.select_related('site', 'batiment').order_by('nom'))
    rows = []
    for module in RefModule.objects.prefetch_related('formations', 'volumes_horaires__formation', 'volumes_horaires__categorie').order_by('intitule'):
        volumes = list(module.volumes_horaires.all())
        if volumes:
            rows.extend((module.intitule, v.formation.intitule, v.categorie.libelle, v.volume_horaire, module.actif) for v in volumes)
        else:
            rows.extend((module.intitule, f.intitule, '', '', module.actif) for f in module.formations.all())
    return rows
