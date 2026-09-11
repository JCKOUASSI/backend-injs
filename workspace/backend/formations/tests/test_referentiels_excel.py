from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from formations.models import RefCategorie, RefFormation, RefModule
from formations.referentiels_excel import export_referentiel_xlsx, import_referentiel_xlsx


class ReferentielsExcelTests(TestCase):
    def _workbook(self, headers, row):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(row)
        stream = BytesIO()
        workbook.save(stream)
        return SimpleUploadedFile('referentiel.xlsx', stream.getvalue())

    def test_import_categories_upserts_by_label(self):
        uploaded = self._workbook(['libelle', 'actif'], ['A', 'non'])
        result = import_referentiel_xlsx('categories', uploaded)
        self.assertEqual(result, {'created': 1, 'updated': 0, 'processed': 1})
        self.assertFalse(RefCategorie.objects.get(libelle='A').actif)

        uploaded = self._workbook(['libelle', 'actif'], ['A', 'oui'])
        result = import_referentiel_xlsx('categories', uploaded)
        self.assertEqual(result, {'created': 0, 'updated': 1, 'processed': 1})
        self.assertTrue(RefCategorie.objects.get(libelle='A').actif)

    def test_import_modules_uses_formation_and_category_labels(self):
        RefFormation.objects.create(intitule='LICENCE STAPS')
        RefCategorie.objects.create(libelle='A')
        uploaded = self._workbook(
            ['intitule', 'formation', 'categorie', 'volume_horaire', 'actif'],
            ['Anatomie', 'LICENCE STAPS', 'A', 24, 'oui'],
        )
        import_referentiel_xlsx('modules', uploaded)
        module = RefModule.objects.get(intitule='ANATOMIE')
        self.assertEqual(module.formations.first().intitule, 'LICENCE STAPS')
        self.assertEqual(float(module.volumes_horaires.get().volume_horaire), 24.0)

    def test_export_generates_openable_xlsx_with_headers(self):
        RefCategorie.objects.create(libelle='A')
        content = export_referentiel_xlsx('categories')
        workbook = load_workbook(BytesIO(content), data_only=True)
        self.assertEqual([cell.value for cell in workbook.active[1]], ['libelle', 'actif'])
        self.assertEqual(workbook.active['A2'].value, 'A')
